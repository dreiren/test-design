#!/usr/bin/env python3
"""Generate Veriff biometric authentication rewards estimation workbook."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.page import PageMargins

OUT = "/workspace/estimates/Veriff_Biometric_Auth_Rewards_Estimation.xlsx"

# Colors
NAVY = "1B365D"
NAVY2 = "2E5090"
WHITE = "FFFFFF"
TEAL = "0F6C6C"
GOLD = "B8860B"
P0 = "D6DCE4"
P1 = "E2EFDA"
P2 = "FFF2CC"
P3 = "E2D5F1"
OPS = "DDEBF7"
LEGAL = "FCE4D6"
TOTAL = "1B365D"
ALT = "F7F9FC"
RED_SOFT = "F4CCCC"
INPUT = "FFF8E7"
COVER = "F4F7FB"
LINK_BLUE = "1F4E79"

thin = Border(
    left=Side(style="thin", color="C5CDD6"),
    right=Side(style="thin", color="C5CDD6"),
    top=Side(style="thin", color="C5CDD6"),
    bottom=Side(style="thin", color="C5CDD6"),
)
thick_bottom = Border(
    left=Side(style="thin", color="C5CDD6"),
    right=Side(style="thin", color="C5CDD6"),
    top=Side(style="thin", color="C5CDD6"),
    bottom=Side(style="medium", color=NAVY),
)

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_navy2 = PatternFill("solid", fgColor=NAVY2)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_p0 = PatternFill("solid", fgColor=P0)
fill_p1 = PatternFill("solid", fgColor=P1)
fill_p2 = PatternFill("solid", fgColor=P2)
fill_p3 = PatternFill("solid", fgColor=P3)
fill_ops = PatternFill("solid", fgColor=OPS)
fill_legal = PatternFill("solid", fgColor=LEGAL)
fill_total = PatternFill("solid", fgColor=TOTAL)
fill_alt = PatternFill("solid", fgColor=ALT)
fill_cover = PatternFill("solid", fgColor=COVER)
fill_input = PatternFill("solid", fgColor=INPUT)
fill_gold = PatternFill("solid", fgColor="FFF3CD")
fill_red = PatternFill("solid", fgColor=RED_SOFT)

font_white_bold = Font(name="Calibri", bold=True, color=WHITE, size=11)
font_white_title = Font(name="Calibri", bold=True, color=WHITE, size=16)
font_white_sub = Font(name="Calibri", color="D6E3F0", size=11)
font_navy_title = Font(name="Calibri", bold=True, color=NAVY, size=16)
font_navy_h = Font(name="Calibri", bold=True, color=NAVY, size=12)
font_navy = Font(name="Calibri", color=NAVY, size=11)
font_bold = Font(name="Calibri", bold=True, size=11, color="1A1A1A")
font_normal = Font(name="Calibri", size=10, color="1A1A1A")
font_small = Font(name="Calibri", size=9, italic=True, color="4A5568")
font_section = Font(name="Calibri", bold=True, size=11, color=NAVY)
font_total = Font(name="Calibri", bold=True, size=11, color=WHITE)
font_num = Font(name="Calibri", size=10, color="1A1A1A")

wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
wrap_c = Alignment(wrap_text=True, vertical="center", horizontal="center")
wrap_m = Alignment(wrap_text=True, vertical="center", horizontal="left")
center = Alignment(vertical="center", horizontal="center")
left_c = Alignment(vertical="center", horizontal="left")

PHASE_FILL = {
    "P0": fill_p0,
    "P1": fill_p1,
    "P2": fill_p2,
    "P3": fill_p3,
    "OPS": fill_ops,
    "LEGAL": fill_legal,
}


def apply_sheet_defaults(ws, freeze="A8"):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:7"
    ws.oddHeader.left.text = "Veriff Biometric Auth — Rewards Redemption Estimation"
    ws.oddFooter.left.text = "Confidential | Solution architecture estimate | Not a statement of work"
    ws.oddFooter.right.text = "Page &P of &N"


def banner(ws, title, subtitle, cols=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    c1 = ws.cell(1, 1, title)
    c1.font = font_white_title
    c1.fill = fill_navy
    c1.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    c2 = ws.cell(2, 1, subtitle)
    c2.font = font_white_sub
    c2.fill = fill_navy2
    c2.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=True)
    for col in range(1, cols + 1):
        ws.cell(1, col).fill = fill_navy
        ws.cell(2, col).fill = fill_navy2
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 36


def note_row(ws, row, text, cols=8, fill=fill_gold):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row, 1, text)
    c.font = font_small
    c.fill = fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
    for col in range(1, cols + 1):
        ws.cell(row, col).fill = fill
    ws.row_dimensions[row].height = 48


def header_row(ws, row, headers, fills=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.font = font_white_bold
        c.fill = fill_navy
        c.alignment = wrap_c
        c.border = thin
    ws.row_dimensions[row].height = 24
    ws.auto_filter.ref = None  # set later


def section_row(ws, row, text, cols, fill):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row, 1, text)
    c.font = font_section
    c.fill = fill
    c.alignment = left_c
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.border = thin
    ws.row_dimensions[row].height = 22
    return row


def data_row(ws, row, values, fills=None, heights=72):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font = font_normal if i <= 2 or not isinstance(v, (int, float)) else font_num
        c.alignment = wrap if i <= 2 else center
        c.border = thin
        if fills and i in fills:
            c.fill = fills[i]
        else:
            c.fill = fill_alt if row % 2 == 0 else fill_white
        if isinstance(v, (int, float)) and i >= 3:
            c.number_format = "0.0"
            c.alignment = center
            c.font = font_num
    ws.row_dimensions[row].height = heights
    return row


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# TAB 1 work items
# Each tuple: (section_key, section_title, section_fill_key, name, desc, min, mid, max)
# ---------------------------------------------------------------------------
MS_SECTIONS = []

tab1 = []

# helper to push
def T1(section, item, desc, mn, mid, mx, fill_key="P1"):
    tab1.append({
        "section": section,
        "item": item,
        "desc": desc,
        "min": mn,
        "mid": mid,
        "max": mx,
        "fill": fill_key,
    })


# === P0 ===
S0 = "P0 — Discovery, vendor, legal & policy (before build)"
T1(S0, "[P0] Solution architecture, identity binding & threat model",
   "Design the multi-brand architecture: Brand sites → Bio-Auth microservice → Veriff (Enrollment integration + Authentication integration). Bind winner_id ↔ logged-in account_id ↔ stable Veriff endUserId (UUID v4) or vendorData ↔ reward_id. Threat model: account takeover after win, prize-link forwarding, selfie spoof / presentation attack, unenrolled users getting Veriff decline 122, webhook spoofing, race conditions (double fulfill), warehouse shipping before decision, small-voucher cost vs fraud. Produce sequence diagrams, state machine, and a policy matrix (which prizes require bio). Deliverable: architecture decision record.",
   16, 24, 40, "P0")

T1(S0, "[P0] Veriff commercial onboarding & dual-integration setup",
   "Contract sandbox + production. Veriff Solutions Engineer must create and connect: (1) Selfie enrollment integration (or Document+Selfie IDV used as enrollment source) and (2) Biometric Authentication integration. Obtain API keys, baseURLs, webhook signing secrets, Customer Portal access, and test-session forcing. Confirm session TTL (7-day expired/abandoned), decision codes, and FaceBlock availability. This is on the critical path — development cannot complete without it.",
   8, 16, 32, "P0")

T1(S0, "[P0] Privacy / legal: biometric DPIA, consent, retention, contest rules",
   "Biometric templates are special-category data in many markets (GDPR Art. 9, similar APAC/LATAM rules). Produce DPIA, purpose limitation (rewards redemption only), consent copy + versioning, retention vs fraud-hold, processor DPA with Veriff, and contest/T&Cs updates stating biometric step-up for defined prizes. Align on lawful basis per market. Hours cover our privacy engineering + legal liaison, not client counsel of record.",
   16, 32, 56, "LEGAL")

T1(S0, "[P0] Business policy workshop: prize taxonomy, thresholds, exceptions, SLA",
   "Workshop with brand, fraud, CS, warehouse, finance. Agree: (a) physical prizes always require bio in MVP; (b) digital voucher high-value threshold per brand/currency; (c) small vouchers stay on existing controls; (d) claim window; (e) retry limits; (f) no-camera / accessibility exceptions; (g) decline = do not fulfill, not 'try a friend'; (h) manual override dual-control rules. Output: signed policy table that the microservice will enforce so websites cannot locally bypass.",
   8, 16, 24, "P0")

T1(S0, "[P0] Cross-brand identity & guest-winner model",
   "Decide how endUserId is minted and reused across brand websites. Guest checkout winners are the weak point: if a fraudster can enroll first, selfie-only bio only proves the same fraudster came back. Require registered, authenticated winner accounts for bio-gated prizes. Map existing user IDs that are not UUID v4 onto vendorData (immutable). Plan collision, brand-prefixed IDs vs global person ID, and what happens if the same person wins on two brands.",
   8, 16, 24, "P0")

# === P1 engineering ===
S1 = "P1 MVP — Microservice engineering (physical + high-value digital gate)"
T1(S1, "[P1] Service bootstrap, CI/CD, environments, multi-tenant brand config",
   "New service: repo, pipelines, lint/test, sandbox/staging/prod, infra (API gateway, secrets, logs). Brand registry: brandId, callback allowlist, threshold config, prize-type flags, Veriff integration IDs, languages, claim SLA. Websites must not hold Veriff API keys — only the microservice does.",
   16, 24, 40, "P1")

T1(S1, "[P1] Veriff credential isolation (enrollment keys ≠ authentication keys)",
   "Store and use separate API keys / baseURLs for enrollment vs authentication integrations. Rotate secrets, restrict by env, never expose to browsers. Include X-AUTH-CLIENT handling and webhook signature verification secrets per integration.",
   8, 12, 20, "P1")

T1(S1, "[P1] Enrollment session orchestration API",
   "POST create-enrollment: validate brand + authenticated user + reward still claimable + consent recorded. Call Veriff POST /v1/sessions on the enrollment integration with the same endUserId/vendorData and a brand callback URL. Persist sessionId, sessionToken metadata, reward_id, purpose=enrollment. Return session URL to the website (InContext or redirect). Do not treat session creation as success.",
   16, 24, 40, "P1")

T1(S1, "[P1] Biometric authentication session orchestration API",
   "POST create-auth only if an approved enrollment exists for this endUserId. Use authentication integration keys. Pass identical endUserId/vendorData (mismatch → Veriff decline 121/128). Bind session to a specific reward_id and redemption_attempt_id so an approved selfie cannot be reused to claim a different prize. Return session URL. Enforce one active session per reward.",
   16, 24, 40, "P1")

T1(S1, "[P1] Enrollment pre-check (avoid Veriff 'no reference' declines)",
   "GET enrollment-status before creating an auth session. If not enrolled, return action=ENROLL instead of starting auth. Prevents wasted Veriff cost and user-facing decline 122 (no reference found). Cache approved enrollment sessionId for matchedSession checks later.",
   8, 12, 20, "P1")

T1(S1, "[P1] Decision webhook receiver + signature verification + idempotency",
   "Public endpoint for Veriff decision webhooks (enrollment and auth). Verify authenticity (HMAC/signature headers per Veriff webhooks guide), reject unsigned payloads, handle retries, and upsert by sessionId/attemptId. Map verification.status (approved/declined/expired/abandoned/resubmission_requested/review) and biometricAuthentication.matchedSessionEndUserId. Webhook is the source of truth — never fulfill from the frontend callback.",
   16, 24, 40, "P1")

T1(S1, "[P1] Decision polling fallback",
   "If webhook is delayed/missed, poll GET /v1/sessions/{id}/decision with backoff until terminal status or TTL. Used by status API when the website shows 'pending'. Include alerting when webhook and poll disagree.",
   8, 12, 20, "P1")

T1(S1, "[P1] Session & redemption state machine",
   "States: NotRequired | ConsentNeeded | NeedsEnrollment | EnrollmentPending | Enrolled | AuthPending | Approved | Declined | FrozenFraud | ExceptionReview | ExpiredClaim | Fulfilled. Transitions are server-side only. Approved enrollment in the same sitting may release the current prize (do not force a second paid auth session immediately). Later redemptions require auth. Expire open holds at claim SLA.",
   12, 20, 32, "P1")

T1(S1, "[P1] Veriff reason-code → business outcome mapping",
   "Map codes: 9001 approved → fulfill gate; 9102 declined; 105/106 suspicious/known fraud → freeze + fraud queue; 120 face mismatch → freeze suspected third-party redeem; 121/128 ID mismatch → engineering incident, do not fulfill; 122 no reference → route to enrollment; expired/abandoned 9104 → retry until SLA; quality/resubmission → guided retry. User-facing copy is generic; ops sees raw codes.",
   8, 16, 24, "P1")

T1(S1, "[P1] Policy / threshold engine (central, not per website)",
   "Input: brandId, prizeType (physical|digital_voucher), faceValue, currency, market. Output: BioRequired=true/false and method=Auth|Enroll|IDV. MVP defaults: physical = required; digital >= HighThreshold = required; digital < HighThreshold = not required. Config-driven so Phase 2 can lower thresholds without website releases. Websites cannot override to skip bio.",
   16, 24, 40, "P1")

T1(S1, "[P1] Atomic redemption lock (only one legit fulfill)",
   "When bio is approved, lock reward_id as claimed by user_id in a single transaction/idempotent consumer. Prevent double-click, double-webhook, and two devices. If already fulfilled, ignore later approved sessions. This is the core 'only a legit winner gets the prize' control together with identity match.",
   16, 28, 48, "P1")

T1(S1, "[P1] Physical prize hold / release integration",
   "Publish events or API to WMS/ops: HOLD on win/claim start; RELEASE_TO_SHIP only after Approved + address captured + still the winner. Include winner name, shipping address, reward SKU, bio sessionId for audit. Default-deny: warehouse must not ship from 'won' status alone.",
   12, 20, 32, "P1")

T1(S1, "[P1] High-value digital voucher issuance gate",
   "Do not generate, display, or email voucher codes until Approved. Call voucher issuer once, store token against user_id, mark irre-issuable. If bio later reversed/fraud-flagged, void unused codes where issuer supports it.",
   12, 20, 32, "P1")

T1(S1, "[P1] Retry, velocity and abuse limits",
   "Cap enrollment/auth attempts per user/reward/device/IP. Cooldown after decline 120/105/106. Block concurrent sessions. Alert fraud ops on burst patterns. Complements Veriff liveness — does not replace it.",
   8, 16, 28, "P1")

T1(S1, "[P1] Brand-to-service authentication & authorization",
   "m2m auth (JWT/mTLS/API keys) from each brand website/BFF. Authorize brandId can only act on its users/rewards (prevent IDOR across brands). Rate-limit per brand.",
   8, 16, 24, "P1")

T1(S1, "[P1] Status query API for pending UX",
   "Websites poll or subscribe: current state, whether to show Veriff, retry allowed, generic user message. Must not leak fraud reason codes or other users' data. Frontend callback from Veriff is only a hint to refresh status.",
   8, 12, 20, "P1")

T1(S1, "[P1] Consent ledger",
   "Store consent timestamp, policy version, brand, locale, user_id, IP before opening camera/Veriff. Without consent, APIs refuse session creation. Support withdrawal (stops future auth; does not auto-unfulfill already shipped goods).",
   8, 12, 20, "P1")

T1(S1, "[P1] Immutable audit trail",
   "Who claimed, which prize, enrollment session, auth session, matchedSessionEndUserId, decision, actor for any override, fulfillment event. Retention aligned with legal. Required for disputes ('my cousin claimed my prize').",
   12, 20, 32, "P1")

T1(S1, "[P1] Observability, cost and funnel metrics",
   "Metrics: sessions created, webhook lag, pass rate, decline reasons, drop-off pending→complete, cost/session (Veriff), time-to-fulfill, exception volume. Dashboards + alerts (webhook failures, spike in 120 mismatch, Veriff 5xx).",
   12, 20, 32, "P1")

T1(S1, "[P1] Veriff outage / degraded-mode policy in code",
   "If Veriff is down: do not silently skip bio for gated prizes. Queue pending, extend claim SLA, show user message, page on-call. Optional client-approved emergency override with dual control. Document RTO.",
   8, 16, 28, "P1")

T1(S1, "[P1] Sandbox harness, webhook simulator, contract tests",
   "Automated tests for all decision codes, idempotent webhooks, lock, policy engine, unenrolled branch. Simulator for local/dev without hitting Veriff. Contract tests against Veriff sandbox.",
   12, 20, 32, "P1")

T1(S1, "[P1] Application security review",
   "Threat review of webhook auth, session IDOR, open redirects on callback, PII in logs, secret scanning, prize-status TOCTOU. Fix findings before production.",
   8, 16, 24, "P1")

T1(S1, "[P1] Service runbooks & API documentation for brand teams",
   "OpenAPI, sequence diagrams, error catalog, how to integrate, how to read states. On-call runbook: missed webhook, stuck pending, Veriff portal lookup.",
   8, 16, 24, "P1")

# === P1 ops ===
S1O = "P1 MVP — Business operations, CS, fraud, warehouse (non-dev)"
T1(S1O, "[P1] Ops / admin console (search, hold, dual-control override)",
   "Internal UI: find winner/reward/session, see Veriff decision, freeze prize, release after bio, and dual-control manual exception (two approvers). Every override is audited with reason. Without this, CS will ask engineering to SQL-update production.",
   24, 40, 64, "OPS")

T1(S1O, "[P1] CS playbooks & macros (unenrolled, decline, camera, expiry)",
   "Scripts for: 'I don't want to use my face', 'camera not working', 'it said it failed', 'I am the winner but my partner tried', 'I never received the email'. Clear rule: CS cannot read out voucher codes or force-ship physical goods. Escalation to fraud vs engineering.",
   8, 16, 24, "OPS")

T1(S1O, "[P1] Fraud operations queue & investigation SOP",
   "Case types: face mismatch after win, known-fraud decline, multiple accounts, office/farm velocity, warehouse address farms. Evidence pack: account history, session media reference via Veriff portal, device, shipping data. Outcome: fulfill / forfeit / law-enforcement referral. Staffing model for peak campaigns.",
   8, 16, 28, "OPS")

T1(S1O, "[P1] Exception path SOP (no camera, disability, coverage gap)",
   "Reasonable-adjustment process: alternate IDV (manual document upload to ops), prize substitution to a non-gated digital voucher, or forfeit with explanation. Never a silent skip of bio for physical/high-value without recorded approval. SLA for exception handling inside the claim window.",
   8, 16, 24, "OPS")

T1(S1O, "[P1] Warehouse / 3PL SOP: default hold until bio_approved",
   "Process mapping with fulfillment: pick-status blocked; packing slip includes verification ID; no 'VIP skip'. Training for 3PL. Test a dry-run hold→release→ship and a hold→forfeit→return-to-stock.",
   8, 12, 20, "OPS")

T1(S1O, "[P1] Winner communications (email/SMS/app): bio required",
   "Templates: you won; physical/high-value requires a one-time face check; privacy summary; claim deadline; what if you are not enrolled; do not share the claim link. Legal review of copy. Avoid sending the voucher in the win email for gated prizes.",
   8, 16, 24, "OPS")

T1(S1O, "[P1] Training: CS, fraud, warehouse, brand ops",
   "Live walkthrough of happy path, decline, unenrolled, and override. Quick-reference one-pagers. Recorded session for campaign surge staff.",
   8, 12, 20, "OPS")

T1(S1O, "[P1] UAT with Veriff sandbox + business UAT",
   "Scripted UAT: enrolled auth success, first-time enrollment, small voucher skip, physical hold, decline mismatch, expired session retry, webhook delay, double submit, wrong account. Include warehouse and CS observers. Exit criteria signed by brand owner.",
   16, 24, 40, "OPS")

T1(S1O, "[P1] Production cutover, campaign hypercare",
   "Feature-flag on for physical + high-value only. On-call, Veriff portal watch, daily pass-rate standup for first campaigns. Rollback = keep holds, do not auto-fulfill.",
   8, 16, 24, "OPS")

T1(S1O, "[P1] Veriff usage, billing reconciliation & threshold tuning review",
   "Finance process: per-session cost vs prize value; kill-switch if small-voucher bio is ever enabled accidentally; monthly review of threshold with fraud loss data.",
   4, 8, 16, "OPS")

# === P2 ===
S2 = "P2 — Expand to medium-value digital, pre-enroll, more brands, fraud tooling"
T1(S2, "[P2] Pre-enrollment at win-notification (before redeem)",
   "Start enrollment when the user is notified they won, not only at the last redeem click. Improves conversion for physical/high-value and reduces claim-SLA crunches. New APIs/events from campaign engine → microservice. Still no fulfill until policy gate is satisfied.",
   16, 24, 40, "P2")

T1(S2, "[P2] Medium-value digital voucher threshold rollout",
   "Config + monitoring to extend bio to the medium band (e.g. between Low and High). Includes fraud vs conversion A/B readout and CS volume forecast. Engineering is small if policy engine was built in P1; ops + comms dominate.",
   8, 16, 24, "P2")

T1(S2, "[P2] Funnel analytics dashboard for brands",
   "Won → notified → consented → enrolled → authenticated → fulfilled, with drop-off and decline mix. Per brand/campaign. Used to decide whether to push bio into small vouchers later.",
   16, 24, 40, "P2")

T1(S2, "[P2] FaceBlock / repeat-fraud feedback loop",
   "On confirmed fraud, add faces to Veriff FaceBlock (if contracted) and internal blocklist of user/device/address. Process + API + ops training. Reduces repeat physical-prize theft.",
   8, 16, 28, "P2")

T1(S2, "[P2] Cross-brand / multi-account duplicate person process",
   "Investigate same biometric or IDV identity winning across brands/accounts (promo abuse). Policy: allow (same person, multiple brands) vs forfeit. May require Veriff SE and extra product rules. Higher estimate if legal identity from IDV is available.",
   16, 28, 48, "P2")

T1(S2, "[P2] Localisation of reason mapping and ops templates",
   "Additional markets/languages for user messages, CS macros, consent, and claim emails. Confirm Veriff flow language parameter per brand.",
   8, 12, 20, "P2")

T1(S2, "[P2] Load, webhook backpressure and peak-campaign testing",
   "Load test session creation + webhook ingest for campaign spikes (many high-value winners at once). Queue depth, duplicate webhooks, Veriff rate limits.",
   8, 16, 24, "P2")

# === P3 ===
S3 = "P3 — Hardening: IDV for top tier, name/address match, privacy at scale"
T1(S3, "[P3] Document + selfie IDV as enrollment for physical / top-tier",
   "Strongest 'legit winner' control: government ID + live selfie at enrollment, then biometric auth on later redemptions. Extra Veriff product, UX, webhook fields (name, DOB), and data-retention design. Recommended for physical prizes if selfie-only MVP shows ATO/third-party claims.",
   24, 40, 64, "P3")

T1(S3, "[P3] Shipping-name vs IDV-name match for physical",
   "Compare Veriff extracted name to shipping recipient; fuzzy match; ops queue on mismatch. Reduces 'winner authenticates, ships to mule'. Includes false-positive handling (married names, character sets).",
   12, 20, 32, "P3")

T1(S3, "[P3] Optional bio on small digital vouchers (fraud-triggered)",
   "Only enable if Phase 2 data shows material loss. Includes conversion-impact review, cost model (Veriff fee vs voucher value — often negative ROI), and feature flag. Default recommendation: keep small vouchers out of bio.",
   8, 16, 24, "P3")

T1(S3, "[P3] Data-subject request handling (access / erasure vs fraud hold)",
   "Process + tooling to export or delete biometric-linked records while retaining fraud evidence where legally required. Coordinate Veriff retention.",
   8, 16, 24, "P3")

T1(S3, "[P3] DR / multi-region webhook availability",
   "Ensure webhook endpoint is highly available during campaigns; replay procedure; RPO/RTO documented. Optional second region or queue buffer.",
   8, 16, 32, "P3")


# ---------------------------------------------------------------------------
# TAB 2 Brand websites
# ---------------------------------------------------------------------------
tab2 = []

def T2(section, item, desc, mn, mid, mx, fill_key="P1"):
    tab2.append({
        "section": section,
        "item": item,
        "desc": desc,
        "min": mn,
        "mid": mid,
        "max": mx,
        "fill": fill_key,
    })

B0 = "P0 — Brand-side discovery"
T2(B0, "[P0] Per-brand journey mapping (win → claim → fulfill)",
   "Inventory each website's current redemption: login, magic links, guest winners, voucher reveal, address form, CMS, feature flags. Identify where a fraudster can claim without being the winner. Output: integration points and pages in scope for MVP (physical + high-value digital only).",
   8, 16, 24, "P0")

T2(B0, "[P0] Content, legal pages and contest-rules updates (brand marketing/legal)",
   "Update T&Cs, prize pages, FAQs: biometric may be required; who is eligible; claim deadline; camera requirement; privacy notice link. Not engineering-heavy but blocks go-live.",
   8, 12, 20, "LEGAL")

B1 = "P1 MVP — Shared integration kit (build once, used by all brand sites)"
T2(B1, "[P1] Shared brand integration kit / BFF client",
   "Reusable client (JS SDK or BFF module) that talks only to the Bio-Auth microservice: getPolicy, getStatus, startEnrollment, startAuth, pollStatus. No Veriff API keys in the browser. Handles correlation IDs, retries, and pending UX. First brand consumes it; others import it.",
   24, 40, 64, "P1")

T2(B1, "[P1] Feature-flagged policy consumption (skip bio for small vouchers)",
   "Redeem button path calls policy engine. If BioRequired=false, keep today's flow. If true, enter bio wizard. Prevents wrapping every small voucher in Veriff (cost and conversion).",
   8, 16, 24, "P1")

T2(B1, "[P1] Winner session binding (stop link forwarding / ATO-lite)",
   "Claim pages require the winning account to be logged in. Bind claim token to user_id server-side. Reject 'I forwarded the email to a friend'. Short-lived links. Re-auth (password/OTP) before bio for high-value. This is as important as the selfie.",
   12, 20, 32, "P1")

T2(B1, "[P1] Enrollment UX (consent → camera → Veriff InContext/redirect → pending)",
   "Unenrolled users on a gated prize: explain why, collect consent, launch Veriff (InContext iframe preferred on web; redirect fallback). Never tell the user they are 'verified' on Veriff's return URL — show Pending until microservice status is Approved. Mobile-web camera guidance (Safari/Chrome).",
   16, 28, 48, "P1")

T2(B1, "[P1] Returning-user biometric authentication UX",
   "Enrolled users: shorter selfie-only Veriff flow. Copy differs from first-time enrollment. Same pending/status rules. Do not start auth if pre-check says unenrolled.",
   12, 20, 32, "P1")

T2(B1, "[P1] Unenrolled-user decision screens",
   "Dedicated screens: You must enroll to claim this physical/high-value prize; small voucher alternative if policy allows substitution; decline-consent consequences (prize held until SLA then forfeited). Accessibility of this branch is a common drop-off.",
   8, 16, 24, "P1")

T2(B1, "[P1] Pending / webhook-wait UX (do not trust Veriff callback)",
   "After camera, poll status with timeout and 'we're confirming it's you' state. Success only when microservice says Approved. Handle refresh, back button, and second tab.",
   8, 16, 24, "P1")

T2(B1, "[P1] Decline, retry, expired, abandoned UX",
   "Friendly, non-accusatory copy for quality failures (retry). Hard stop + CS path for mismatch/fraud without revealing fraud typology. Expired session → start new session, not reuse URL. Attempt counter visible where allowed.",
   8, 16, 28, "P1")

T2(B1, "[P1] Device / camera capability detection & fallback CTA",
   "Detect missing camera, insecure context (non-HTTPS), in-app browsers that break getUserMedia. Prompt switch to supported browser or use exception/CS path. Critical on mobile web.",
   8, 12, 20, "P1")

T2(B1, "[P1] Physical prize claim UI (address AFTER successful bio)",
   "Do not collect shipping address before identity gate if that leaks to fraud ops load. After Approved: address form, restricted-country list, confirmation that recipient is the winner. Show 'we will ship only after verification' earlier in the journey.",
   12, 20, 32, "P1")

T2(B1, "[P1] High-value digital voucher reveal (code only after Approved)",
   "Remove voucher code from win page/email for gated amounts. Reveal once in authenticated page + optional email after Approved. Clipboard + screenshot warning. Prevent HTML cache of the code.",
   8, 16, 24, "P1")

T2(B1, "[P1] Consent, privacy notice and T&Cs capture UI",
   "Versioned checkboxes, links to privacy notice, withdraw-consent link. Block Veriff launch without consent. Store on microservice via API.",
   8, 12, 20, "P1")

T2(B1, "[P1] Accessibility and inclusive-design pass",
   "WCAG for the new wizard, keyboard, screen reader, motion/camera alternatives pointing at exception path. Avoid trapping users in a camera modal.",
   8, 16, 24, "P1")

T2(B1, "[P1] Analytics events (site funnel)",
   "Track policy_required, consent_accepted/denied, veriff_launched, pending, approved, declined, abandoned, fulfill_shown — without sending biometrics. Align names with microservice metrics.",
   4, 8, 16, "P1")

T2(B1, "[P1] UX copy deck + localisation for bio screens",
   "Source copy in brand voice; translate for each live locale. Include camera tips, privacy one-liners, and failure messages. Review with legal.",
   8, 16, 24, "P1")

T2(B1, "[P1] QA on reference brand (desktop + mobile web cameras)",
   "Real devices: iOS Safari, Android Chrome, in-app Instagram/Facebook browsers, desktop webcam absent. All decision paths. Regression on small-voucher redeem (must still work without bio).",
   16, 28, 40, "P1")

B1B = "P1 MVP — Per additional brand website (unit cost after the first brand)"
T2(B1B, "[P1] Per additional brand site: wire-up, callbacks, config, UAT",
   "UNIT COST — multiply by (N brands − 1). Import shared kit, brandId/keys, allowlisted callback domains, feature flags, CMS content, locale, prize-type mapping, staging UAT, production hypercare. Assumes similar stack; max hours if the brand is a legacy/custom platform.",
   16, 32, 48, "P1")

T2(B1B, "[P1] Per additional brand: CS FAQ, macros, campaign ops checklist",
   "UNIT COST — multiply by (N brands − 1). Brand-specific prize names, CS hours, and warehouse SKUs in the SOP. Not code, but required for a safe launch.",
   4, 8, 16, "OPS")

B2 = "P2 — Brand expansion"
T2(B2, "[P2] Pre-enroll CTA on win notification surfaces (web + email)",
   "Win page and email: 'Verify now so you can claim faster' for gated prizes. Deep link into enrollment kit while prize remains on hold.",
   8, 16, 24, "P2")

T2(B2, "[P2] Medium-threshold digital voucher UI rollout",
   "Extend gated reveal to medium band; watch conversion. A/B or staged brand rollout.",
   8, 12, 20, "P2")

T2(B2, "[P2] Brand theming / Veriff flow language QA",
   "Pass lang parameter, brand colors if contracted, and visual QA per brand guidelines.",
   8, 16, 28, "P2")

B3 = "P3 — Brand hardening"
T2(B3, "[P3] IDV document-capture UX for physical / top-tier",
   "Longer flow (ID + selfie). Country document hints, lighting tips, resume-later inside claim SLA. Higher drop-off — needs progress saver.",
   16, 28, 48, "P3")

T2(B3, "[P3] Step-up login (OTP / password) immediately before bio",
   "Reduce account-takeover that then passes face match because the thief enrolled after stealing the session. Especially if enrollment can happen post-win.",
   8, 16, 24, "P3")


# ---------------------------------------------------------------------------
# TAB 3 requirements
# ---------------------------------------------------------------------------
reqs = [
    ("P0", "Y", "Signed Veriff commercial agreement covering Biometric Authentication + Enrollment (and optional IDV), sandbox and production, estimated session volumes, FaceBlock, and data processing terms.",
     "Without a live Veriff account and the two linked integrations (enrollment + authentication), the microservice cannot create sessions or receive decisions. Volume estimates also drive the cost vs small-voucher policy."),
    ("P0", "Y", "Veriff Solutions Engineer to provision and connect a Selfie Enrollment integration and a Biometric Authentication integration (separate API keys and baseURLs).",
     "Veriff requires an approved enrollment session before authentication. Connecting the two integrations is a vendor configuration step, not something engineering can do from the public API alone."),
    ("P0", "Y", "Veriff Customer Portal users (engineering, fraud ops, maybe CS lead) for sandbox and production, with webhook URL fields writable.",
     "Webhooks, test-session forcing, and investigation of declined redemptions are done in the portal. Missing access blocks UAT and freeze/fraud handling."),
    ("P0", "Y", "Sandbox and production API keys, baseURLs, and webhook signature secrets delivered via a secret manager (not email/Slack).",
     "Needed to call POST /sessions and to verify that decision webhooks are actually from Veriff. Keys must never be placed on brand websites."),
    ("P0", "Y", "List of brand websites in scope: legal entity, production/staging URLs, owners, tech stack, locales, and whether web-only or also native apps.",
     "Microservice is multi-tenant. Callback allowlists, CORS, lang parameters, and the per-brand integration estimate depend on this list."),
    ("P0", "Y", "Stable user identity scheme per brand: is user_id UUID v4? If not, the immutable identifier to send as vendorData. Guest vs registered winner policy.",
     "Veriff endUserId must be UUID v4; otherwise vendorData (immutable, non-semantic). If guest winners exist, selfie-only bio cannot prove they are the intended recipient — we need a registered account policy for gated prizes."),
    ("P0", "Y", "Existing winner / campaign / redemption APIs and data model (how a win is stored, claim tokens, voucher issuer, fulfillment statuses).",
     "The bio service must lock and release the real prize record. We cannot build a side database that the warehouse and voucher issuer ignore."),
    ("P0", "Y", "Prize catalog: physical vs digital voucher, face values, currencies, SKUs, average monthly volumes, and historical fraud/loss on each band.",
     "Drives MVP thresholds (physical + high-value only) and Veriff cost modelling. Small-voucher volume is usually too high to put behind bio in MVP."),
    ("P0", "Y", "Client-approved threshold table: High (MVP in), Medium (Phase 2), Low (bio not required). Per brand/market if needed.",
     "Policy engine is configuration. Without signed numbers we cannot estimate conversion impact or finish UAT cases."),
    ("P0", "Y", "Claim window / SLA (e.g. 7 or 14 days) and what happens on expiry (forfeit, redraw, CS exception).",
     "Must match Veriff session TTL (unstarted sessions expire in 7 days) and warehouse reservation time."),
    ("P0", "Y", "Fulfillment / 3PL process for physical prizes today: statuses, who can release a shipment, address validation, restricted countries.",
     "MVP success depends on default-hold. If 3PL ships on 'won', bio will not protect the prize regardless of software quality."),
    ("P0", "Y", "Digital voucher issuer details: API, whether codes can be generated lazily, voided, and whether win emails already contain the code.",
     "Gated digital prizes must not reveal the code before Approved. If emails already contain codes, marketing/ops must change templates before go-live."),
    ("P0", "Y", "Current authentication model on each site (session, SSO, magic link, OTP) and account-recovery flow.",
     "Bio proves the person matches the enrolled face, not that the account wasn't stolen before enrollment. We need to know how weak login/recovery is to design step-up."),
    ("P0", "Y", "Whether any prior KYC/IDV/selfie exists (including other vendors) and if Veriff can migrate reference images.",
     "Migration can skip mass re-enrollment. If none exist, every gated winner is an unenrolled user and must go through enrollment at claim (or pre-enroll)."),
    ("P0", "Y", "Markets / countries in go-live and any biometric or promotions-law constraints (GDPR, consent age, photography bans, prize regulations).",
     "DPIA, consent UX, Veriff coverage, and whether IDV documents are acceptable in those markets. Affects legal hours and exception SOP."),
    ("P0", "Y", "Client privacy counsel contact and whether a DPIA / ROPA update is already in flight. Retention policy for biometric-linked records.",
     "We cannot set retention or consent copy without the controller's policy. Blocks P0 legal workstream."),
    ("P0", "Y", "Brand consent / privacy notice copy owner and required legal review SLA.",
     "Camera cannot open without consented purpose text. Late legal copy is a common go-live slip."),
    ("P0", "Y", "CS, fraud, and warehouse owners + surge staffing model for campaign peaks.",
     "Exception and mismatch queues are business operations, not a developer backlog. Need named owners to write SOPs and staff UAT."),
    ("P0", "N", "Historical fraud cases on rewards (ATO, farmed accounts, physical not-received, voucher resale).",
     "Tunes thresholds, velocity limits, and whether selfie-only enrollment is acceptable or IDV should be pulled into MVP for physical."),
    ("P1", "Y", "Staging environments of each in-scope brand site that can reach the bio microservice and Veriff sandbox (HTTPS, cameras allowed).",
     "UAT of InContext/redirect and webhooks requires real browser camera flows on staging, not mocks only."),
    ("P1", "Y", "Allowlisted callback / return URLs per brand (prod + staging) and who can request DNS/WAF changes.",
     "Veriff redirects users back to the site. Wrong or blocked callbacks look like 'bio is broken'. WAF must allow Veriff webhook IPs/headers to the microservice."),
    ("P1", "Y", "Decision on Veriff UX: InContext SDK vs full redirect, and default language per brand.",
     "Frontend kit implementation differs. Redirect is simpler but higher drop-off; InContext is better UX but more QA on iframe/camera."),
    ("P1", "Y", "Manual-override policy: who may release a prize without approved bio, dual-control, and which prize types are never overridable.",
     "Admin console is dangerous without a written policy. Physical prizes should rarely (if ever) skip bio."),
    ("P1", "Y", "Exception policy for users without a camera, with disabilities, or in unsupported browsers/devices.",
     "Unenrolled / unable-to-enroll is a guaranteed real-world path. Need an approved alternative (manual ID check, substitute prize, or forfeit) before launch."),
    ("P1", "Y", "Retry policy: max attempts after quality decline vs immediate freeze after mismatch/fraud codes.",
     "Coded into the state machine. Too generous → fraud; too strict → CS volume and prize complaints."),
    ("P1", "Y", "Sample test winners per brand (physical, high-value digital, small digital) on staging, including users with and without webcams.",
     "UAT scripts cannot run on production campaigns. Need seed data and permission to complete Veriff sandbox sessions."),
    ("P1", "N", "Brand design tokens / Veriff flow customization options (logo, color) if purchased.",
     "Affects frontend theming hours and vendor config. Can launch with default Veriff UI in MVP."),
    ("P1", "Y", "Production feature-flag owner and campaign calendar (first physical / high-value draws after go-live).",
     "Cutover should land before a high-value drop, with hypercare. Do not enable on a mega small-voucher drop first."),
    ("P1", "N", "Existing customer-comms platform templates access (Braze/SFMC/etc.) for win and reminder emails.",
     "Gated prizes must not email voucher codes early. Physical win emails must mention the face check and deadline."),
    ("P2", "N", "Whether the same natural person may win across multiple brands, and the source-of-truth person ID if yes.",
     "Needed for duplicate-face / cross-brand abuse rules in Phase 2. If unknown, we treat brands as isolated identities in MVP."),
    ("P2", "N", "Appetite to use Veriff FaceBlock and who is authorised to add faces.",
     "Repeat physical-prize fraudsters. Requires contract add-on and an ops process."),
    ("P3", "N", "Whether Document+Selfie IDV is approved for physical/top-tier (extra vendor cost, higher drop-off, stronger legal-identity bind).",
     "Selfie-only enrollment cannot prove the enrollee is the intended winner if the account was already taken over. IDV is the recommended end-state for physical prizes."),
    ("P3", "N", "Name-match and address-match tolerance rules for shipping vs IDV payload.",
     "Only relevant if IDV is in scope. Needs ops false-positive policy (hyphenated names, character sets)."),
    ("P0", "Y", "Budget owner for Veriff per-session fees vs prize COGS, and a kill-switch authority if cost/conversion is worse than fraud savings.",
     "Putting bio on small vouchers can cost more than the voucher. Finance must own the threshold decision, not engineering."),
    ("P1", "Y", "WAF / public endpoint ownership for the microservice webhook URL (TLS cert, Veriff webhook delivery requirements).",
     "Veriff must reach our webhook; we must verify signatures. Network/security team usually owns this."),
    ("P0", "N", "Native iOS/Android apps in scope, or web-only for MVP.",
     "Native needs Veriff native SDKs and extra hours not included in the website kit. Default estimate assumes web (including mobile web)."),
    ("P1", "Y", "On-call / incident contacts on client side (site down, campaign live, warehouse blocked).",
     "Hypercare is joint. A frozen physical-prize queue at 3PL is a business incident, not only an HTTP 500."),
]


# ---------------------------------------------------------------------------
# Process / decision flow rows
# ---------------------------------------------------------------------------
flow = [
    ("0", "All", "User, Website", "User opens claim for a reward", "Continue to account check", "—", "P1"),
    ("1", "Control", "Website + existing auth", "Is the user logged in as the winning account (claim token bound to user_id)?", "Continue. Optionally step-up OTP for high-value/physical.", "Force login / recovery. Do NOT start Veriff. Do NOT show voucher or address form. A forwarded email must fail.", "P1"),
    ("2", "Policy", "Microservice policy engine", "Load prize type, face value, currency, brand. Is biometric step-up required?", "Physical: YES. Digital >= High threshold: YES (MVP). Digital medium: YES in P2. Digital small: NO.", "If NO: fulfill with existing controls (account session, rate limits). Do not call Veriff (cost/conversion).", "P1"),
    ("3", "Consent", "Website + Microservice", "Has the user accepted current biometric consent / privacy version for this purpose?", "Record consent ledger (timestamp, version, brand, locale).", "Show consent. If refused: hold prize until claim SLA then forfeit (or offer ungated substitute if policy allows). No camera.", "P1"),
    ("4", "Enrollment status", "Microservice", "Does this endUserId have an approved Veriff enrollment (reference template exists)?", "Go to AUTHENTICATION path (step 8).", "Go to ENROLLMENT path (step 5). Do not create an auth session (avoids decline 122 'no reference').", "P1"),
    ("5", "Enrollment readiness", "Website", "Can this device complete a camera flow (HTTPS, camera, supported browser)?", "Create ENROLLMENT session via microservice (Veriff enrollment integration).", "Offer: switch device/browser, or Exception queue (disability / no camera). No silent skip for physical/high-value.", "P1"),
    ("6", "Enrollment capture", "User + Veriff", "User completes enrollment: MVP selfie-only (minimum) OR Document+Selfie IDV (recommended for physical / top-tier, P3 or pull-forward).", "Veriff returns decision via WEBHOOK (source of truth). Frontend callback only refreshes status.", "User abandons: session → abandoned/expired; allow new session until claim SLA.", "P1"),
    ("7", "Enrollment decision", "Microservice webhook", "Enrollment verification.status?", "approved: mark Enrolled; for THIS sitting, identity captured → FULFILL GATE (step 11). Do not charge a second auth session immediately. declined quality: retry. declined 105/106/fraud: FREEZE + fraud ops. review/pending: wait.", "Never fulfill on 'session created' or on redirect back to the site.", "P1"),
    ("8", "Authentication start", "Microservice", "Create AUTH session on the authentication integration with the SAME endUserId/vendorData and bind to this reward_id / attempt_id.", "Return session URL to website. One active session per reward.", "If vendorData/endUserId would not match enrollment, stop (prevent 121/128). Repair identity mapping.", "P1"),
    ("9", "Authentication capture", "User + Veriff", "User takes live selfie (liveness + anti-spoof). Session auto-submits via SDK.", "Wait for webhook / poll.", "Abandoned/expired: new session, retry limits apply.", "P1"),
    ("10", "Authentication decision", "Microservice webhook", "Auth status + biometricAuthentication.matchedSessionEndUserId == winner endUserId?", "approved + match: FULFILL GATE. declined 120 mismatch: FREEZE (third party / ATO). 105/106: FREEZE fraud. 122: treat as unenrolled. 121/128: incident, do not fulfill. quality: retry.", "No fulfill. CS sees generic failure; ops sees codes.", "P1"),
    ("11", "Fulfill gate", "Microservice lock", "Is reward still unfulfilled and still assigned to this user (atomic lock)?", "Yes: branch by prize type (12a/12b).", "Already fulfilled or reassigned: ignore webhook, alert. This prevents double win.", "P1"),
    ("12a", "Physical fulfill", "Microservice + WMS + User", "Capture/confirm shipping address AFTER approval. Optional P3: IDV name vs recipient name. Emit RELEASE_TO_SHIP with bio sessionId on the packing record.", "Warehouse ships. Notify winner. Audit.", "Name/address mismatch or fraud freeze: do not ship; ops case. Default-deny at 3PL if bio flag missing.", "P1"),
    ("12b", "Digital high-value fulfill", "Microservice + voucher issuer", "Generate/reveal voucher once to this account. Email only after Approved. Mark irre-issuable.", "User sees code in authenticated page.", "Issuer failure: keep Approved, retry issue, do not require another selfie. Do not email codes from the original 'you won' template.", "P1"),
    ("13", "Fraud freeze", "Fraud ops", "Mismatch / known fraud / velocity hit.", "Prize remains held. Investigate. Outcomes: forfeit, fulfill after dual-control (rare), FaceBlock (P2).", "CS cannot override a fraud freeze alone.", "P1"),
    ("14", "Exception queue", "CS + authorised approver", "No camera, disability, Veriff outage beyond SLA, coverage gap.", "Documented alternative (manual ID check, substitute ungated prize, extend SLA). Dual-control if releasing physical/high-value without Veriff approved.", "No undocumented skip. Record in admin console.", "P1"),
    ("15", "Claim SLA expiry", "Microservice + ops", "Still not enrolled/approved when claim window ends?", "Forfeit (or redraw per policy). Release warehouse stock. Notify user.", "Do not silently fulfill. Do not keep holds forever.", "P1"),
    ("16", "Veriff outage", "On-call", "Cannot create sessions or receive decisions.", "Keep gated prizes on HOLD. User messaging. Optional dual-control emergency policy. Never auto-skip bio.", "Small ungated vouchers continue.", "P1"),
]


assumptions = [
    ("Scope", "MVP gates physical prizes and high-value digital vouchers only. Small-value digital vouchers stay on existing account controls because volume is high and Veriff per-session cost plus conversion drop usually exceeds fraud savings."),
    ("Architecture", "One Bio-Auth microservice owns Veriff keys, webhooks, policy, enrollment registry, and fulfillment unlock. Multiple brand websites integrate as clients. Websites never call Veriff directly."),
    ("Identity", "Gated prizes require a registered, authenticated winning account. Guest-link winners are out of MVP unless converted to an account before bio."),
    ("Enrollment", "Veriff biometric authentication cannot run without a prior approved enrollment. Unenrolled users are a first-class path (enroll-at-claim), not an error."),
    ("Strength", "Selfie-only enrollment proves returning-person match, not legal identity. If the account is already taken over, the thief enrolls their own face. IDV (document+selfie) is recommended for physical prizes and can be Phase 3 or pulled into MVP if ATO risk is high."),
    ("Source of truth", "Fulfillment is allowed only after a verified Veriff decision webhook (or poll of the same decision) AND an atomic prize lock. The Veriff frontend callback is not authorization."),
    ("Same-sitting", "A successful enrollment during claim may release that prize without an immediate second authentication session (saves vendor cost/UX). Subsequent redemptions require authentication."),
    ("Hours basis", "Hours are blended delivery effort (engineering, QA, BA/architecture, CS/fraud/warehouse ops design, privacy liaison). They exclude Veriff vendor fees, 3PL rate cards, and client legal counsel of record."),
    ("Brands", "Tab 2 includes a full first-brand build. 'Per additional brand' rows are unit costs. Summary sheet lets you set N brands."),
    ("Stack", "Assumes web + mobile web. Native iOS/Android Veriff SDKs are not in the base hours; declare apps in client requirements to add a delta."),
    ("Team", "Mid estimate assumes a squad that has built API/webhook integrations before, timely Veriff SE responses, and staging camera access. Max includes vendor delay, legal delay, and a legacy brand platform."),
    ("3-point", "Min = best case, clear inputs, 1 modern brand. Mid = expected. Max = delayed vendor/legal, 3+ messy brands, or IDV pulled into MVP. PERT expected = (Min + 4×Mid + Max) / 6."),
    ("Non-dev", "CS playbooks, warehouse hold SOP, communications, training, and dual-control exceptions are in the estimate because they are required to actually stop a non-winner receiving a prize."),
    ("Out of scope MVP", "Replacing existing login; full KYC for all users at registration; putting bio on every small voucher; building a new warehouse WMS; native apps unless declared."),
]

risks = [
    ("Critical", "3PL ships on 'won' instead of 'bio_approved'", "Physical prizes leave the warehouse before the selfie. Bio becomes theatre.", "Contractual default-hold, packing-slip flag, dry-run UAT with 3PL."),
    ("Critical", "Voucher codes already in 'you won' emails", "High-value digital is claimed by whoever has the inbox.", "Change templates before flag-on; lazy-generate codes."),
    ("Critical", "Guest winners + selfie-only enrollment", "Fraudster opens the claim link, enrolls their face, and is now the 'legit' biometric user.", "Require registered account; consider IDV for physical; bind claim token to user_id."),
    ("High", "Veriff SE / dual integration not ready", "Entire P1 blocked; hours burn on mocks.", "P0 vendor checkpoint before sprinting website UX."),
    ("High", "Legal/DPIA delay on biometrics", "UX built, cannot go live in EU/other markets.", "Start DPIA in P0; selfie consent copy in parallel."),
    ("High", "Trusting frontend callback", "Attacker fakes 'success' query param and redeems.", "Server-side webhook/poll only; status API."),
    ("High", "Unenrolled users treated as hard fail", "Mass CS contacts and prize complaints; Veriff 122 declines.", "Pre-check enrollment; enroll-at-claim UX; exception SOP."),
    ("High", "Account takeover then enroll", "Thief passes bio because they are the enrolled face.", "Step-up login; optional IDV; speed-to-enroll after win; anomaly on new device."),
    ("Medium", "Camera / in-app browser failures", "Mobile-web drop-off, especially from social apps.", "Detect unsupported browsers; 'open in Safari/Chrome'; exception path."),
    ("Medium", "Bio on small vouchers", "Negative ROI; conversion crash.", "Keep Phase 1 thresholds high; finance kill-switch."),
    ("Medium", "Manual override without dual control", "CS becomes the bypass for fraud.", "Admin console permissions + audit."),
    ("Medium", "Webhook missed at peak", "Users stuck pending, CS volume.", "Poll fallback, alerts, Veriff retry, load test."),
    ("Low", "Name mismatch false positives (P3)", "Legit winners blocked on shipping.", "Fuzzy match + ops queue."),
]


def write_cover(wb):
    ws = wb.active
    ws.title = "0-How to use"
    apply_sheet_defaults(ws, freeze="A1")
    ws.freeze_panes = "A1"
    banner(ws, "Veriff Biometric Authentication — Rewards Redemption",
           "Solution architecture estimation pack  |  Microservice used by multiple brand websites  |  Physical prizes + digital vouchers  |  Chat/estimate only — no deployment",
           cols=6)
    set_col_widths(ws, [28, 28, 22, 22, 22, 55])

    ws.row_dimensions[3].height = 10
    note_row(ws, 4,
             "How to use: Tabs 1–2 are 3-point hour estimates (Min / Mid / Max) plus PERT expected. Tab 3 is client inputs that unblock the estimate. Tab 4 rolls up MVP vs later phases and lets you set brand count. Tab 5 is the recommended process/decision flow. Tab 6 lists assumptions and risks. Hours include engineering AND business operations. They exclude Veriff vendor fees.",
             cols=6, fill=fill_gold)

    # Architecture
    r = 6
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "Recommended architecture (logical)").font = font_navy_h
    ws.cell(r, 1).fill = fill_p0
    for c in range(1, 7):
        ws.cell(r, c).fill = fill_p0
    r = 7
    arch = (
        "Brand Website A / B / C  →  (m2m auth)  Bio-Auth Microservice  →  Veriff Enrollment integration  +  Veriff Authentication integration\n"
        "Microservice owns: policy/thresholds, endUserId registry, session state, webhook (source of truth), atomic prize lock, hold/release to WMS & voucher issuer.\n"
        "Brand sites own: login-as-winner, consent UI, Veriff InContext/redirect using session URL only, pending UX, address form / voucher reveal AFTER Approved.\n"
        "Fulfillment systems (3PL / voucher issuer) must NOT release on 'won'. They release only on microservice Fulfilled/Release event."
    )
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    cell = ws.cell(r, 1, arch)
    cell.alignment = wrap
    cell.font = font_normal
    ws.row_dimensions[r].height = 78

    r = 9
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "Recommended MVP and phases").font = font_navy_h
    ws.cell(r, 1).fill = fill_p1
    for c in range(1, 7):
        ws.cell(r, c).fill = fill_p1

    phases = [
        ("Phase", "What goes live", "Why this cut", "Bio required?", "Unenrolled users", "Prize types"),
        ("P0 Discovery", "Architecture, Veriff dual-integration, DPIA/consent, signed policy table, identity model.",
         "Vendor + legal + policy are the critical path. Building UX first wastes rework.",
         "n/a", "Policy defined", "Taxonomy only"),
        ("P1 MVP (recommend)", "Microservice + 1..N brand kit. Gate PHYSICAL prizes and HIGH-VALUE digital vouchers. Small vouchers unchanged. Enroll-at-claim. Warehouse default-hold. CS/fraud SOPs. Admin dual-control.",
         "Physical goods are irreversible and attractive to fraud. High-value digital codes leak via inbox/ATO. Small vouchers are high-volume / low-value — bio would dominate Veriff cost and hurt conversion.",
         "Physical = always. Digital >= High threshold = yes. Digital small = no.",
         "First-class enroll-at-claim (consent → enrollment session). Same-sitting approved enrollment can release THAT prize. Exception SOP for no camera.",
         "Physical + high-value digital"),
        ("P2 Expand", "Medium-value digital; pre-enroll at win notification; more brands; funnel dashboard; FaceBlock; load testing.",
         "Once pass-rate and CS load are known, extend coverage where fraud savings > cost+drop-off.",
         "Add medium digital band via config (no website rewrite if P1 policy engine is done).",
         "Pre-enroll so they are not seeing bio for the first time at the last click.",
         "+ medium digital"),
        ("P3 Harden", "Document+Selfie IDV for physical/top-tier; shipping name match; privacy DSR tooling; optional small-voucher bio only if data supports it (default: still skip).",
         "Selfie-only cannot bind legal identity. IDV is how you prove the enrollee is the intended winner, not an ATO.",
         "Physical/top-tier step-up to IDV. Small vouchers remain ungated unless fraud data says otherwise.",
         "IDV enrollment; name/address mismatch queue.",
         "Stronger physical"),
    ]
    for i, rowv in enumerate(phases):
        rr = 10 + i
        for c, v in enumerate(rowv, 1):
            cell = ws.cell(rr, c, v)
            cell.alignment = wrap
            cell.border = thin
            if i == 0:
                cell.font = font_white_bold
                cell.fill = fill_navy
            elif i == 2:
                cell.font = font_normal
                cell.fill = fill_p1
            else:
                cell.font = font_normal
                cell.fill = fill_white if i % 2 else fill_alt
        ws.row_dimensions[rr].height = 88 if i else 22

    r = 16
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "What 'only a legit winner' actually requires (controls stacked)").font = font_navy_h
    ws.cell(r, 1).fill = fill_ops
    for c in range(1, 7):
        ws.cell(r, c).fill = fill_ops
    controls = [
        "1. Prize assigned to a registered account (not a bearer magic link).",
        "2. Claim-time login (and step-up) as that account — stops casual link forwarding.",
        "3. Enrollment binds a live face (liveness) to that account's endUserId — with IDV if you also need legal name.",
        "4. Authentication on later claims matches the same face + same endUserId (Veriff matchedSessionEndUserId checked server-side).",
        "5. Atomic lock so a prize can be fulfilled once.",
        "6. Warehouse / voucher issuer default-deny until the microservice says release.",
        "7. Ops freeze on mismatch/fraud; no CS-only override for physical/high-value.",
    ]
    for i, t in enumerate(controls):
        rr = 17 + i
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
        cell = ws.cell(rr, 1, t)
        cell.font = font_normal
        cell.alignment = left_c
        ws.row_dimensions[rr].height = 18

    r = 25
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "Unenrolled users (explicit product behaviour, not an afterthought)").font = font_navy_h
    ws.cell(r, 1).fill = fill_p2
    for c in range(1, 7):
        ws.cell(r, c).fill = fill_p2
    unenr = (
        "Most winners will NOT have a Veriff biometric template on day one. The microservice must pre-check enrollment and branch:\n"
        "• Small digital, bio not required → redeem as today.\n"
        "• Gated prize, not enrolled, camera OK, consent given → Enrollment session (not Auth). On approved enrollment, release this prize.\n"
        "• Gated prize, not enrolled, no camera / refused consent → Exception SOP or forfeit at SLA. Never a raw Veriff decline 122 as the UX.\n"
        "• Later prize, already enrolled → Authentication session only.\n"
        "Do not silently skip the gate; do not force Auth on an empty template."
    )
    r = 26
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    cell = ws.cell(r, 1, unenr)
    cell.alignment = wrap
    cell.font = font_normal
    ws.row_dimensions[r].height = 90

    r = 28
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1, "Workbook tabs").font = font_navy_h
    tabs = [
        ("1-Microservice Bio Auth", "Work items for the shared Veriff microservice, policy engine, webhooks, locks, admin, CS/fraud/warehouse ops."),
        ("2-Brand Websites", "Work items for the shared website kit plus unit cost per additional brand. Includes UX for unenrolled users and gated voucher/physical flows."),
        ("3-Client Requirements", "What we need from the client/developer counterparts and why. Blocking vs optional, by phase."),
        ("4-Phased Estimate Summary", "Hour roll-ups, PERT expected, brand-count calculator, MVP vs full."),
        ("5-Process Decision Flow", "End-to-end decision table used for estimation workshops and later user stories."),
        ("6-Assumptions & Risks", "Scope cuts and the failure modes that would still let a non-winner get the prize."),
    ]
    ws.cell(29, 1, "Tab").font = font_white_bold
    ws.cell(29, 1).fill = fill_navy
    ws.merge_cells(start_row=29, start_column=2, end_row=29, end_column=6)
    ws.cell(29, 2, "Contents").font = font_white_bold
    ws.cell(29, 2).fill = fill_navy
    for c in range(2, 7):
        ws.cell(29, c).fill = fill_navy
    for i, (n, d) in enumerate(tabs):
        rr = 30 + i
        ws.cell(rr, 1, n).font = font_bold
        ws.cell(rr, 1).border = thin
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
        ws.cell(rr, 2, d).font = font_normal
        ws.cell(rr, 2).alignment = wrap_m
        for c in range(1, 7):
            ws.cell(rr, c).border = thin
        ws.row_dimensions[rr].height = 28

    ws.row_dimensions[36].height = 36
    ws.merge_cells("A36:F36")
    ws["A36"] = "Estimates are for planning / solutioning. Re-baseline after P0 (vendor access + signed threshold table). Not a commercial statement of work."
    ws["A36"].font = font_small
    ws["A36"].fill = fill_input
    ws["A36"].alignment = wrap_m
    for c in range(1, 7):
        ws.cell(36, c).fill = fill_input


def write_work_sheet(wb, title, full_title, subtitle, items, extra_note):
    ws = wb.create_sheet(title)
    apply_sheet_defaults(ws, freeze="A8")
    banner(ws, full_title, subtitle, cols=6)
    note_row(ws, 3, extra_note, cols=6, fill=fill_gold)
    note_row(ws, 4,
             "Min = best case (clear requirements, Veriff sandbox ready, one modern brand).  Mid = expected.  Max = vendor/legal delay, messy identity, extra brands or IDV pulled into MVP.  "
             "Expected (PERT) = (Min + 4×Mid + Max) / 6.  Hours are blended (dev + QA + BA + ops design).  Colour: grey P0, green P1 MVP, yellow P2, purple P3, blue ops, orange legal/privacy.",
             cols=6, fill=fill_cover)
    ws.row_dimensions[5].height = 8
    headers = ["Work items", "Description", "Min (hrs)", "Mid (hrs)", "Max (hrs)", "Expected PERT (hrs)"]
    header_row(ws, 6, headers)
    ws.auto_filter.ref = "A6:F6"
    ws.auto_filter.add_sort_condition("A6:F6")

    set_col_widths(ws, [52, 88, 14, 14, 14, 20])

    # Group by section preserving order
    sections = []
    for it in items:
        if not sections or sections[-1] != it["section"]:
            sections.append(it["section"])

    row = 7
    data_rows = []
    unit_rows = []
    current_section = None
    for it in items:
        if it["section"] != current_section:
            current_section = it["section"]
            fill = PHASE_FILL.get(it["fill"], fill_p0)
            section_row(ws, row, current_section, 6, fill)
            row += 1
        values = [it["item"], it["desc"], it["min"], it["mid"], it["max"]]
        data_row(ws, row, values, heights=84)
        pert = ws.cell(row, 6)
        pert.value = f'=ROUND((C{row}+4*D{row}+E{row})/6,1)'
        pert.number_format = "0.0"
        pert.font = font_num
        pert.alignment = center
        pert.border = thin
        pf = PHASE_FILL.get(it["fill"], fill_white)
        for c in range(3, 7):
            ws.cell(row, c).fill = pf
        ws.cell(row, 1).font = font_bold
        data_rows.append(row)
        if "Per additional brand" in it["item"]:
            unit_rows.append(row)
        row += 1

    def total_line(row, label, formula_c, fill=fill_total, font=font_total):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        t = ws.cell(row, 1, label)
        t.font = font
        t.fill = fill
        t.alignment = left_c
        ws.cell(row, 2).fill = fill
        ws.cell(row, 2).border = thin
        for col, letter in [(3, "C"), (4, "D"), (5, "E"), (6, "F")]:
            expr = formula_c.replace("{col}", letter)
            c = ws.cell(row, col, expr)
            c.font = font
            c.fill = fill
            c.number_format = "0.0"
            c.alignment = center
            c.border = thin
        ws.row_dimensions[row].height = 24
        return row

    first_data = data_rows[0]
    last_data = data_rows[-1]
    total_label = (
        "TOTAL (all rows on this tab, including one extra-brand unit)"
        if unit_rows else
        "TOTAL (all phases on this tab)"
    )
    total_row = total_line(row, total_label, f"=SUM({{col}}{first_data}:{{col}}{last_data})")
    row += 1
    if unit_rows:
        unit_sum = "+".join(f"{{col}}{r}" for r in unit_rows)
        total_line(row, "Of which: extra-brand UNIT COST (multiply by N−1 on Tab 4)",
                   f"={unit_sum}", fill=fill_p2, font=font_bold)
        row += 1
        total_line(row, "FIRST BRAND ONLY (tab total minus extra-brand unit)",
                   f"=SUM({{col}}{first_data}:{{col}}{last_data})-({unit_sum})",
                   fill=fill_p1, font=font_bold)
        row += 1
    row += 1

    # Phase subtotals using SUMIF on item prefix is hard; compute in Python
    ws.cell(row, 1, "Phase subtotals (from item tags)").font = font_navy_h
    row += 1
    header_row(ws, row, ["Phase", "Work item count", "Min (hrs)", "Mid (hrs)", "Max (hrs)", "Expected PERT (hrs)"])
    row += 1
    phase_order = [("P0", "P0 Discovery"), ("P1", "P1 MVP"), ("P2", "P2 Expand"), ("P3", "P3 Harden")]
    phase_start = row
    for key, label in phase_order:
        subset = [it for it in items if it["item"].startswith(f"[{key}]")]
        ws.cell(row, 1, label).font = font_bold
        ws.cell(row, 1).border = thin
        ws.cell(row, 2, len(subset)).border = thin
        ws.cell(row, 2).alignment = center
        for col, field in [(3, "min"), (4, "mid"), (5, "max")]:
            cell = ws.cell(row, col, sum(it[field] for it in subset))
            cell.number_format = "0.0"
            cell.border = thin
            cell.alignment = center
            cell.fill = PHASE_FILL[key]
        ws.cell(row, 6, f"=ROUND((C{row}+4*D{row}+E{row})/6,1)")
        ws.cell(row, 6).number_format = "0.0"
        ws.cell(row, 6).border = thin
        ws.cell(row, 6).alignment = center
        ws.cell(row, 6).fill = PHASE_FILL[key]
        ws.cell(row, 1).fill = PHASE_FILL[key]
        ws.cell(row, 2).fill = PHASE_FILL[key]
        row += 1
    # ops/legal split
    for key, label, fk in [
        ("OPS", "Of which tagged Business Ops (subset of above)", "OPS"),
        ("LEGAL", "Of which tagged Legal/Privacy (subset of above)", "LEGAL"),
    ]:
        subset = [it for it in items if it["fill"] == fk]
        ws.cell(row, 1, label).font = font_bold
        for col in range(1, 7):
            ws.cell(row, col).fill = PHASE_FILL[fk]
            ws.cell(row, col).border = thin
        ws.cell(row, 2, len(subset)).alignment = center
        ws.cell(row, 3, sum(it["min"] for it in subset)).number_format = "0.0"
        ws.cell(row, 4, sum(it["mid"] for it in subset)).number_format = "0.0"
        ws.cell(row, 5, sum(it["max"] for it in subset)).number_format = "0.0"
        ws.cell(row, 6, f"=ROUND((C{row}+4*D{row}+E{row})/6,1)").number_format = "0.0"
        for col in range(2, 7):
            ws.cell(row, col).alignment = center
        row += 1

    ws.cell(row + 1, 1, "MVP recommendation on this tab = P0 + P1 rows. P2/P3 are sequenced after MVP evidence (pass-rate, CS load, fraud loss).").font = font_small
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)

    ws.oddHeader.left.text = full_title
    return ws, total_row, first_data, last_data


def write_reqs(wb):
    ws = wb.create_sheet("3-Client Requirements")
    apply_sheet_defaults(ws, freeze="A7")
    banner(ws, "Developer / client requirements we need",
           "Inputs that unblock architecture, hours, and go-live. Blocking items in P0/P1 should be requested immediately. This is not a technical design of Veriff — it is the client-side dependency list.",
           cols=4)
    note_row(ws, 3,
             "Give this list to the client PM. Items marked Blocking = Y stop P1 build or UAT if missing. 'When needed' is the latest phase the item must exist; earlier is better. Engineering can proceed on mocks only until Veriff keys + brand identity + prize taxonomy arrive.",
             cols=4, fill=fill_gold)
    headers = ["When needed", "Blocking?", "Requirements", "Why we need it"]
    header_row(ws, 5, headers)
    ws.auto_filter.ref = "A5:D5"
    set_col_widths(ws, [16, 14, 72, 78])

    # sort: phase then blocking
    phase_rank = {"P0": 0, "P1": 1, "P2": 2, "P3": 3}
    ordered = sorted(reqs, key=lambda x: (phase_rank[x[0]], 0 if x[1] == "Y" else 1, x[2]))
    row = 6
    for phase, block, req, why in ordered:
        ws.cell(row, 1, phase).alignment = center
        ws.cell(row, 1).font = font_bold
        ws.cell(row, 1).fill = PHASE_FILL[phase]
        ws.cell(row, 2, block).alignment = center
        ws.cell(row, 2).font = font_bold
        ws.cell(row, 2).fill = fill_red if block == "Y" else fill_p1
        ws.cell(row, 3, req).alignment = wrap
        ws.cell(row, 3).font = font_normal
        ws.cell(row, 4, why).alignment = wrap
        ws.cell(row, 4).font = font_normal
        for c in range(1, 5):
            ws.cell(row, c).border = thin
            if c >= 3:
                ws.cell(row, c).fill = fill_white if row % 2 else fill_alt
        ws.row_dimensions[row].height = 78
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row, 1, "Minimum set to start P1 engineering: Veriff sandbox keys + linked enrollment/auth integrations, brand list, user id scheme, prize catalog + High threshold, winner/redemption API access, staging site with HTTPS.").font = font_small
    ws.row_dimensions[row].height = 36


def write_summary(wb, n1, n2):
    ws = wb.create_sheet("4-Phased Estimate Summary", 4)
    apply_sheet_defaults(ws, freeze="A8")
    banner(ws, "Phased estimation summary",
           "Roll-up of Tabs 1 and 2. Change the yellow inputs (brand count, whether to include P2/P3). MVP = P0 + P1. Additional brands use the unit-cost rows from Tab 2.",
           cols=8)
    note_row(ws, 3,
             "PERT expected is the planning number to use in a SOW discussion: (Min + 4×Mid + Max) / 6. Re-baseline after P0. Do not average Min and Max without the Mid weight — Mid is the expected delivery path.",
             cols=8)

    # Inputs
    ws["A5"] = "INPUTS (edit yellow cells)"
    ws["A5"].font = font_white_bold
    ws["A5"].fill = fill_teal
    ws.merge_cells("A5:C5")
    ws["B5"].fill = fill_teal
    ws["C5"].fill = fill_teal

    ws["A6"] = "Number of brand websites in MVP"
    ws["B6"] = 2
    ws["B6"].fill = fill_input
    ws["B6"].font = font_bold
    ws["B6"].alignment = center
    ws["B6"].border = thin
    ws["B6"].number_format = "0"
    ws["C6"] = "First brand is in the base Tab 2 hours. Extra brands add the unit-cost rows."
    ws["C6"].font = font_small
    ws.merge_cells("C6:H6")

    ws["A7"] = "Include P2 hours in 'Full recommended'?"
    ws["B7"] = "Yes"
    ws["B7"].fill = fill_input
    ws["B7"].font = font_bold
    ws["B7"].alignment = center
    ws["B7"].border = thin
    ws["C7"] = "P2 = medium vouchers, pre-enroll, fraud tooling."
    ws["C7"].font = font_small
    ws.merge_cells("C7:H7")

    ws["A8"] = "Include P3 hours in 'Full recommended'?"
    ws["B8"] = "Yes"
    ws["B8"].fill = fill_input
    ws["B8"].font = font_bold
    ws["B8"].alignment = center
    ws["B8"].border = thin
    ws["C8"] = "P3 = IDV for physical/top-tier, name match, DSR. Stronger 'legit winner' bind."
    ws["C8"].font = font_small
    ws.merge_cells("C8:H8")

    dv1 = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=False)
    dv1.add("B6")
    ws.add_data_validation(dv1)
    dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    dv2.add("B7")
    dv2.add("B8")
    ws.add_data_validation(dv2)

    set_col_widths(ws, [42, 18, 16, 16, 16, 20, 22, 40])

    # Compute python phase totals for MS and Brands (first brand)
    def phase_sums(items):
        out = {}
        for key in ("P0", "P1", "P2", "P3"):
            sub = [it for it in items if it["item"].startswith(f"[{key}]")]
            out[key] = {
                "min": sum(it["min"] for it in sub),
                "mid": sum(it["mid"] for it in sub),
                "max": sum(it["max"] for it in sub),
                "n": len(sub),
            }
        return out

    ms = phase_sums(tab1)
    br = phase_sums(tab2)

    # Incremental brand rows — identified by name
    inc_items = [it for it in tab2 if "Per additional brand" in it["item"]]
    inc = {
        "min": sum(it["min"] for it in inc_items),
        "mid": sum(it["mid"] for it in inc_items),
        "max": sum(it["max"] for it in inc_items),
    }
    # Base brand hours = all tab2 minus ONE copy of incremental (incremental is already in P1 totals once as unit cost)
    # Interpretation: tab2 P1 includes 1× incremental row as the template. For N brands:
    #   brand_hours = (tab2 totals) + (N-1)*inc   would double-count the unit row already in the sheet.
    # Better: treat the unit-cost rows as NOT part of the first brand. First brand = tab2 minus inc rows.
    def minus_inc(phase_key):
        sub = [it for it in tab2 if it["item"].startswith(f"[{phase_key}]") and "Per additional brand" not in it["item"]]
        return {
            "min": sum(it["min"] for it in sub),
            "mid": sum(it["mid"] for it in sub),
            "max": sum(it["max"] for it in sub),
        }

    br_base = {k: minus_inc(k) for k in ("P0", "P1", "P2", "P3")}
    # incremental is P1
    inc_p1 = inc

    # Write lookup table starting at row 10 (static numbers from tabs — documented)
    ws["A10"] = "BASE HOURS (from Tabs 1–2, first brand only; incremental brand rows excluded from base)"
    ws["A10"].font = font_white_bold
    ws["A10"].fill = fill_navy
    ws.merge_cells("A10:H10")
    for c in range(1, 9):
        ws.cell(10, c).fill = fill_navy

    headers = ["Workstream / phase", "Count of items", "Min (hrs)", "Mid (hrs)", "Max (hrs)", "Expected PERT (hrs)", "Notes", ""]
    header_row(ws, 11, headers)

    rows_meta = []
    r = 12
    blocks = [
        ("Microservice P0 Discovery", ms["P0"], "Vendor, architecture, identity, policy workshop", "P0"),
        ("Microservice P1 MVP", ms["P1"], "Service, Veriff, lock, admin, CS/fraud/warehouse ops", "P1"),
        ("Microservice P2 Expand", ms["P2"], "Pre-enroll, medium band, FaceBlock, analytics", "P2"),
        ("Microservice P3 Harden", ms["P3"], "IDV enrollment, name match, DSR, DR", "P3"),
        ("Brand sites P0 Discovery", br_base["P0"], "Journey mapping + T&Cs (first brand)", "P0"),
        ("Brand sites P1 MVP (first brand)", br_base["P1"], "Shared kit + first website UX/QA (excludes extra-brand unit)", "P1"),
        ("Brand sites P2 Expand (first brand)", br_base["P2"], "Pre-enroll CTA, medium UI, theming", "P2"),
        ("Brand sites P3 Harden (first brand)", br_base["P3"], "IDV UX, step-up login", "P3"),
        ("UNIT extra brand (P1 ops+eng)", {"min": inc_p1["min"], "mid": inc_p1["mid"], "max": inc_p1["max"], "n": len(inc_items)},
         "Multiply by (N − 1). Wire-up + CS checklist.", "P1"),
    ]
    # item counts
    def count(items, key, exclude_inc=False):
        n = 0
        for it in items:
            if not it["item"].startswith(f"[{key}]"):
                continue
            if exclude_inc and "Per additional brand" in it["item"]:
                continue
            n += 1
        return n

    counts = [
        count(tab1, "P0"),
        count(tab1, "P1"),
        count(tab1, "P2"),
        count(tab1, "P3"),
        count(tab2, "P0", True),
        count(tab2, "P1", True),
        count(tab2, "P2", True),
        count(tab2, "P3", True),
        len(inc_items),
    ]

    named_rows = {}
    for i, (label, sums, note, pk) in enumerate(blocks):
        ws.cell(r, 1, label).font = font_bold
        ws.cell(r, 2, counts[i]).alignment = center
        ws.cell(r, 3, sums["min"]).number_format = "0.0"
        ws.cell(r, 4, sums["mid"]).number_format = "0.0"
        ws.cell(r, 5, sums["max"]).number_format = "0.0"
        ws.cell(r, 6, f"=ROUND((C{r}+4*D{r}+E{r})/6,1)").number_format = "0.0"
        ws.cell(r, 7, note).font = font_small
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        for c in range(1, 9):
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = PHASE_FILL[pk]
            if c in (2, 3, 4, 5, 6):
                ws.cell(r, c).alignment = center
        named_rows[label] = r
        r += 1

    r = 22
    ws["A22"] = "SCALED TOTALS (formulas use B6 brand count and B7/B8 flags)"
    ws["A22"].font = font_white_bold
    ws["A22"].fill = fill_teal
    ws.merge_cells("A22:H22")
    for c in range(1, 9):
        ws.cell(22, c).fill = fill_teal

    header_row(ws, 23, ["Package", "How it is calculated", "Min (hrs)", "Mid (hrs)", "Max (hrs)", "Expected PERT (hrs)", "Person-weeks (PERT / 32h)", "Use for"])

    # Row refs
    r_ms_p0, r_ms_p1, r_ms_p2, r_ms_p3 = 12, 13, 14, 15
    r_br_p0, r_br_p1, r_br_p2, r_br_p3 = 16, 17, 18, 19
    r_inc = 20

    def pkg_row(row, name, how, minf, midf, maxf, use):
        ws.cell(row, 1, name).font = font_bold
        ws.cell(row, 2, how).alignment = wrap
        ws.cell(row, 2).font = font_small
        ws.cell(row, 3, minf).number_format = "0.0"
        ws.cell(row, 4, midf).number_format = "0.0"
        ws.cell(row, 5, maxf).number_format = "0.0"
        ws.cell(row, 6, f"=ROUND((C{row}+4*D{row}+E{row})/6,1)").number_format = "0.0"
        ws.cell(row, 7, f"=ROUND(F{row}/32,1)").number_format = "0.0"
        ws.cell(row, 8, use).alignment = wrap
        ws.cell(row, 8).font = font_small
        for c in range(1, 9):
            ws.cell(row, c).border = thin
            ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical="center",
                                                  horizontal="center" if c >= 3 and c <= 7 else "left")
        ws.row_dimensions[row].height = 48

    extra = f'MAX(0,$B$6-1)'
    pkg_row(
        24, "MVP (P0+P1) — recommended start",
        "MS P0+P1 + Brand P0+P1 + extra brands × unit",
        f"=C{r_ms_p0}+C{r_ms_p1}+C{r_br_p0}+C{r_br_p1}+{extra}*C{r_inc}",
        f"=D{r_ms_p0}+D{r_ms_p1}+D{r_br_p0}+D{r_br_p1}+{extra}*D{r_inc}",
        f"=E{r_ms_p0}+E{r_ms_p1}+E{r_br_p0}+E{r_br_p1}+{extra}*E{r_inc}",
        "SOW / sprint-0 proposal. Physical + high-value digital only.",
    )
    for c in range(1, 9):
        ws.cell(24, c).fill = fill_p1

    pkg_row(
        25, "P2 increment",
        "MS P2 + Brand P2 (first brand; extra brand P2 not separately unit-costed — absorb in extra-brand max)",
        f"=C{r_ms_p2}+C{r_br_p2}",
        f"=D{r_ms_p2}+D{r_br_p2}",
        f"=E{r_ms_p2}+E{r_br_p2}",
        "After MVP metrics: medium vouchers + pre-enroll.",
    )
    for c in range(1, 9):
        ws.cell(25, c).fill = fill_p2

    pkg_row(
        26, "P3 increment",
        "MS P3 + Brand P3",
        f"=C{r_ms_p3}+C{r_br_p3}",
        f"=D{r_ms_p3}+D{r_br_p3}",
        f"=E{r_ms_p3}+E{r_br_p3}",
        "Stronger legit-winner bind (IDV + name match).",
    )
    for c in range(1, 9):
        ws.cell(26, c).fill = fill_p3

    pkg_row(
        27, "Full recommended (flags in B7/B8)",
        "MVP + P2 if Yes + P3 if Yes",
        '=C24+IF(B7="Yes",C25,0)+IF(B8="Yes",C26,0)',
        '=D24+IF(B7="Yes",D25,0)+IF(B8="Yes",D26,0)',
        '=E24+IF(B7="Yes",E25,0)+IF(B8="Yes",E26,0)',
        "Upper planning envelope if all phases are approved.",
    )
    for c in range(1, 9):
        ws.cell(27, c).fill = fill_navy
        ws.cell(27, c).font = font_total if c != 2 else Font(name="Calibri", size=9, color=WHITE, italic=True)
    ws.cell(27, 1).font = font_total
    ws.cell(27, 8).font = Font(name="Calibri", size=9, color=WHITE)

    # Chart data
    ws["A29"] = "MVP mix (Mid hrs, first 2 brands default) — illustration"
    ws["A29"].font = font_navy_h
    ws["A30"] = "Microservice P0+P1"
    ws["B30"] = f"=D{r_ms_p0}+D{r_ms_p1}"
    ws["A31"] = "Brand kit + first site P0+P1"
    ws["B31"] = f"=D{r_br_p0}+D{r_br_p1}"
    ws["A32"] = "Extra brands (N-1)×unit"
    ws["B32"] = f"={extra}*D{r_inc}"
    for r in range(30, 33):
        ws.cell(r, 2).number_format = "0.0"
        ws.cell(r, 1).font = font_normal

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "MVP Mid hours mix"
    chart.y_axis.title = None
    chart.x_axis.title = "Hours"
    data = Reference(ws, min_col=2, min_row=30, max_row=32)
    cats = Reference(ws, min_col=1, min_row=30, max_row=32)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.width = 18
    chart.height = 7
    ws.add_chart(chart, "D29")

    ws["A40"] = "Planning notes"
    ws["A40"].font = font_navy_h
    notes = [
        "Person-weeks use 32 productive hours/week (meetings, env waits, vendor). Calendar elapsed time is longer (legal, Veriff SE, 3PL, sequential UAT).",
        "Do not start Tab 2 pixel-perfect UX until P0 identity + threshold table are signed — otherwise enrollment vs auth branching will be rebuilt.",
        "If native apps are in scope, add a separate delta (Veriff iOS/Android SDKs, review, camera permissions) — not in these hours.",
        "If IDV is mandated for MVP physical prizes, add Microservice P3 IDV + Brand P3 IDV UX into the MVP package (use Max or pull those two rows forward).",
        "Veriff session fees are OPEX, not in these hours. Finance should compare fee × expected gated winners vs current fraud loss on physical/high-value.",
    ]
    for i, n in enumerate(notes):
        rr = 41 + i
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=8)
        ws.cell(rr, 1, "• " + n).font = font_normal
        ws.cell(rr, 1).alignment = wrap
        ws.row_dimensions[rr].height = 32


def write_flow(wb):
    ws = wb.create_sheet("5-Process Decision Flow")
    apply_sheet_defaults(ws, freeze="A7")
    banner(ws, "Recommended process / decision flow",
           "Use this as the estimation workshop script and later as epic/user-story backbone. Webhook is the only fulfill signal. Unenrolled users are a designed branch.",
           cols=7)
    note_row(ws, 3,
             "Actors: Website (brand), Microservice (policy + Veriff + lock), Veriff, User, Warehouse/voucher issuer, CS/Fraud ops. "
             "MVP applies this flow only when policy says bio required (physical + high-value digital). Small vouchers skip to existing fulfill after step 2 = NO.",
             cols=7, fill=fill_gold)
    headers = ["Step", "Control type", "Actor", "Process / decision", "If Yes / Success", "If No / Fail / Else", "Phase"]
    header_row(ws, 5, headers)
    ws.auto_filter.ref = "A5:G5"
    set_col_widths(ws, [10, 18, 28, 55, 48, 52, 12])

    row = 6
    for step, ctype, actor, decision, yes, no, phase in flow:
        ws.cell(row, 1, step).alignment = center
        ws.cell(row, 1).font = font_bold
        ws.cell(row, 2, ctype).alignment = wrap_c
        ws.cell(row, 3, actor).alignment = wrap
        ws.cell(row, 4, decision).alignment = wrap
        ws.cell(row, 5, yes).alignment = wrap
        ws.cell(row, 6, no).alignment = wrap
        ws.cell(row, 7, phase).alignment = center
        ws.cell(row, 7).fill = PHASE_FILL[phase]
        ws.cell(row, 7).font = font_bold
        for c in range(1, 8):
            ws.cell(row, c).border = thin
            ws.cell(row, c).font = font_normal if c != 1 else font_bold
            if c != 7:
                # highlight freeze/fulfill
                fill = fill_white if row % 2 else fill_alt
                if ctype in ("Fraud freeze",) or step == "13":
                    fill = fill_red
                if step in ("11", "12a", "12b"):
                    fill = fill_p1
                if step == "4":
                    fill = fill_p2
                if c == 1:
                    ws.cell(row, c).fill = fill_navy
                    ws.cell(row, c).font = font_white_bold
                else:
                    ws.cell(row, c).fill = fill
        ws.row_dimensions[row].height = 78
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    ws.cell(row, 1, "Same-sitting rule: approved ENROLLMENT may unlock the current gated prize. A later prize for the same user uses AUTHENTICATION only. Do not run enroll+auth back-to-back for one prize (double Veriff fee, worse UX).").font = font_small
    ws.row_dimensions[row].height = 32
    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    ws.cell(row, 1, "Mermaid (for Confluence / PR / workshop whiteboard) — also in the accompanying chat note").font = font_navy_h
    row += 1
    mermaid = """flowchart TD
  A[User claims reward] --> B{Logged in as winning account?}
  B -->|No| Z1[Login / recovery — no Veriff]
  B -->|Yes| C{Policy: physical OR digital >= High?}
  C -->|No small voucher| Z2[Existing fulfill — no Veriff]
  C -->|Yes gated| D{Biometric consent current?}
  D -->|No| D1{Accepts consent?}
  D1 -->|No| Z3[Hold until SLA then forfeit / substitute]
  D1 -->|Yes| E
  D -->|Yes| E{Approved enrollment exists?}
  E -->|No| F{Camera / supported browser?}
  F -->|No| X[Exception queue]
  F -->|Yes| G[Create ENROLLMENT session]
  G --> H{Enrollment webhook}
  H -->|approved| K[Fulfill gate]
  H -->|quality fail| G
  H -->|fraud / mismatch| FZ[Fraud freeze]
  H -->|abandon / expire| SLA{Claim SLA left?}
  SLA -->|Yes| G
  SLA -->|No| Z4[Forfeit]
  E -->|Yes| I[Create AUTH session same endUserId]
  I --> J{Auth webhook + matched user}
  J -->|approved match| K
  J -->|quality fail| I
  J -->|122 no reference| E
  J -->|120 / 105 / 106| FZ
  J -->|abandon / expire| SLA2{Claim SLA left?}
  SLA2 -->|Yes| I
  SLA2 -->|No| Z4
  K{Atomic prize still free?}
  K -->|No| Z5[Ignore — already fulfilled]
  K -->|Yes physical| P[Address then WMS RELEASE_TO_SHIP]
  K -->|Yes digital high| V[Issue voucher once then reveal]
"""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    cell = ws.cell(row, 1, mermaid)
    cell.alignment = wrap
    cell.font = Font(name="Consolas", size=8, color="1A1A1A")
    cell.fill = fill_cover
    ws.row_dimensions[row].height = 420


def write_risks(wb):
    ws = wb.create_sheet("6-Assumptions & Risks")
    apply_sheet_defaults(ws, freeze="A6")
    banner(ws, "Assumptions and risks",
           "If an assumption is false, re-estimate. Risks listed here are the ways a non-winner still gets the prize even after Veriff is integrated.",
           cols=4)
    ws["A4"] = "Assumptions"
    ws["A4"].font = font_white_bold
    ws["A4"].fill = fill_navy
    ws.merge_cells("A4:D4")
    for c in range(1, 5):
        ws.cell(4, c).fill = fill_navy
    header_row(ws, 5, ["Theme", "Assumption", "", ""])
    ws.merge_cells("B5:D5")
    set_col_widths(ws, [18, 40, 40, 42])
    row = 6
    for theme, text in assumptions:
        ws.cell(row, 1, theme).font = font_bold
        ws.cell(row, 1).fill = fill_p0
        ws.cell(row, 1).border = thin
        ws.cell(row, 1).alignment = wrap_m
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row, 2, text).alignment = wrap
        ws.cell(row, 2).font = font_normal
        for c in range(2, 5):
            ws.cell(row, c).border = thin
            ws.cell(row, c).fill = fill_white if row % 2 else fill_alt
        ws.row_dimensions[row].height = 48
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row, 1, "Risks — how a non-winner still gets the prize").font = font_white_bold
    ws.cell(row, 1).fill = fill_navy
    for c in range(1, 5):
        ws.cell(row, c).fill = fill_navy
    row += 1
    header_row(ws, row, ["Severity", "Risk", "How it bypasses bio", "Mitigation in this solution"])
    row += 1
    sev_fill = {"Critical": fill_red, "High": fill_legal, "Medium": fill_p2, "Low": fill_p1}
    for sev, risk, bypass, mit in risks:
        ws.cell(row, 1, sev).font = font_bold
        ws.cell(row, 1).fill = sev_fill[sev]
        ws.cell(row, 1).alignment = center
        ws.cell(row, 2, risk).alignment = wrap
        ws.cell(row, 3, bypass).alignment = wrap
        ws.cell(row, 4, mit).alignment = wrap
        for c in range(1, 5):
            ws.cell(row, c).border = thin
            ws.cell(row, c).font = font_normal if c != 1 else font_bold
            if c > 1:
                ws.cell(row, c).fill = fill_white if row % 2 else fill_alt
        ws.row_dimensions[row].height = 56
        row += 1

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row, 1, "End-state control stack for 'only a legit winner'").font = font_navy_h
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row, 1,
            "Account bind + claim-time login + (IDV enrollment for physical) + biometric auth on later claims + atomic fulfill lock + 3PL/voucher default-deny + dual-control exceptions + fraud freeze on mismatch. "
            "Removing any layer re-opens a bypass. Veriff is necessary but not sufficient.").alignment = wrap
    ws.row_dimensions[row].height = 48


def main():
    wb = Workbook()
    write_cover(wb)
    write_work_sheet(
        wb,
        "1-Microservice Bio Auth",
        "Tab 1 — Microservice Veriff Biometric Authentication",
        "Shared service used by all brand websites. Includes engineering, Veriff integration, policy/lock, AND business operations (CS, fraud, warehouse, comms, training).",
        tab1,
        "Do not put Veriff API keys on websites. Enrollment integration and Authentication integration are separate Veriff products that must be linked. Unenrolled users must be pre-checked — never start Auth first.",
    )
    write_work_sheet(
        wb,
        "2-Brand Websites",
        "Tab 2 — Brand websites",
        "Shared integration kit + first brand UX, then a unit cost per additional brand. Covers unenrolled UX, physical vs high-value digital, and the rule: never fulfill on the Veriff redirect callback.",
        tab2,
        "Small-voucher pages must keep working WITHOUT bio (regression). Physical address and high-value voucher codes appear only AFTER microservice Approved. Multiply 'Per additional brand' rows by (N − 1) — the summary tab does this.",
    )
    write_reqs(wb)
    write_summary(wb, len(tab1), len(tab2))
    write_flow(wb)
    write_risks(wb)

    # print areas / freeze already set
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
        ws.page_setup.horizontalCentered = False

    wb.save(OUT)
    print("Wrote", OUT)
    print("Tab1 items", len(tab1), "Tab2 items", len(tab2), "Reqs", len(reqs), "Flow", len(flow))
    def tot(items):
        return sum(i["min"] for i in items), sum(i["mid"] for i in items), sum(i["max"] for i in items)
    print("Tab1 hours min/mid/max", tot(tab1))
    print("Tab2 hours min/mid/max", tot(tab2))


if __name__ == "__main__":
    main()
