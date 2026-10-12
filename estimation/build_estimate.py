#!/usr/bin/env python3
"""Build the Veriff biometric rewards-redemption estimation workbook."""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "/workspace/estimation/Veriff-Biometric-Rewards-Estimation.xlsx"

NAVY = "1F4E79"
TEAL = "0E6B67"
INK = "1D1D1F"
MUTED = "5C6770"
PALE = "F4F7FB"
PALE_TEAL = "E6F3F2"
PALE_NAVY = "E7EEF6"
GREEN = "E5F4EA"
AMBER = "FFF4D6"
ORANGE = "FDE8D8"
RED = "F8D6D6"
WHITE = "FFFFFF"
YELLOW = "FFF2CC"
LINE = "D5DDE6"
TOTAL = "163A5F"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
thick_top = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="medium", color=NAVY),
    bottom=Side(style="thin", color=LINE),
)

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_total = PatternFill("solid", fgColor=TOTAL)
fill_pale = PatternFill("solid", fgColor=PALE)
fill_pale_navy = PatternFill("solid", fgColor=PALE_NAVY)
fill_pale_teal = PatternFill("solid", fgColor=PALE_TEAL)
fill_green = PatternFill("solid", fgColor=GREEN)
fill_amber = PatternFill("solid", fgColor=AMBER)
fill_orange = PatternFill("solid", fgColor=ORANGE)
fill_red = PatternFill("solid", fgColor=RED)
fill_yellow = PatternFill("solid", fgColor=YELLOW)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_section = PatternFill("solid", fgColor="D6E3F0")

font_title = Font(name="Calibri", size=20, bold=True, color=NAVY)
font_h = Font(name="Calibri", size=14, bold=True, color=NAVY)
font_h2 = Font(name="Calibri", size=12, bold=True, color=TEAL)
font_white = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_white_s = Font(name="Calibri", size=10, bold=True, color=WHITE)
font_body = Font(name="Calibri", size=11, color=INK)
font_small = Font(name="Calibri", size=10, color=INK)
font_muted = Font(name="Calibri", size=10, italic=True, color=MUTED)
font_bold = Font(name="Calibri", size=11, bold=True, color=INK)
font_link = Font(name="Calibri", size=11, color="0B5CAB", underline="single")
font_mono = Font(name="Consolas", size=10, color=INK)

wrap = Alignment(wrap_text=True, vertical="top")
wrap_mid = Alignment(wrap_text=True, vertical="center")
left_mid = Alignment(vertical="center", wrap_text=True)

PHASE_FILL = {
    "MVP": fill_green,
    "Phase 2": fill_pale_navy,
    "Phase 3": fill_amber,
}


def apply_header_footer(ws, subtitle):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.oddHeader.left.text = "Veriff biometric authentication — rewards redemption"
    ws.oddHeader.right.text = subtitle
    ws.oddFooter.left.text = "Planning estimate for workshops. Effort hours, not a quote or a schedule."
    ws.oddFooter.right.text = "Page &P of &N"
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.7
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.sheet_properties.tabColor = NAVY


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def paint(cell, font=None, fill=None, alignment=None, border=None, fmt=None):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt


def section_bar(ws, row, start, end, text, fill=fill_teal):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    cell = ws.cell(row, start, text)
    paint(cell, font_white, fill, Alignment(vertical="center", indent=1))
    for col in range(start, end + 1):
        paint(ws.cell(row, col), fill=fill, border=thin)
    ws.row_dimensions[row].height = 22
    return row


def write_title(ws, row, col, text, merge_end):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_end)
    cell = ws.cell(row, col, text)
    paint(cell, font_title, alignment=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = 28
    return row


# ---------------------------------------------------------------------------
# Line items
# Columns: work item, description, min, mid, max, phase, section, workstream,
#          basis, include, id, notes
# ---------------------------------------------------------------------------

MS = [
    (
        "Current-state discovery of win, redemption, and fulfillment",
        "Map how the pilot brand records a win, signs the winner in, issues a digital voucher, and ships a physical prize. Name the system of record for the winner, the prize, the voucher code, and the shipment. Outcome is a signed as-is flow and the integration points the biometric gate must call.",
        16, 24, 40, "MVP", "Discovery", "Discovery and design", "Once", "Yes", "MS-01",
        "Re-forecast the build items toward Min if this discovery finds a clean redemption API.",
    ),
    (
        "Redemption threat model and control matrix",
        "Walk account takeover, shared logins, an address change after the win, voucher interception, multi-account farming, replay of an old Veriff result, presentation attacks (screen, print, or streamed video), and a support agent releasing a prize after a social-engineering call. Each case gets a control, an owner, and an MVP or later tag.",
        12, 20, 32, "MVP", "Identity and policy", "Security and privacy", "Once", "Yes", "MS-02",
        "Pairs with the decision flow. The matrix is the security acceptance list for go-live.",
    ),
    (
        "Prize-tier policy",
        "Write the MVP rule the business will sign: a biometric check for every physical prize, and for digital vouchers at or above a face value the client sets. Smaller digital vouchers stay on today's redemption path. The threshold is configuration per brand and campaign, with a named owner for every change.",
        8, 16, 24, "MVP", "Identity and policy", "Discovery and design", "Once", "Yes", "MS-03",
        "Use the cost model to set the number. Do not hard-code a currency amount.",
    ),
    (
        "Journeys for enrolled winners and winners with no biometric",
        "Specify both primary journeys. A winner with no approved enrollment completes consent and a document-and-selfie enrollment before that prize can be released. A winner who is already enrolled completes a selfie biometric authentication only. Include consent refusal, camera denial, resubmission, expiry, abandonment, a fraud decline, and finish-later.",
        16, 24, 40, "MVP", "Discovery", "Discovery and design", "Once", "Yes", "MS-04",
        "Treat 'not enrolled' as the launch default. Most winners will be in that state on day one.",
    ),
    (
        "Biometric subject identity standard",
        "Define an immutable UUID v4 endUserId for each brand plus winner account. The microservice generates and stores it, and sends the same value on every enrollment and authentication session. Veriff returns it unchanged. A later session with a different identifier is declined (reason 128, or 121 if vendor data is used). Winner account ids must never be recycled. MVP keeps a separate subject per brand.",
        8, 12, 20, "MVP", "Identity and policy", "Discovery and design", "Once", "Yes", "MS-05",
        "Prefer endUserId. Use vendorData only if Veriff requires a non-UUID identifier. vendorData cannot be changed once sent.",
    ),
    (
        "Release authority and trust boundaries",
        "Record the architecture decision. Only this microservice may authorize a prize release. The browser return from Veriff is untrusted. Decision webhooks are accepted only after HMAC verification. A missing brand configuration, a bad signature, or a Veriff outage leaves the prize unredeemed. Each prize can be released once.",
        12, 16, 28, "MVP", "Release control", "Security and privacy", "Once", "Yes", "MS-06",
        "Brand websites display status. They do not decide it.",
    ),
    (
        "Veriff solution design with their Solutions Engineer",
        "Specify the two linked integrations Veriff must configure: a selfie enrollment integration, and a biometric authentication integration connected to it, in sandbox and in production. Enrollment is document and selfie, so the stored template is tied to a document, not to a selfie alone. Confirm webhook URLs and whether existing approved Veriff sessions can be migrated into enrollment.",
        12, 20, 32, "MVP", "Enrollment and authentication", "Discovery and design", "Once", "Yes", "MS-07",
        "These are design and coordination hours. Calendar time waiting on Veriff is not included.",
    ),
    (
        "Privacy, consent, retention, and regional constraints",
        "Design explicit consent, a purpose limited to in-scope prize redemption, the minimum age, and retention. Persist the decision, reason code, session ids, and consent record. Keep face images, document images, and biometric templates at Veriff under the client contract. Counsel owns the notice text and the privacy impact assessment; this item is the technical design and the review cycles.",
        12, 20, 36, "MVP", "Identity and policy", "Security and privacy", "Once", "Yes", "MS-08",
        "Regional biometric rules can change scope. The go/no-go is a client legal decision listed on the requirements tab.",
    ),
    (
        "Operating model, dual control, and hold SLAs",
        "Write the RACI across support, fraud, warehouse, the voucher issuer, and engineering. A manual release requires two different named roles. Define what the winner is told while a claim is held, how often overrides are sampled, and how a shared inbox is kept off the approval path.",
        12, 20, 32, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-09",
        "Build of the approval API is MS-22. This item is the operating design those screens must enforce.",
    ),
    (
        "API contract and sequence diagrams",
        "Publish the contract brand backends will call: create or resume a challenge, read its status, and receive signed callbacks. Include the idempotency key, the error catalog, and the sequences for a first enrollment and for a later authentication.",
        12, 20, 32, "MVP", "Discovery", "Discovery and design", "Once", "Yes", "MS-10",
        "Status values must match the decision-flow sheet so every brand uses the same language.",
    ),
    (
        "Multi-tenant service foundation",
        "Stand up one service that later brands share. Include brand configuration, separate secret storage for enrollment and authentication credentials, sandbox and production, service-to-service authentication from brand backends, correlation ids, and health checks.",
        24, 32, 56, "MVP", "Platform and security", "Backend engineering", "Once", "Yes", "MS-11",
        "One Veriff customer with linked enrollment and authentication integrations is the assumed shape. Confirm in MS-07.",
    ),
    (
        "Veriff session client",
        "Create sessions on the correct integration and base URL, always send the stored endUserId, set the callback, and map transport errors. When a legal name is on the frozen prize (the shipping name, or a name collected for a high-value voucher), send that name on the enrollment session so Veriff can return a name mismatch (reason 104).",
        16, 24, 40, "MVP", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-12",
        "Authentication sessions use the authentication integration keys. Enrollment sessions use the enrollment keys.",
    ),
    (
        "Enrollment orchestration for winners with no biometric",
        "When the subject has no approved enrollment, start a document-and-selfie session and store it against the redemption challenge. An approval completed for this challenge can release this prize after the binding checks, and it becomes the reference for later claims. Do not start biometric authentication with no reference on file; Veriff declines that case with reason 122.",
        20, 36, 56, "MVP", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-13",
        "A second selfie immediately after enrollment is a Phase 2 tightening, not the MVP. The enrollment capture is the liveness and document check for this claim.",
    ),
    (
        "Biometric authentication for an already enrolled winner",
        "Create the authentication session with the authentication credentials and the same endUserId. Return the session URL and the challenge id to the brand site. Track attempts. If the decision is reason 122 (no reference), route the winner back to enrollment.",
        16, 28, 44, "MVP", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-14",
        "This is the short path for the second and later in-scope prizes.",
    ),
    (
        "Decision webhook intake",
        "Verify the HMAC signature with the integration secret, acknowledge quickly, and process duplicates safely. Treat verification.status as the source of truth, because expired and abandoned can both arrive with code 9104. Read biometricAuthentication matched session and matched end-user id. Keep the challenge pending through resubmission; for SDK sessions Veriff can delay the resubmission_requested webhook by about five minutes. The decision webhook is the only automated release trigger. Event webhooks (started 7001, submitted 7002) are optional progress signals.",
        16, 28, 44, "MVP", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-15",
        "Unsigned or unknown payloads are stored as failures and never move a prize.",
    ),
    (
        "Decision reconciliation",
        "Poll GET /sessions/{id}/decision when a webhook is missing or late. Align Veriff's seven-day session lifetime with the shorter business hold operations agree. A poll that times out leaves the prize unredeemed.",
        8, 16, 24, "MVP", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-16",
        "Required because a public webhook can be missed even when signature checks are correct.",
    ),
    (
        "Redemption challenge state machine",
        "Represent consent, enrollment in progress, authentication in progress, resubmission, fraud hold, name-mismatch hold, retryable decline, expiry, abandonment, dual-control release, rejection, and one-time released. Reject illegal transitions and a second release of the same prize.",
        20, 32, 52, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-17",
        "The state machine is the control. UI copy must not invent extra states.",
    ),
    (
        "Binding rules and reason-code policy",
        "Release only when brand, winner account, prize, challenge, and endUserId still match the frozen payout. Hold the prize when the matched enrollment belongs to a different end user. Fraud-stop reasons, with no self-serve retry: 105, 106, 108, 109, 110, 113, 124, and 127. Treat 121 and 128 as integration defects and alert engineering. Treat 104 as a name-mismatch hold. Treat 120 as a limited retry. Treat 122 as not enrolled.",
        16, 28, 44, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-18",
        "Reason meanings are taken from Veriff's decision-code tables. Confirm the enabled checks with the Solutions Engineer.",
    ),
    (
        "Retry, cooldown, and same-session resubmission",
        "For resubmission_requested (9103), send the winner back to the same session URL. Veriff allows a limited number of resubmissions and then declines (reason 539); a further try needs a new session. Give quality failures a small self-serve budget. Fraud holds have no self-serve retry.",
        8, 16, 24, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-19",
        "A new session after a fraud decline is an operations decision, not a website button.",
    ),
    (
        "Prize freeze and exactly-once release",
        "Starting an in-scope challenge locks the physical recipient, address, and phone, and blocks voucher minting. After binding checks pass, emit one release. A failed or abandoned check leaves the prize won and unredeemed. The website success page cannot mint a code or create a shipment.",
        16, 24, 40, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-20",
        "Fulfillment systems must ignore any request that does not carry this release.",
    ),
    (
        "Signed callbacks to the brand, warehouse, and voucher issuer",
        "Deliver at-least-once callbacks with an idempotency key for pending, enrollment required, released, declined, and needs a manual decision. The warehouse and the voucher issuer act on this message. A repeated delivery must not ship a second parcel or mint a second code.",
        16, 24, 40, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-21",
        "Depends on sandbox APIs from the client requirements list.",
    ),
    (
        "Operations case API, timeline, and dual approval",
        "Authorised staff can search by brand, winner, prize, and Veriff session, and can see the status and reason code. The API does not return face or document images. A manual release stores both approvers, the reason, and an evidence reference, and it is rejected when the same person submits both steps.",
        16, 28, 48, "MVP", "Business operations", "Backend engineering", "Once", "Yes", "MS-22",
        "Enforces the operating model in MS-09. Image access, if ever allowed, is a later legal decision.",
    ),
    (
        "Consent record, withdrawal, and deletion",
        "Store the notice version, time, brand, and subject. Withdrawal stops new biometric challenges and leaves the prize unredeemed. Deletion removes our metadata and requests Veriff session deletion where the contract allows. A scheduled job applies the retention period counsel sets.",
        12, 20, 32, "MVP", "Platform and security", "Backend engineering", "Once", "Yes", "MS-23",
        "A withdrawn consent is not a fraud decline. The winner can be offered the non-biometric policy only if the client defined one.",
    ),
    (
        "Observability and alerts",
        "Dashboards for the decision mix, reason codes, webhook lag, time to release, signature failures, and identifier-mismatch spikes. Alert on those spikes and on any attempt to release without an approval. Logs carry correlation ids and omit images and document numbers.",
        8, 16, 24, "MVP", "Platform and security", "Backend engineering", "Once", "Yes", "MS-24",
        "The pilot scorecard in MS-39 reads from these metrics.",
    ),
    (
        "Abuse protection on the brand-facing API",
        "Rate-limit create-challenge, allow one open challenge per prize, reject replayed requests, and issue short-lived service credentials per brand. A brand cannot read another brand's winners.",
        8, 16, 24, "MVP", "Platform and security", "Security and privacy", "Once", "Yes", "MS-25",
        "Website session security stays with each brand and is a client requirement.",
    ),
    (
        "Threshold configuration and kill switch",
        "Configure, per brand and campaign: physical prizes in scope, the digital face-value threshold, a pilot allow-list, and a kill switch. Missing configuration fails closed. Every change is audited.",
        8, 12, 20, "MVP", "Release control", "Backend engineering", "Once", "Yes", "MS-26",
        "Phase 2 adds business-user screens. MVP can be configuration managed by engineering, with the audit trail still required.",
    ),
    (
        "Persistence, backup, and retention",
        "Transactional storage for subjects and challenges, a tested backup restore, and the retention job. The schema has no column for a raw selfie, a document image, or a biometric template.",
        8, 16, 24, "MVP", "Platform and security", "Backend engineering", "Once", "Yes", "MS-27",
        "Retention duration comes from the client legal requirement.",
    ),
    (
        "Automated tests for the decision matrix",
        "Cover every Veriff status, the shared 9104 code for expired and abandoned, duplicate webhooks, a mismatched endUserId, a second release, reason 122 routing back to enrollment, and each fraud-stop reason. Tests must fail if a browser-supplied status can release a prize.",
        24, 40, 56, "MVP", "Quality", "Quality assurance", "Once", "Yes", "MS-28",
        "Fixtures follow Veriff's published webhook samples, including the biometricAuthentication object.",
    ),
    (
        "Sandbox proving and go-live evidence",
        "Veriff's test integration does not decide sessions automatically; a tester sets the result in the Customer Portal. Script approved, name mismatch, no reference, face mismatch, fraud decline, resubmission, expired, and abandoned. Save the evidence pack for go-live.",
        12, 20, 36, "MVP", "Quality", "Quality assurance", "Once", "Yes", "MS-29",
        "Needs portal access and test users from the client.",
    ),
    (
        "Environments, deployment, and webhook exposure",
        "Continuous delivery for sandbox and production, secret rotation, TLS, and a stable public decision-webhook URL. Infrastructure definitions cover the service and its datastore.",
        12, 20, 36, "MVP", "Platform and security", "Backend engineering", "Once", "Yes", "MS-30",
        "Webhook host ownership is a client requirement.",
    ),
    (
        "Security review and remediation window",
        "Review HMAC verification, log redaction, who can authorize a release, and cross-brand access. The maximum includes time to fix findings from that review. Support for a client-commissioned penetration test is included; the test fee and the tester's time are not.",
        8, 16, 32, "MVP", "Platform and security", "Security and privacy", "Once", "Yes", "MS-31",
        "Schedule the review before pilot traffic, with a fix window still inside the pilot plan.",
    ),
    (
        "Engineering runbooks",
        "Operate a webhook backlog, a Veriff outage, a stuck challenge, and key rotation. Show how to prove a prize was not released. These runbooks are for engineering; business procedures are separate items.",
        6, 10, 16, "MVP", "Quality", "Backend engineering", "Once", "Yes", "MS-32",
        "Store them where the pilot on-call rota will actually look.",
    ),
    (
        "Capacity test of the gate",
        "Test create-challenge, status reads, and webhook bursts at the client's stated peak. The winner's wait inside Veriff is outside this test. The acceptance target is that our gate stays correct under duplicate delivery.",
        6, 12, 20, "MVP", "Quality", "Quality assurance", "Once", "Yes", "MS-33",
        "Peak volumes are a client input on the requirements tab.",
    ),
    (
        "Fraud operations procedure and training",
        "Job aid and a training session for fraud reasons, name mismatch, a duplicate face (124), and a suspected mule. Include the path used when a shipment or a voucher was created in error. Complete training before the pilot campaign.",
        10, 16, 28, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-34",
        "Analysts see reason codes. They do not get a gallery of biometric images in the MVP.",
    ),
    (
        "Customer support procedure and training",
        "Scripts for a first-time enrollment, camera and lighting help, resubmission coaching, and a winner asking where a physical prize is while the claim is held. The script uses the same wording for a fraud decline and a quality decline.",
        8, 14, 24, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-35",
        "Support can escalate. Support cannot release a prize alone.",
    ),
    (
        "Fulfillment and voucher-issuer procedure",
        "Walk the real warehouse screens and the voucher partner's screens. A pick, pack, or ship action requires a biometric release id. A voucher is minted only on the release event, and the partner demonstrates how a mistaken code is voided.",
        8, 14, 24, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-36",
        "This is a process change, not only an API change. Budget partner time on the client side.",
    ),
    (
        "Veriff cost reconciliation",
        "A weekly match of Veriff session volume to challenges and to released prizes, so finance can see cost per released prize during the pilot. The digital threshold is revisited from this report, not from the original illustration.",
        6, 10, 16, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-37",
        "The client enters unit fees on the cost model. This item is the operating routine.",
    ),
    (
        "Promotion terms, winner notices, and consent-copy cycle",
        "Workshops and implementation support while the client updates promotion rules, the winner message, and the consent notice. Engineering does not author the legal text. The maximum includes one rework cycle after counsel review.",
        8, 16, 28, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-38",
        "Go-live waits on the approved text. Track that dependency on the requirements tab.",
    ),
    (
        "Pilot control review",
        "Agree the scorecard before launch: enrollment completion, drop-off, decline mix, manual releases, contacts from winners who were declined in error, time to ship, and cost per released prize. Agree what evidence would later justify lowering the digital threshold.",
        6, 10, 16, "MVP", "Business operations", "Business operations", "Once", "Yes", "MS-39",
        "Phase 3 does not start from opinion. It starts from this scorecard.",
    ),
    (
        "Pilot hypercare",
        "Engineering effort for the pilot window: defect fixes, webhook issues, and an on-call rota. The hours are effort. They are not the calendar length of hypercare.",
        12, 24, 40, "MVP", "Quality", "Backend engineering", "Once", "Yes", "MS-40",
        "Pair with the business hypercare owners named by the client.",
    ),
    (
        "Second-brand onboarding kit",
        "A checklist, secret loading, a smoke test, and a tenant check so the next brand does not repeat MVP discovery. Pairs with the per-brand website line.",
        12, 20, 32, "Phase 2", "Scale", "Backend engineering", "Once", "Yes", "MS-41",
        "Build this from the pilot's actual steps, after the pilot has run.",
    ),
    (
        "Admin screens for threshold and campaign opt-in",
        "Authorised business users can change a threshold or opt a campaign in or out without a deployment. Every change stores the actor and the previous value.",
        16, 28, 44, "Phase 2", "Scale", "Backend engineering", "Once", "Yes", "MS-42",
        "MVP can run on audited engineering configuration if the client wants a smaller first release.",
    ),
    (
        "Fraud and operations dashboard",
        "A reviewed view of reason codes, brands, campaigns, override rate, and time to release. The dashboard does not show biometric images.",
        12, 24, 40, "Phase 2", "Business operations", "Backend engineering", "Once", "Yes", "MS-43",
        "Feeds the same scorecard started in the pilot.",
    ),
    (
        "Assisted redemption for winners who cannot complete capture",
        "Design a supervised alternative for an accessibility need or an unsupported document, with dual control and retained evidence. The design has to be at least as strict as the Veriff path, so the support channel does not become the easier door.",
        12, 20, 36, "Phase 2", "Business operations", "Business operations", "Once", "Yes", "MS-44",
        "Blocked until the client writes the policy on the requirements tab. Hours cover design, training, and the small workflow, not a call-centre operation.",
    ),
    (
        "Early enrollment from the win notification",
        "Let a winner start document-and-selfie enrollment from the win message, before redemption day, using the same endUserId. Enrollment alone does not release the prize. Redemption of an in-scope prize still requires a passing decision, which will be the short selfie if enrollment already succeeded.",
        8, 16, 28, "Phase 2", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-45",
        "This reduces redemption-day drop-off for physical prizes. It is optional for the MVP.",
    ),
    (
        "Staff review of Veriff status 9121",
        "Veriff sends review only when that setting is enabled, and the client's own staff must then approve or decline. Train the reviewers, minimise who can open a session, and put a clock on the hold. Leave this setting off in the MVP unless that team already exists.",
        8, 16, 28, "Phase 2", "Business operations", "Business operations", "Once", "Yes", "MS-46",
        "If a 9121 arrives unexpectedly in the MVP, the state machine holds the prize and does not release it.",
    ),
    (
        "Cross-account duplicate-face policy",
        "With Veriff and counsel, decide how reason 124 (face already in the collection) and the velocity reasons apply across accounts and, later, across brands. Implement the agreed hold or allow rule. A shared integration can treat the same legitimate person on two brands as a duplicate; resolve that before a second brand goes live if those checks are on.",
        16, 28, 44, "Phase 3", "Identity and policy", "Security and privacy", "Once", "Yes", "MS-47",
        "MVP already holds a prize when 124 is returned. This item decides the business rule and any extra matching.",
    ),
    (
        "Targeted extension to an abused low-value campaign",
        "A governed way to turn the biometric check on for one small-voucher campaign that the scorecard shows is being abused, while every other small voucher stays on the current path.",
        6, 12, 20, "Phase 3", "Identity and policy", "Discovery and design", "Once", "Yes", "MS-48",
        "Depends on MS-39. Do not enable this from launch volume alone.",
    ),
    (
        "Re-enrollment and subject repair",
        "A controlled recovery when an enrollment was deleted or the wrong document was bound. Re-enrollment cannot silently replace the person behind a winner account. A previous fraud decline, or an account that already redeemed, goes through the fraud procedure first.",
        12, 24, 40, "Phase 3", "Enrollment and authentication", "Backend engineering", "Once", "Yes", "MS-49",
        "Subject ids are still immutable. Repair links a reviewed new enrollment; it does not edit the old endUserId in place.",
    ),
]

BR = [
    (
        "Pilot redemption UX specification",
        "Specify the pilot screens for an in-scope prize and for a small voucher that stays on the current path. The first-time copy states that a photo ID and a selfie are required once, and that a later claim is a selfie. Include camera permission, finish later, and a held prize. Outcome is an approved UX spec.",
        12, 20, 32, "MVP", "Discovery", "Discovery and design", "Once", "Yes", "BW-01",
        "One pilot brand. Further brands are BW-20 after the shared kit exists.",
    ),
    (
        "Consent and explanation screens",
        "Show the counsel-approved notice, record accept or decline, and open Veriff only after acceptance. A decline leaves the prize unredeemed and tells the winner how to come back.",
        8, 12, 20, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-02",
        "The notice text is client-supplied. This item is the screen and the consent hand-off.",
    ),
    (
        "Enrollment UI for a winner with no biometric",
        "Open the enrollment session URL issued by the microservice and support document and selfie capture in the browser. Handle a denied camera permission, a closed browser, and the return to the site. Design this as the main launch path.",
        16, 28, 44, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-03",
        "Uses Veriff's web flow. Native SDK work is the optional BW-23 line.",
    ),
    (
        "Authentication UI for an enrolled winner",
        "Open the shorter selfie session. Tell the winner that the stored enrollment will be used and that a new document is not required for this claim.",
        12, 20, 36, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-04",
        "Shown only when the microservice reports an approved enrollment.",
    ),
    (
        "Status screens driven by the challenge API",
        "Render pending, resubmission, held, declined, expired, and released from the microservice status only. After Veriff redirects back, ignore any success hint in the URL. Keep the pending screen in place through the resubmission delay.",
        12, 20, 32, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-05",
        "A released status is the only state that may reveal a code or a dispatch confirmation.",
    ),
    (
        "Physical prize: show the locked recipient and address",
        "Show the name, address, and phone that will be checked and then shipped. Allow an edit only before the challenge starts, or after a failed challenge returns to draft. After release, changes follow the support procedure.",
        8, 16, 24, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-06",
        "The name on this screen is the name sent to Veriff for the document match when the client requires it.",
    ),
    (
        "High-value digital voucher reveal",
        "Render the code, barcode, or wallet action only after the challenge status is released. The pending page must not contain the code in the markup.",
        6, 12, 20, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-07",
        "Below-threshold vouchers are unchanged and are covered by the regression item.",
    ),
    (
        "Regression of small-voucher redemption",
        "Prove that a digital reward under the threshold still redeems as it does today and does not call the biometric service.",
        6, 10, 16, "MVP", "Quality", "Quality assurance", "Once", "Yes", "BW-08",
        "Include at least one regression case per pilot reward type that is out of scope.",
    ),
    (
        "User-facing error copy",
        "Plain language for a quality retry, a hold, and a generic unsuccessful check. The site uses one decline message for fraud and for a mismatch, so the page does not teach which rule fired. Support tools keep the finer reason.",
        6, 10, 18, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-09",
        "Copy is reviewed with support training so the two channels match.",
    ),
    (
        "Funnel analytics",
        "Record in-scope impression, consent, enrollment started, authentication started, abandoned, and released, with brand and campaign ids. The events contain no biometric payload and no document data.",
        6, 10, 16, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-10",
        "These events feed the pilot scorecard together with microservice metrics.",
    ),
    (
        "Locale, brand theme, and return URLs",
        "Apply the pilot language and theme to our pages, and set the Veriff locale and callback URLs. Veriff's own capture screens keep their product layout.",
        6, 10, 16, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-11",
        "Additional languages beyond the pilot set are extra and should be added as a change.",
    ),
    (
        "Winner-message deep link",
        "The win email or SMS opens this prize for the winning account. A forwarded link does not start a challenge until that account is signed in.",
        8, 12, 20, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-12",
        "Message content itself is the terms-and-notices item on the microservice tab.",
    ),
    (
        "Help content placed on the pilot site",
        "Place the approved FAQ: what will be asked the first time, lighting and camera tips, which documents are accepted, and what happens to the prize if the winner stops halfway.",
        6, 12, 20, "MVP", "Business operations", "Business operations", "Once", "Yes", "BW-13",
        "Legal text remains client-owned. Hours include one content revision.",
    ),
    (
        "Device matrix on real phones",
        "Exercise iOS Safari, Android Chrome, and the pilot's in-app webview if the brand opens the site inside an app. Camera permission and the return from Veriff are the cases that must pass on hardware, not only in a desktop emulator.",
        12, 20, 32, "MVP", "Quality", "Quality assurance", "Once", "Yes", "BW-14",
        "Confirm the webview in the channel-scope requirement. A native app is BW-23.",
    ),
    (
        "Integration into the existing rewards page",
        "Wire the pilot redemption page so it creates a challenge and branches on prize tier. Hours cover routing, the signed-in session, and coupling to that codebase. They do not repeat the screens in the enrollment, authentication, status, address, or voucher items.",
        16, 28, 48, "MVP", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-15",
        "Min assumes a modular page. Max assumes a tightly coupled legacy redemption template.",
    ),
    (
        "Pilot user-acceptance test",
        "Business and QA run enrolled, not-enrolled, name mismatch, held, expired, and small-voucher scripts against sandbox decisions. Engineering support during the test is included. A second build is not.",
        12, 20, 32, "MVP", "Quality", "Quality assurance", "Once", "Yes", "BW-16",
        "Depends on portal access and the business scripts signed in discovery.",
    ),
    (
        "Pilot support training and macros",
        "The pilot support team dry-runs a winner with no biometric, a declined check, and a winner asking where a physical prize is while the claim is held.",
        6, 10, 18, "MVP", "Business operations", "Business operations", "Once", "Yes", "BW-17",
        "Uses the master scripts from the microservice operations items.",
    ),
    (
        "Campaign operations checklist",
        "Show the pilot marketing team how a prize is marked physical, or how a digital face value is entered, so the website and the microservice apply the same rule.",
        4, 8, 14, "MVP", "Business operations", "Business operations", "Once", "Yes", "BW-18",
        "A prize with a missing type or value fails closed and stays unredeemed until data is fixed.",
    ),
    (
        "Warehouse floor card for the pilot",
        "A one-page aid for the pilot desk: a pick ticket without a biometric release id is not shipped. The card points at the master fulfillment procedure.",
        4, 6, 12, "MVP", "Business operations", "Business operations", "Once", "Yes", "BW-19",
        "Translate the card if the warehouse language differs from the website.",
    ),
    (
        "Each additional brand website",
        "Apply the pilot pattern to one more brand: theme, locales, callbacks, prize-tier wiring, analytics, acceptance test, and support macros. This does not build another microservice. If the brand runs a different rewards codebase, estimate it as a new pilot integration instead of this row.",
        16, 32, 56, "Phase 2", "Scale", "Frontend engineering", "Per additional brand", "Yes", "BW-20",
        "Enter the number of extra brands on the Phases sheet. The pilot is already inside the MVP total.",
    ),
    (
        "Shared front-end kit taken from the pilot",
        "After the pilot, package the status client, the event names, and a reference page so the next brand has one pattern to follow.",
        8, 16, 28, "Phase 2", "Scale", "Frontend engineering", "Once", "Yes", "BW-21",
        "Do this once, then use the per-brand row.",
    ),
    (
        "Holding pages for name mismatch and staff review",
        "A held state for a name mismatch and, if enabled, for Veriff review. The page has no self-serve bypass and tells the winner the claim is still theirs while it is checked.",
        6, 10, 16, "Phase 2", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-22",
        "MVP already has a generic held screen. This item is the specific copy and state once those policies are on.",
    ),
    (
        "Native iOS and Android apps, if they are in scope",
        "Integrate Veriff's native SDK, prepare the store copy for camera use, and run the release trains for each app. Leave this out of the estimate until the client confirms native apps redeem prizes. An in-app webview is already in the pilot device matrix.",
        24, 40, 72, "Phase 2", "Scale", "Frontend engineering", "Once", "No", "BW-23",
        "Excluded from the scenario totals. Set Include in rollup to Yes to pull it in, or add the Mid hours by hand.",
    ),
    (
        "Win-message prompt to enroll early",
        "From the win notification, offer enrollment before redemption day. The prize is still released only at redemption, after a passing decision.",
        6, 10, 16, "Phase 2", "Pilot website", "Frontend engineering", "Once", "Yes", "BW-24",
        "Pairs with the microservice early-enrollment item.",
    ),
    (
        "Opt-in experience for one abused small-voucher campaign",
        "Show the biometric step only on the campaign operations opted in. The copy explains why this reward asks for a check when other small vouchers do not.",
        6, 10, 18, "Phase 3", "Scale", "Frontend engineering", "Once", "Yes", "BW-25",
        "Depends on the Phase 3 policy item and the pilot scorecard.",
    ),
]

# requirements: id, requirement, why, needed before, recommended default
REQS = [
    (
        "CR-01",
        "A commercial agreement that includes Veriff selfie enrollment and Veriff Biometric Authentication, plus a volume expectation for the pilot.",
        "Authentication cannot be switched on against an ID-verification-only contract. Pricing also sets the digital-voucher threshold, because a check that costs more than the voucher is a poor default for high-volume rewards.",
        "Before discovery sign-off",
        "Confirm both products are in the order form. Enter the unit fees on the Cost model tab. Leave the fee cells blank in this file; do not send fees by email to the delivery team unless the client wants them stored there.",
    ),
    (
        "CR-02",
        "A named Veriff Solutions Engineer, and written confirmation they will create a selfie enrollment integration and a biometric authentication integration connected to it, in sandbox and production.",
        "Veriff documents this link as a prerequisite. The authentication integration can match a face only against an approved enrollment in the connected integration. The delivery team cannot create that link from API keys alone.",
        "Before build",
        "Enrollment type: document and selfie. Selfie-only enrollment is not recommended, because a stolen account could bind an attacker's face and then pass later checks.",
    ),
    (
        "CR-03",
        "Customer Portal access for sandbox and production, with the smallest role that can open sessions, set sandbox decisions, and read reason codes.",
        "Veriff's test integration does not return automatic decisions. Testers set the status in the portal. Production access is also how operations will investigate a held prize.",
        "Before sandbox",
        "Name two administrators and the investigators. Do not share one login. MVP does not require a wide group to view document or face images.",
    ),
    (
        "CR-04",
        "API keys, HMAC secrets, and base URLs for enrollment and for authentication, separately for sandbox and production, delivered through a secret channel.",
        "Session creation and webhook verification fail closed without them. Mixing the enrollment key with the authentication key, or sandbox with production, produces declines that look like user fraud (reasons 121, 122, and 128).",
        "Before sandbox",
        "Four credential sets. Store them in the client's secret manager. Do not paste them into tickets, chat, or this spreadsheet.",
    ),
    (
        "CR-05",
        "The stable winner-account identifier, confirmation it is never recycled when an account is closed, and agreement that we mint a UUID v4 endUserId from it.",
        "Veriff ties the biometric template to endUserId (UUID v4) or to vendorData. A changed or reused identifier declines with reason 128 or 121, or matches the wrong person. vendorData is immutable once sent.",
        "Before build",
        "Use endUserId. Keep one subject per brand and account in the MVP. Do not use email address as the identifier.",
    ),
    (
        "CR-06",
        "Whether any winners already have an approved Veriff document-and-selfie session, and whether those sessions may be migrated into enrollment.",
        "Migration can remove a second document capture for people already verified. It is a Solutions Engineer activity, and it still needs the same endUserId on the old session and the new one.",
        "Before build",
        "If there is no existing Veriff population, say so. The journey then assumes every winner starts with no biometric.",
    ),
    (
        "CR-07",
        "Accepted document types and countries, and the minimum age for a winner who redeems an in-scope prize.",
        "Veriff will request a resubmission or decline unsupported, expired, or out-of-policy documents (including reason codes such as 206 and 207). The website FAQ and the support script need the same list. Age is a legal rule, not an engineering guess.",
        "Before UAT",
        "Publish the list counsel and Veriff agree. Winners under the minimum age need a business outcome, such as a guardian process or an unscored prize rule.",
    ),
    (
        "CR-08",
        "Signed prize-tier rule: every physical prize is in scope; digital vouchers are in scope at or above a named face value and currency; a named role may change the value.",
        "The microservice fails closed when type or value is missing. Without a signed rule, every brand campaign will invent its own, and small vouchers will be pulled in by accident or high-value prizes will be left out.",
        "Before build",
        "Set the first number with finance using the Cost model. Revisit it after the pilot scorecard, not during a campaign.",
    ),
    (
        "CR-09",
        "The business definition of a legitimate winner for this program.",
        "A biometric match proves the person at the device is the person enrolled on that account. It proves they are the original entrant only when the document is also checked against a name the promotion already holds. These are different controls and different amounts of work.",
        "Before discovery sign-off",
        "MVP: bind the prize to the winning account, and match the document to the legal name on a physical claim and on any high-value digital claim where a name is collected. State in the terms that the winner is the account holder who passes that check.",
    ),
    (
        "CR-10",
        "Confirmation that the physical-prize shipping name, address, and phone may be frozen for the claim and that the shipping name may be sent to Veriff on the enrollment session.",
        "The name match (reason 104) is the control that ties the document to the person who will receive the parcel. If the name cannot be sent, a different person with a genuine document can enroll on a taken-over account.",
        "Before build",
        "Freeze the shipping payload when the biometric challenge starts. Later address changes go through support, after release, and are re-checked.",
    ),
    (
        "CR-11",
        "For a digital voucher at or above the threshold: whether the winner must enter a legal name before enrollment.",
        "Many voucher wins have only an email. Without a name, enrollment still stops a later account takeover and raises the cost of farming, but it does not tie the prize to a named entrant.",
        "Before build",
        "Collect the legal name for in-scope digital prizes and send it to Veriff, consistent with the physical-prize rule.",
    ),
    (
        "CR-12",
        "The policy for a winner who cannot complete camera or document capture.",
        "A policy left to the contact centre becomes the bypass. Fraud will use the bypass if it is easier than Veriff. Accessibility still needs a designed path.",
        "Before go-live",
        "MVP: the prize stays unredeemed and support follows the hold script. Phase 2: a dual-controlled assisted path (MS-44) once this policy is signed. There is no single-agent release.",
    ),
    (
        "CR-13",
        "System-of-record diagrams and sandbox APIs for the winner, the prize, voucher mint, voucher void, and shipment creation.",
        "The gate has to freeze and then release real downstream actions. Without a void and a hold, a mistaken code or parcel cannot be stopped, and the biometric result arrives too late.",
        "Before build",
        "Provide a sandbox where mint and ship can be called twice so exactly-once behaviour can be tested.",
    ),
    (
        "CR-14",
        "Written confirmation from the voucher issuer that a code is minted only on our release event, and a demonstration of the void path.",
        "If the issuer mints at the moment of the win, the biometric check no longer protects the code. The winner, or an attacker in the account, already has it.",
        "Before UAT",
        "Mint on release. Keep the code out of emails that are sent before the decision.",
    ),
    (
        "CR-15",
        "Written confirmation from warehouse operations that pick, pack, and ship will not run without a biometric release id.",
        "A parcel picked before the decision ships to the address on file even when the biometric check later fails. The control is then too late.",
        "Before UAT",
        "Add the release id to the pick ticket. Use the pilot floor card. Hold the stock in a 'won, not released' state.",
    ),
    (
        "CR-16",
        "How the brand website knows the session is the winning account, including session lifetime and any existing step-up.",
        "Veriff checks the face against the endUserId we send. If an anonymous or different user can start the challenge, we enroll or authenticate the wrong person onto the winner's prize.",
        "Before build",
        "The brand backend, not the browser, tells the microservice who the winner is, over the service credential.",
    ),
    (
        "CR-17",
        "The pilot brand: domains, return-URL owners, languages, and the first campaign that will use the check.",
        "The MVP is one brand and one campaign pattern. A vague 'all websites' scope removes the ability to rehearse warehouse, support, and Veriff together.",
        "Before discovery sign-off",
        "One pilot brand. Choose the campaign with physical prizes or higher-value vouchers, not the highest-volume small voucher.",
    ),
    (
        "CR-18",
        "The list of later brands, for scale planning only.",
        "Each extra brand uses the per-brand website row. A brand on a different codebase should be priced as another pilot integration. This list stops that surprise at phase 2.",
        "Before phase 2",
        "Leave these brands out of the MVP build. Enter the count on the Phases sheet when the pilot pattern is stable.",
    ),
    (
        "CR-19",
        "Channel scope: responsive website, in-app webview, native iOS, native Android.",
        "The base estimate uses Veriff's web flow in the browser. A native app needs the native SDK, store disclosures, and another test matrix, which is the optional line BW-23.",
        "Before discovery sign-off",
        "Assume responsive web, plus the pilot webview if one exists. Keep native apps off until confirmed.",
    ),
    (
        "CR-20",
        "Counsel-approved consent notice, privacy notice, retention period, regional go/no-go, and the Veriff data-processing terms.",
        "A face template is biometric data. Several privacy regimes require a defined purpose, a notice, and a retention limit before collection. The technical design stores decisions, not images, but counsel still has to approve the purpose and the regions.",
        "Before UAT",
        "Purpose: releasing an in-scope prize to the winning account. Retention: the shortest period that still covers a prize dispute. Regions that have not been cleared stay on the kill switch.",
    ),
    (
        "CR-21",
        "Who, if anyone, may view a document image or a face image.",
        "Opening images in the Veriff portal is a separate privacy decision from storing a reason code. A wide group with image access is a new data exposure and is not required to release a prize.",
        "Before go-live",
        "MVP: investigators see status and reason code in our tool. Image access stays limited to the minimum Veriff role needed for a disputed fraud case.",
    ),
    (
        "CR-22",
        "Two named roles for manual release, and the rule that one person cannot fill both roles.",
        "The dual-control API needs real people. A shared mailbox or a single team lead collapses the control the first time a winner pressures support.",
        "Before go-live",
        "Role A proposes, role B approves, neither is the winner's case owner alone. Sample the releases each week during the pilot.",
    ),
    (
        "CR-23",
        "Support and fraud operating hours, and the sentence the winner will see while a prize is held.",
        "A held physical prize creates contacts immediately. The copy and the staffing need to exist before the campaign, or the team will invent a release shortcut under volume.",
        "Before go-live",
        "Tell the winner the prize is reserved and under check, with the claim deadline. Do not promise same-day shipment once a check has started.",
    ),
    (
        "CR-24",
        "The claim deadline for a won prize, distinct from Veriff's seven-day session expiry.",
        "Veriff expires an unstarted session and abandons an unfinished one after seven days. That is a session rule. The business still has to say whether the prize entitlement remains until the campaign claim deadline, and how many times the winner may start again.",
        "Before build",
        "Entitlement lasts until the published claim deadline. A new Veriff session can be started until that deadline, except after a fraud-stop decline.",
    ),
    (
        "CR-25",
        "The fraud controls that already protect small vouchers and must keep running.",
        "The MVP deliberately leaves small vouchers off Veriff. If the project accidentally replaces velocity limits, device checks, or one-code rules, the high-volume rewards become weaker than they are today.",
        "Before UAT",
        "List the current controls and add them to the regression test. The biometric service is an added gate for in-scope prizes, not a replacement for account security.",
    ),
    (
        "CR-26",
        "Non-production winners, devices, and permission for testers to set sandbox decisions in the Veriff portal.",
        "End-to-end proof needs a real camera path and a controlled Veriff result. Production documents of real customers are not a test method.",
        "Before UAT",
        "Provide at least two pilot accounts and two phones (iOS and Android). Use sample or tester identities in sandbox only.",
    ),
    (
        "CR-27",
        "Security constraints: outbound access to Veriff, inbound webhook exposure, log retention, and any penetration-test expectation.",
        "The decision webhook must be reachable by Veriff and must be something we can verify. A blocked egress path or a forbidden public URL delays the pilot after the build is done. Logging rules decide whether reason codes may be stored.",
        "Before build",
        "Public HTTPS webhook on a host the client controls. Logs store ids and reason codes, not images or full document numbers.",
    ),
    (
        "CR-28",
        "The kill-switch owner and the incident contacts, including a Veriff outage.",
        "Fail closed means in-scope prizes pause when Veriff or the gate is unavailable. Someone must own the customer message and the decision to leave the switch off.",
        "Before go-live",
        "Name a business owner and an engineering owner. Default behaviour is to pause release, not to bypass the check.",
    ),
    (
        "CR-29",
        "Volume assumptions: monthly physical redemptions, in-scope digital redemptions, small digital redemptions, peak factor, and the share of winners with no enrollment.",
        "Capacity, support staffing, warehouse hold times, and the Veriff bill all depend on this split. The first month's 'not enrolled' share will be close to all winners if biometrics are new.",
        "Before discovery sign-off",
        "Enter the volumes on the Cost model. Plan support for a first month dominated by document-and-selfie enrollment.",
    ),
    (
        "CR-30",
        "A prize catalogue extract with type (physical or digital) and face value for the pilot campaign.",
        "The threshold rule needs real prizes. A catalogue that cannot distinguish a mug from a high-value voucher cannot be gated safely, because missing data fails closed and blocks legitimate winners.",
        "Before build",
        "Add an explicit prize type and numeric face value. Do not infer value from the prize name.",
    ),
    (
        "CR-31",
        "How long a decision record must be kept to defend a prize dispute, and who signs that period.",
        "The retention job will delete decision metadata when this period ends. Too short, and a dispute cannot be explained. Too long, and biometric-related records outlive the purpose.",
        "Before go-live",
        "Counsel sets the period. Engineering configures it. Images remain under the Veriff terms rather than in our database.",
    ),
    (
        "CR-32",
        "A question to Veriff: can duplicate-face (reason 124) and velocity or farming checks (108, 109, 110, 113) be enabled, and is the face collection scoped per integration or wider?",
        "These are the checks that stop one document or one face enrolling many winner accounts. If the collection is wider than one brand, a legitimate customer who wins on two brands can be declined once phase 2 adds brands.",
        "Before build",
        "For the single-brand MVP, enable the duplicate and farming checks if Veriff offers them. Resolve cross-brand scope before the second brand, in the Phase 3 policy item.",
    ),
    (
        "CR-33",
        "The owner for promotion-terms changes and the date those terms will be live relative to the first in-scope claim.",
        "Collecting a document and a face for a prize needs a purpose the winner can read in the rules and the notice. Shipping the feature before the terms is a legal defect, not a technical one.",
        "Before go-live",
        "Terms and consent are live before the first in-scope campaign opens. MS-38 covers delivery support, not the legal authoring.",
    ),
    (
        "CR-34",
        "Finance owner for the Veriff bill during the pilot.",
        "The threshold is a business control. Without an owner, nobody will compare session fees with prizes actually released, and the small-voucher exception will never be revisited with evidence.",
        "Before go-live",
        "Use the weekly reconciliation routine (MS-37) and the Cost model. The owner can lower or raise the threshold after the pilot review.",
    ),
    (
        "CR-35",
        "In-app webview user agent and camera behaviour, if the pilot site is opened inside a brand app.",
        "Webviews are a common place for the Veriff redirect or the camera to fail while mobile Safari works. Finding this in production creates a false decline wave and support volume.",
        "Before UAT",
        "If there is no app, record that. If there is an app webview, include it in BW-14. A native SDK remains optional.",
    ),
]

PREBUILD = [
    (
        "D1",
        "What does a legitimate winner mean?",
        "Account holder who passes biometrics, with the document matched to a name we already hold for physical and high-value digital prizes.",
        "A match to the original entrant without a name on file cannot be proven at redemption. Say so in the terms if the business accepts account binding only.",
    ),
    (
        "D2",
        "Which brand and campaign are the pilot?",
        "One brand. A campaign that includes physical prizes or higher-value vouchers.",
        "Small-voucher campaigns stay out of the pilot even if they are the loudest stakeholder.",
    ),
    (
        "D3",
        "What is the digital face-value threshold?",
        "A number finance sets from Veriff's fee, expected drop-off, and prize value. Physical prizes are always in scope.",
        "The Cost model is the worksheet. The engineering configuration must match the signed number.",
    ),
    (
        "D4",
        "Which documents and which minimum age are accepted?",
        "The list counsel and Veriff agree, published to support and to the FAQ.",
        "Winners who cannot meet it follow D6. They do not get a quiet exception.",
    ),
    (
        "D5",
        "Is a legal name required before a high-value digital voucher is enrolled?",
        "Yes. Collect it and send it with the enrollment session, as with the shipping name on a physical prize.",
        "Skipping the name is a conscious weakening of the control. Record it if the business chooses it.",
    ),
    (
        "D6",
        "What happens when a winner cannot complete biometrics?",
        "The prize stays unredeemed. Support uses the hold script. A dual-controlled assisted path is Phase 2, after a signed policy.",
        "A single support agent must not be able to release the prize.",
    ),
    (
        "D7",
        "Is Veriff review (status 9121) enabled?",
        "Off for the MVP, unless a trained review team already exists. An unexpected 9121 holds the prize.",
        "9121 means the client's staff decide. Veriff does not decide those sessions for you.",
    ),
    (
        "D8",
        "Are duplicate-face and farming checks enabled?",
        "Yes for the pilot brand, if Veriff can enable reasons such as 124, 108, 109, 110, and 113. Confirm the collection scope before brand two.",
        "A global face collection plus a second brand can decline a legitimate two-brand customer. That is a phase-2 decision, not an accident.",
    ),
    (
        "D9",
        "How long does the prize entitlement last?",
        "Until the published claim deadline. Veriff sessions still expire or abandon after seven days and may be restarted, except after a fraud-stop decline.",
        "Do not equate a seven-day Veriff expiry with loss of the prize.",
    ),
    (
        "D10",
        "Does the first approved enrollment release that same claim?",
        "Yes in the MVP. The document-and-selfie enrollment is the check for a winner who has no biometric yet. Later claims use selfie authentication against that enrollment.",
        "Requiring a second selfie in the same sitting is a Phase 2 tightening if the pilot scorecard shows hand-off fraud.",
    ),
]

FLOW = [
    (
        "1",
        "Brand website",
        "Winner asks to redeem a prize they have won.",
        "Is this prize in biometric scope? Physical prizes are in scope. A digital voucher is in scope when its face value is at or above the threshold, or the campaign was explicitly opted in.",
        "Continue to step 2.",
        "Use today's redemption. Do not call Veriff. Keep the existing small-voucher controls.",
        "No",
        "MVP",
    ),
    (
        "2",
        "Brand website",
        "The prize is in scope.",
        "Is the signed-in session the winning account, as asserted by the brand backend?",
        "Continue to step 3.",
        "Ask the winner to sign in to the winning account. Do not create a Veriff session for an anonymous or different user.",
        "No",
        "MVP",
    ),
    (
        "3",
        "Microservice",
        "The winning account is known.",
        "Is there already an open challenge for this prize?",
        "Resume that challenge. Do not create a second one.",
        "Create one challenge bound to brand, winner, prize, and a hash of the payout details.",
        "No",
        "MVP",
    ),
    (
        "4",
        "Microservice and brand",
        "A challenge exists.",
        "Can the payout be frozen? Physical: recipient name, address, and phone. Digital: block minting. If policy requires a legal name and none is stored, collect it now.",
        "Freeze succeeds. Continue to step 5.",
        "Fail closed. The prize stays won and unredeemed until the payout data can be frozen.",
        "No",
        "MVP",
    ),
    (
        "5",
        "Brand website",
        "Payout is frozen.",
        "Has this winner accepted the current biometric notice for this purpose?",
        "Continue to step 6.",
        "Show the notice. If the winner declines, close the challenge as consent declined. The prize stays unredeemed.",
        "No",
        "MVP",
    ),
    (
        "6",
        "Microservice",
        "Consent is on file.",
        "Does this endUserId already have an approved enrollment?",
        "Go to step 9, the selfie authentication path.",
        "Go to step 7, the enrollment path. This is the expected path at launch.",
        "No",
        "MVP",
    ),
    (
        "7",
        "Microservice and Veriff",
        "No approved enrollment.",
        "Create a document-and-selfie enrollment session with the enrollment credentials, the stored endUserId, and the legal name when we have one. The website opens the session URL.",
        "Winner is inside the Veriff flow. Wait for a decision at step 8.",
        "If the session cannot be created, fail closed and let the winner retry later. Do not release.",
        "No",
        "MVP",
    ),
    (
        "8",
        "Microservice",
        "An enrollment decision arrives by signed webhook or by reconciliation.",
        "What is verification.status, and the reason code when there is one?",
        "Approved (9001): mark the subject enrolled and go to step 11. Resubmission requested (9103): keep the same session URL and coach the winner; a quality resubmission is not a release. Review (9121): hold for staff, and only if that setting was enabled.",
        "Fraud-stop reasons 105, 106, 108, 109, 110, 113, 124, 127: hold, no self-serve retry. Name mismatch 104: operations hold. Other document declines: a limited new session, then operations. Expired or abandoned: the winner may start again until the claim deadline. Reason 122 on a later authentication: return to this enrollment step.",
        "Only after step 11",
        "MVP",
    ),
    (
        "9",
        "Microservice and Veriff",
        "An approved enrollment already exists.",
        "Create a biometric authentication session with the authentication credentials and the same endUserId. The website opens the selfie flow.",
        "Wait for the authentication decision at step 10.",
        "If session creation fails, fail closed. Do not reuse an old approval for a new prize.",
        "No",
        "MVP",
    ),
    (
        "10",
        "Microservice",
        "An authentication decision arrives.",
        "Is the status approved, and does the matched end user equal our subject?",
        "Yes: go to step 11. Resubmission, review, expiry, and abandonment follow the same handling as step 8.",
        "Approved but the matched end user differs: fraud hold, do not release. Reason 120, face does not match the reference: limited retry. Reasons 121 or 128: stop and raise an integration incident, no blind retry. Reason 122: go back to enrollment. Fraud-stop reasons: hold, no self-serve retry.",
        "Only after step 11",
        "MVP",
    ),
    (
        "11",
        "Microservice",
        "Veriff returned an approval for this challenge.",
        "Do brand, winner account, prize, challenge, endUserId, and the frozen payout hash still match?",
        "Go to step 12 and release once.",
        "Do not release. Raise a security alert. Leave the prize unredeemed.",
        "No",
        "MVP",
    ),
    (
        "12",
        "Microservice, then warehouse or voucher issuer",
        "Binding checks passed.",
        "Emit one signed release with an idempotency key.",
        "Physical: the warehouse may pick and ship the frozen address. Digital: the issuer mints the code, and only then may the website display it.",
        "A repeated callback, a second browser tab, or a replayed webhook must not create a second parcel or a second code.",
        "Yes, once",
        "MVP",
    ),
    (
        "13",
        "Microservice",
        "A dependency is missing or unhealthy.",
        "Is the webhook unsigned, the brand config missing, the endUserId missing, the kill switch on, or Veriff unavailable?",
        "Fail closed. Tell the winner they can retry. The prize entitlement remains until the claim deadline.",
        "This step is the exception path. A healthy signed approval continues through steps 11 and 12.",
        "No",
        "MVP",
    ),
    (
        "14",
        "Fraud or operations",
        "A person needs a decision the automated rules will not make.",
        "Is this a dual-controlled manual release with two different named roles, an allowed reason, and an evidence reference?",
        "Release once through the same state machine, so fulfillment still sees one release id.",
        "A winner-facing button, a single agent, or a shared mailbox cannot release the prize.",
        "Only after both approvals",
        "MVP",
    ),
    (
        "15",
        "Microservice",
        "The same winner later redeems another in-scope prize.",
        "Is the enrollment still approved, and is consent still current?",
        "Use step 9 only. Redemption day is a selfie.",
        "If consent was withdrawn or there is no enrollment, start again at step 5 or step 7. A previous fraud-stop decline is cleared only through the fraud procedure, not by a new enrollment.",
        "No",
        "MVP",
    ),
]

STATUSES = [
    (
        "approved",
        "9001",
        "Decision webhook",
        "Veriff verified the session. The session is finished.",
        "Move to binding checks. Release only when step 11 passes.",
        "None, unless binding fails.",
        "Only after binding checks",
    ),
    (
        "declined",
        "9102",
        "Decision webhook",
        "Veriff did not verify the session. Read the reason code before choosing a retry.",
        "Apply the reason-code table. A declined session is finished; a retry needs a new session, except where policy forbids a retry.",
        "Fraud-stop reasons go to the fraud queue. Name mismatch goes to the hold queue.",
        "No",
    ),
    (
        "resubmission_requested",
        "9103",
        "Decision webhook",
        "Something was missing or the capture quality was too low. The same session URL stays open.",
        "Keep the challenge pending and send the winner to the same session URL. For SDK sessions the webhook can be delayed by about five minutes, so the pending screen stays up.",
        "Support may coach lighting and document placement. Support does not release.",
        "No",
    ),
    (
        "expired",
        "9104",
        "Decision webhook",
        "The winner never entered the flow. Veriff expires the session seven days after creation. Code 9104 is shared with abandoned, so read the status text.",
        "Mark the challenge expired. The winner may start a new session until the claim deadline.",
        "No fraud case solely because a session expired.",
        "No",
    ),
    (
        "abandoned",
        "9104",
        "Decision webhook",
        "The winner started and did not finish. Veriff abandons the session seven days after creation. The status text distinguishes this from expired.",
        "Same as expired: entitlement remains, a new session is allowed until the claim deadline, except after a fraud-stop on a different attempt.",
        "No fraud case solely because a session was abandoned.",
        "No",
    ),
    (
        "review",
        "9121",
        "Decision webhook, only if enabled",
        "Veriff did not decide. The client's own staff must decide. This status is off unless Veriff enables it.",
        "Hold the prize. Do not treat review as an approval. MVP recommendation is to leave the setting off.",
        "If it arrives anyway, the hold queue owns it. Phase 2 describes a staffed review.",
        "No",
    ),
    (
        "started",
        "7001",
        "Event webhook, optional",
        "The winner opened the Veriff flow. There is no decision.",
        "Progress only. Ignore it for release.",
        "None.",
        "No",
    ),
    (
        "submitted",
        "7002",
        "Event webhook, optional",
        "The winner submitted captures. Veriff has not decided.",
        "Show pending. Release only from a later decision webhook or a decision poll.",
        "None.",
        "No",
    ),
]

REASONS = [
    ("104", "declined", "Name entered does not match the document.", "Hold for operations. Often a middle name, a married name, or a nickname. Not an automatic fraud ban.", "No", "MVP, whenever a name was sent"),
    ("105", "declined", "Suspicious behaviour, including a face shown from a screen or a print.", "Fraud hold. No self-serve retry.", "No", "MVP"),
    ("106", "declined", "Known fraud, including streamed images.", "Fraud hold. No self-serve retry.", "No", "MVP"),
    ("108", "declined", "Velocity or abuse: duplicated end-user.", "Fraud hold. Relevant to multi-account farming.", "No", "MVP when Veriff has the check enabled"),
    ("109", "declined", "Velocity or abuse: duplicated device.", "Fraud hold.", "No", "MVP when enabled"),
    ("110", "declined", "Velocity or abuse: duplicated ID.", "Fraud hold. Strong signal for many accounts using one document.", "No", "MVP when enabled"),
    ("113", "declined", "Suspicious behaviour: identity farming.", "Fraud hold.", "No", "MVP when enabled"),
    ("120", "declined", "The portrait does not match the reference photo.", "Limited self-serve retry, then a hold. This is the authentication mismatch.", "Limited", "MVP, authentication path"),
    ("121", "declined", "vendorData does not match the enrollment.", "Integration defect if vendorData is in use. Stop blind retries and alert engineering.", "No", "MVP"),
    ("122", "declined", "No enrollment reference for this endUserId or vendorData.", "Route to document-and-selfie enrollment. Do not show this as a fraud decline.", "Not a retry of the same path; start enrollment", "MVP"),
    ("124", "declined", "Face already exists in the face collection.", "Hold. Do not auto-release. Confirms a duplicate-face hit. Scope of the collection must be confirmed before a second brand.", "No", "MVP when enabled; policy in Phase 3"),
    ("127", "declined", "Face matches a blocklist.", "Fraud hold.", "No", "MVP when enabled"),
    ("128", "declined", "endUserId does not match the enrollment.", "Integration defect. Alert engineering. Do not mint a new endUserId to 'fix' it.", "No", "MVP"),
    ("539", "declined", "Resubmission limit exceeded.", "The session is finished. A further try, if policy allows, is a new session.", "Only via a new session, and not after a fraud-stop", "MVP"),
    ("201–217", "resubmission_requested", "Missing capture, glare, blur, unsupported or poorly framed document, and similar quality problems.", "Coach the winner and keep the same session URL. Examples include 202 face not clearly visible, 203 full document not visible, 206 document type not supported, 207 document expired, 214 selfie quality.", "Yes, inside the same session, up to Veriff's limit", "MVP"),
]


def style_input(cell):
    paint(cell, font_bold, fill_yellow, Alignment(vertical="center", horizontal="center"), thin, "#,##0.00")


def main():
    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Veriff biometric authentication — rewards redemption estimate"
    wb.properties.subject = "Solution work items for project estimation"
    wb.properties.creator = "Solution architecture planning baseline"
    wb.properties.description = (
        "Effort-hour estimate for a shared Veriff biometric microservice "
        "and brand-website redemption flows. Not a quote."
    )

    ws_start = wb.active
    ws_start.title = "Start Here"
    build_start(ws_start)

    ws_ms = wb.create_sheet("1. Bio Auth Microservice")
    build_items(
        ws_ms,
        "Tab 1 — Microservice: Veriff Biometric Authentication",
        "Shared service. All brands call it. It is the only component that may authorize a prize release.",
        MS,
        "MicroItems",
    )
    ws_br = wb.create_sheet("2. Brand Websites")
    build_items(
        ws_br,
        "Tab 2 — Brand websites",
        "Pilot brand experience, plus a repeatable cost for each later brand. Web flow, not native apps.",
        BR,
        "BrandItems",
    )
    ws_req = wb.create_sheet("3. Client Requirements")
    build_requirements(ws_req)
    ws_flow = wb.create_sheet("4. Decision Flow")
    build_flow(ws_flow)
    ws_phase = wb.create_sheet("5. Phases and MVP")
    build_phases(ws_phase)
    ws_cost = wb.create_sheet("6. Cost model")
    build_cost(ws_cost)
    ws_assume = wb.create_sheet("7. Assumptions")
    build_assumptions(ws_assume)

    # Cross-check printed for the authoring run.
    def tot(rows, phase=None, basis="Once", include="Yes"):
        mins = mids = maxs = 0
        for r in rows:
            if phase and r[5] != phase:
                continue
            if r[8] != basis or r[9] != include:
                continue
            mins += r[2]
            mids += r[3]
            maxs += r[4]
        return mins, mids, maxs

    print("MS MVP", tot(MS, "MVP"))
    print("MS P2", tot(MS, "Phase 2"))
    print("MS P3", tot(MS, "Phase 3"))
    print("BR MVP", tot(BR, "MVP"))
    print("BR P2 once", tot(BR, "Phase 2"))
    print("BR P2 per", tot(BR, "Phase 2", "Per additional brand"))
    print("BR P3", tot(BR, "Phase 3"))
    print("BR optional", tot(BR, "Phase 2", "Once", "No"))

    wb.save(OUT)
    print(OUT)


ITEM_HEADERS = [
    "Work items",
    "Description",
    "Min (hrs)",
    "Mid (hrs)",
    "Max (hrs)",
    "Phase",
    "Section",
    "Workstream",
    "Basis",
    "Include in rollup",
    "ID",
    "Notes",
]


def build_items(ws, title, subtitle, rows, table_name):
    apply_header_footer(ws, title)
    set_widths(
        ws,
        {
            "A": 46,
            "B": 78,
            "C": 12,
            "D": 12,
            "E": 12,
            "F": 14,
            "G": 28,
            "H": 24,
            "I": 24,
            "J": 18,
            "K": 12,
            "L": 42,
        },
    )
    write_title(ws, 1, 1, title, 12)
    ws.merge_cells("A2:L2")
    paint(ws.cell(2, 1, subtitle), font_muted, alignment=wrap)
    ws.row_dimensions[2].height = 20
    ws.merge_cells("A3:L3")
    paint(
        ws.cell(
            3,
            1,
            "Hours are effort hours for the work described, across the named workstream. Min is a clean estate and an experienced team. Mid is the planning number. Max is legacy coupling, extra review cycles, or unclear downstream APIs. Change a number and the phase totals below will follow.",
        ),
        font_small,
        fill_pale,
        wrap,
    )
    ws.row_dimensions[3].height = 32

    header_row = 5
    for col, header in enumerate(ITEM_HEADERS, 1):
        cell = ws.cell(header_row, col, header)
        paint(cell, font_white_s, fill_navy, Alignment(wrap_text=True, vertical="center", horizontal="center"), thin)
    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = "A6"

    for i, row in enumerate(rows):
        r = header_row + 1 + i
        values = [row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8], row[9], row[10], row[11]]
        stripe = fill_white if i % 2 == 0 else fill_pale
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            paint(cell, font_small, stripe, wrap, thin)
            if c in (3, 4, 5):
                cell.number_format = "0"
                cell.alignment = Alignment(vertical="center", horizontal="center")
                cell.font = font_bold
        phase_cell = ws.cell(r, 6)
        phase_cell.fill = PHASE_FILL.get(row[5], stripe)
        phase_cell.alignment = Alignment(vertical="center", horizontal="center")
        phase_cell.font = font_bold
        if row[9] == "No":
            ws.cell(r, 10).fill = fill_orange
        desc_lines = max(1, (len(row[1]) // 95) + 1)
        note_lines = max(1, (len(row[11]) // 48) + 1)
        item_lines = max(1, (len(row[0]) // 42) + 1)
        ws.row_dimensions[r].height = min(78, max(36, 15 * max(desc_lines, note_lines, item_lines)))

    last = header_row + len(rows)
    tab = Table(displayName=table_name, ref=f"A{header_row}:L{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
    ws.add_table(tab)

    # Rollup. Formulas follow any hour edited inside the table.
    roll = last + 2
    labels = [
        (roll + 1, "MVP, one-off (planning scope)"),
        (roll + 2, "Phase 2, one-off"),
        (roll + 3, "Phase 3, one-off"),
        (roll + 4, "Each additional brand (do not add until you count brands)"),
        (roll + 5, "Excluded optional rows"),
    ]
    filters = [
        ('"MVP"', '"Once"', '"Yes"'),
        ('"Phase 2"', '"Once"', '"Yes"'),
        ('"Phase 3"', '"Once"', '"Yes"'),
        ('"Phase 2"', '"Per additional brand"', '"Yes"'),
        ('"Phase 2"', '"Once"', '"No"'),
    ]
    section_bar(ws, roll, 1, 2, "Totals  ·  formulas follow edited hours")
    for col, text in enumerate(["Min (hrs)", "Mid (hrs)", "Max (hrs)"], 3):
        paint(ws.cell(roll, col, text), font_white_s, fill_teal, Alignment(horizontal="center", vertical="center"), thin)

    for (r, label), (phase, basis, include) in zip(labels, filters):
        paint(ws.cell(r, 1, label), font_bold, fill_pale_navy, left_mid, thin)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        paint(ws.cell(r, 2), fill=fill_pale_navy, border=thin)
        for col, hour_col in enumerate(["Min (hrs)", "Mid (hrs)", "Max (hrs)"], 3):
            formula = (
                f'=SUMIFS({table_name}[{hour_col}],{table_name}[Phase],{phase},'
                f'{table_name}[Basis],{basis},{table_name}[Include in rollup],{include})'
            )
            cell = ws.cell(r, col, formula)
            paint(cell, font_bold, fill_yellow, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        ws.row_dimensions[r].height = 22

    note_row = roll + 7
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    paint(
        ws.cell(
            note_row,
            1,
            "The Min–Max range is the uncertainty. Do not add a second contingency on top. "
            "Per-brand hours are excluded from the one-off totals; multiply them on the Phases tab. "
            "Rows with Include in rollup = No (native apps) stay out of the scenarios until you change that flag.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[note_row].height = 32
    ws.oddHeader.right.text = title
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.sheet_view.zoomScale = 110


def build_requirements(ws):
    apply_header_footer(ws, "Client requirements")
    ws.sheet_properties.tabColor = "C65911"
    set_widths(ws, {"A": 46, "B": 62, "C": 24, "D": 55, "E": 12})
    write_title(ws, 1, 1, "Tab 3 — Developer and delivery requirements from the client", 5)
    ws.merge_cells("A2:E2")
    paint(
        ws.cell(
            2,
            1,
            "These are inputs the delivery team cannot invent. A missing row blocks the phase in the Needed-before column. Recommended defaults are a starting position for the workshop, and the client can override them.",
        ),
        font_small,
        fill_pale,
        wrap,
    )
    ws.row_dimensions[2].height = 32
    headers = ["Requirements", "Why we need it", "Needed before", "Recommended default", "ID"]
    for col, header in enumerate(headers, 1):
        paint(
            ws.cell(4, col, header),
            font_white_s,
            fill_navy,
            Alignment(wrap_text=True, vertical="center", horizontal="center"),
            thin,
        )
    ws.row_dimensions[4].height = 22
    ws.freeze_panes = "A5"
    for i, (rid, req, why, when, default) in enumerate(REQS):
        r = 5 + i
        values = [req, why, when, default, rid]
        stripe = fill_white if i % 2 == 0 else fill_pale
        for c, value in enumerate(values, 1):
            paint(ws.cell(r, c, value), font_small, stripe, wrap, thin)
        when_fill = {
            "Before discovery sign-off": fill_orange,
            "Before build": fill_amber,
            "Before sandbox": fill_pale_teal,
            "Before UAT": fill_pale_navy,
            "Before go-live": fill_green,
            "Before phase 2": fill_pale,
        }.get(when, stripe)
        ws.cell(r, 3).fill = when_fill
        ws.cell(r, 3).font = font_bold
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        lines = max(len(req) // 52, len(why) // 70, len(default) // 62) + 1
        ws.row_dimensions[r].height = min(96, max(48, 15 * lines))
    last = 4 + len(REQS)
    tab = Table(displayName="ClientReqs", ref=f"A4:E{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    ws.add_table(tab)
    foot = last + 2
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=5)
    paint(
        ws.cell(
            foot,
            1,
            "Order of use: settle CR-09 (what 'legitimate winner' means), CR-08 (which prizes), CR-02 (Veriff configuration), and CR-17 (the pilot) before estimating the build at Mid versus Min. Credentials never belong in this workbook.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[foot].height = 32
    ws.print_title_rows = "1:4"


def build_flow(ws):
    apply_header_footer(ws, "Decision flow")
    ws.sheet_properties.tabColor = TEAL
    widths = {"A": 8, "B": 28, "C": 36, "D": 42, "E": 42, "F": 42, "G": 22, "H": 16}
    set_widths(ws, widths)
    write_title(ws, 1, 1, "Recommended process and decision flow", 8)
    ws.merge_cells("A2:H2")
    paint(
        ws.cell(
            2,
            1,
            "Use this as the build specification for the state machine and the operating procedures. The website displays the result. The microservice is the only place a release is authorized.",
        ),
        font_small,
        fill_pale,
        wrap,
    )
    ws.row_dimensions[2].height = 28

    section_bar(ws, 4, 1, 8, "One-screen flow")
    diagram = (
        "Winner asks to redeem\n"
        "        |\n"
        "        v\n"
        "Prize in biometric scope?  ---- no ---->  today's small-voucher redemption. Do not call Veriff.\n"
        "        | yes   (physical prize, or digital face value at or above the threshold)\n"
        "        v\n"
        "Signed in as the winning account?  ---- no ---->  sign in. Do not start Veriff.\n"
        "        | yes\n"
        "        v\n"
        "Freeze shipping details or block voucher mint. Collect a legal name if this prize requires one.\n"
        "        |\n"
        "        v\n"
        "Current biometric consent?  ---- no ---->  show the notice. A decline leaves the prize unredeemed.\n"
        "        | yes\n"
        "        v\n"
        "Approved enrollment for this endUserId?\n"
        "        | no                                      | yes\n"
        "        v                                         v\n"
        "Document + selfie enrollment                Selfie biometric authentication\n"
        "(first claim; expected at launch)           (later claims)\n"
        "        |                                         |\n"
        "        +--------------------+--------------------+\n"
        "                             v\n"
        "               Signed decision, then binding checks\n"
        "               (same winner, same prize, same frozen payout)\n"
        "                             |\n"
        "              fail closed <--+--> pass\n"
        "              prize stays         one release\n"
        "              unredeemed          physical: warehouse may ship\n"
        "                                  digital: issuer may mint, then the site may show the code\n"
    )
    ws.merge_cells("A5:H5")
    paint(ws.cell(5, 1, diagram), font_mono, fill_pale_teal, Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[5].height = 280

    section_bar(ws, 7, 1, 8, "Decisions to lock before build")
    pre_headers = ["ID", "Decision", "Recommended position", "If the business chooses otherwise"]
    for col, header in enumerate(pre_headers, 1):
        paint(ws.cell(8, col, header), font_white_s, fill_navy, Alignment(wrap_text=True, vertical="center", horizontal="center"), thin)
    ws.merge_cells("D8:H8")
    for col in range(4, 9):
        paint(ws.cell(8, col), fill=fill_navy, border=thin)
    for i, (rid, decision, rec, otherwise) in enumerate(PREBUILD):
        r = 9 + i
        paint(ws.cell(r, 1, rid), font_bold, fill_amber, Alignment(horizontal="center", vertical="center"), thin)
        paint(ws.cell(r, 2, decision), font_small, fill_white, wrap, thin)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        paint(ws.cell(r, 3, rec), font_small, fill_green, wrap, thin)
        for col in range(3, 6):
            paint(ws.cell(r, col), border=thin, fill=fill_green)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        paint(ws.cell(r, 6, otherwise), font_small, fill_orange, wrap, thin)
        for col in range(6, 9):
            paint(ws.cell(r, col), border=thin, fill=fill_orange)
        ws.row_dimensions[r].height = 48

    flow_header = 20
    section_bar(ws, flow_header, 1, 8, "Runtime steps")
    flow_cols = [
        "Step",
        "Actor",
        "Entry",
        "Question",
        "Continue",
        "Exception",
        "Release the prize?",
        "Phase",
    ]
    for col, header in enumerate(flow_cols, 1):
        paint(
            ws.cell(flow_header + 1, col, header),
            font_white_s,
            fill_navy,
            Alignment(wrap_text=True, vertical="center", horizontal="center"),
            thin,
        )
    ws.row_dimensions[flow_header + 1].height = 22
    for i, step in enumerate(FLOW):
        r = flow_header + 2 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        for c, value in enumerate(step, 1):
            paint(ws.cell(r, c, value), font_small, stripe, wrap, thin)
        release = step[6]
        if release.startswith("Yes") or release.startswith("Only after both"):
            ws.cell(r, 7).fill = fill_green
        elif release.startswith("Only after step"):
            ws.cell(r, 7).fill = fill_amber
        else:
            ws.cell(r, 7).fill = fill_red
        ws.cell(r, 7).font = font_bold
        ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[r].height = 72
    flow_last = flow_header + 1 + len(FLOW)
    tab = Table(displayName="DecisionSteps", ref=f"A{flow_header + 1}:H{flow_last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    ws.add_table(tab)

    status_row = flow_last + 2
    section_bar(ws, status_row, 1, 8, "Veriff status handling  ·  release never keys off the numeric code alone")
    status_headers = [
        "Status",
        "Code",
        "Arrives on",
        "Meaning in this program",
        "System action",
        "Operations action",
        "Release?",
        "",
    ]
    for col, header in enumerate(status_headers, 1):
        if header:
            paint(ws.cell(status_row + 1, col, header), font_white_s, fill_navy, Alignment(wrap_text=True, vertical="center", horizontal="center"), thin)
    ws.merge_cells(start_row=status_row + 1, start_column=7, end_row=status_row + 1, end_column=8)
    paint(ws.cell(status_row + 1, 8), fill=fill_navy, border=thin)
    for i, status in enumerate(STATUSES):
        r = status_row + 2 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        for c, value in enumerate(status, 1):
            paint(ws.cell(r, c, value), font_small, stripe, wrap, thin)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        release_fill = fill_amber if status[6].startswith("Only") else fill_red
        paint(ws.cell(r, 7), fill=release_fill, font=font_bold)
        paint(ws.cell(r, 8), border=thin, fill=release_fill)
        ws.cell(r, 7).font = font_bold
        ws.row_dimensions[r].height = 48

    reason_header = status_row + 2 + len(STATUSES) + 1
    section_bar(ws, reason_header, 1, 8, "Reason codes that change the path")
    reason_headers = ["Reason", "Status", "Veriff meaning", "Our handling", "Self-serve retry", "When"]
    for col, header in enumerate(reason_headers, 1):
        paint(ws.cell(reason_header + 1, col, header), font_white_s, fill_navy, Alignment(wrap_text=True, vertical="center", horizontal="center"), thin)
    ws.merge_cells(start_row=reason_header + 1, start_column=6, end_row=reason_header + 1, end_column=8)
    for col in range(6, 9):
        paint(ws.cell(reason_header + 1, col), fill=fill_navy, border=thin)
    for i, reason in enumerate(REASONS):
        r = reason_header + 2 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        for c, value in enumerate(reason, 1):
            paint(ws.cell(r, c, value), font_small, stripe, wrap, thin)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        for col in range(6, 9):
            paint(ws.cell(r, col), border=thin, fill=stripe)
        ws.cell(r, 1).font = font_bold
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 36
    note = reason_header + 2 + len(REASONS) + 1
    ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=8)
    paint(
        ws.cell(
            note,
            1,
            "Sources for the status and reason behaviour: Veriff developer documentation for Biometric Authentication, verification session status codes, and verification session decision codes, as published when this workbook was written. "
            "Confirm enabled checks, face-collection scope, and webhook settings with the Solutions Engineer during MS-07. Product settings can differ by integration.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[note].height = 36
    ws.print_title_rows = "1:2"
    ws.page_setup.fitToWidth = 1
    ws.freeze_panes = "A4"


def build_phases(ws):
    apply_header_footer(ws, "Phases and MVP")
    ws.sheet_properties.tabColor = "548235"
    set_widths(
        ws,
        {
            "A": 28,
            "B": 22,
            "C": 55,
            "D": 18,
            "E": 16,
            "F": 16,
            "G": 16,
            "H": 42,
        },
    )
    write_title(ws, 1, 1, "Phases and the MVP recommendation", 8)
    ws.merge_cells("A2:H2")
    paint(
        ws.cell(
            2,
            1,
            "Start with physical prizes and higher-value digital vouchers, on one brand. Small vouchers are the volume problem, and they are the wrong first place to put document capture. The shared microservice is built once so the next brand is a website integration, not a second platform.",
        ),
        font_small,
        fill_green,
        wrap,
    )
    ws.row_dimensions[2].height = 36

    section_bar(ws, 4, 1, 8, "Scenario calculator  ·  yellow cells are inputs  ·  hour cells are formulas")
    ws["A5"] = "Additional brands beyond the pilot"
    ws["C5"] = 0
    paint(ws["A5"], font_bold, fill_pale, left_mid, thin)
    ws.merge_cells("A5:B5")
    paint(ws["B5"], fill=fill_pale, border=thin)
    style_input(ws["C5"])
    ws["C5"].number_format = "0"
    ws["C5"].comment = Comment(
        "The pilot brand is already inside the MVP total. Enter how many further brands use the BW-20 pattern. A brand with a different codebase is not this number; price it as another pilot.",
        "Estimate",
        width=240,
        height=80,
    )
    paint(
        ws.cell(5, 4, "Pilot is already in the MVP. This number multiplies the 'each additional brand' row only."),
        font_muted,
        alignment=wrap,
    )
    ws.merge_cells("D5:H5")
    ws.row_dimensions[5].height = 28

    scenario_headers = [
        (1, "Scenario"),
        (2, "What it includes"),
        (3, "Use when"),
        (4, "Min (hrs)"),
        (5, "Mid (hrs)"),
        (6, "Max (hrs)"),
        (7, "Planning note"),
    ]
    for col, header in scenario_headers:
        paint(ws.cell(7, col, header), font_white_s, fill_navy, Alignment(horizontal="center", vertical="center", wrap_text=True), thin)
    ws.merge_cells("G7:H7")
    paint(ws.cell(7, 8), fill=fill_navy, border=thin)

    # Helper formulas pulling named tables.
    def sumifs(table, hour, phase, basis, include):
        return (
            f'SUMIFS({table}[{hour}],{table}[Phase],"{phase}",{table}[Basis],"{basis}",'
            f'{table}[Include in rollup],"{include}")'
        )

    def both(hour, phase, basis="Once", include="Yes"):
        return f"{sumifs('MicroItems', hour, phase, basis, include)}+{sumifs('BrandItems', hour, phase, basis, include)}"

    scenarios = [
        (
            "A. MVP pilot",
            "Phase 1 only. One brand. Physical prizes and digital vouchers at or above the threshold. Enrollment for winners with no biometric, selfie authentication after that, fail closed, dual control.",
            "The estimate to fund first.",
            both("Min (hrs)", "MVP"),
            both("Mid (hrs)", "MVP"),
            both("Max (hrs)", "MVP"),
            "Mid is the number to take into a quote. Move toward Min only after discovery shows a clean redemption API.",
        ),
        (
            "B. MVP plus Phase 2",
            "Scenario A, the shared Phase 2 items, and the additional brands entered above. Native apps stay out.",
            "After the pilot scorecard supports more brands.",
            f'{both("Min (hrs)", "MVP")}+{both("Min (hrs)", "Phase 2")}+{sumifs("BrandItems", "Min (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            f'{both("Mid (hrs)", "MVP")}+{both("Mid (hrs)", "Phase 2")}+{sumifs("BrandItems", "Mid (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            f'{both("Max (hrs)", "MVP")}+{both("Max (hrs)", "Phase 2")}+{sumifs("BrandItems", "Max (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            "If a later brand has a different codebase, replace one BW-20 unit with a BW-15-sized integration.",
        ),
        (
            "C. Through Phase 3",
            "Scenario B plus duplicate-identity policy, repair of enrollments, and the ability to opt in one abused small-voucher campaign.",
            "Only after the pilot has evidence of abuse or of cross-account collisions.",
            f'{both("Min (hrs)", "MVP")}+{both("Min (hrs)", "Phase 2")}+{both("Min (hrs)", "Phase 3")}+{sumifs("BrandItems", "Min (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            f'{both("Mid (hrs)", "MVP")}+{both("Mid (hrs)", "Phase 2")}+{both("Mid (hrs)", "Phase 3")}+{sumifs("BrandItems", "Mid (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            f'{both("Max (hrs)", "MVP")}+{both("Max (hrs)", "Phase 2")}+{both("Max (hrs)", "Phase 3")}+{sumifs("BrandItems", "Max (hrs)", "Phase 2", "Per additional brand", "Yes")}*C5',
            "Phase 3 is a policy and engineering increment. It is not a promise to check every small voucher.",
        ),
    ]
    for i, (name, what, when, mn, md, mx, note) in enumerate(scenarios):
        r = 8 + i
        paint(ws.cell(r, 1, name), font_bold, fill_green if i == 0 else fill_pale_navy, wrap_mid, thin)
        paint(ws.cell(r, 2, what), font_small, fill_white, wrap, thin)
        paint(ws.cell(r, 3, when), font_small, fill_white, wrap, thin)
        for col, formula in enumerate([mn, md, mx], 4):
            cell = ws.cell(r, col, f"={formula}")
            paint(cell, font_bold, fill_yellow, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        paint(ws.cell(r, 7, note), font_small, fill_pale, wrap, thin)
        paint(ws.cell(r, 8), border=thin, fill=fill_pale)
        ws.row_dimensions[r].height = 78

    # Component breakdown
    section_bar(ws, 12, 1, 8, "Component breakdown  ·  formulas")
    for col, header in enumerate(["Component", "Basis", "Phase", "Min (hrs)", "Mid (hrs)", "Max (hrs)"], 1):
        paint(ws.cell(13, col, header), font_white_s, fill_navy, Alignment(horizontal="center", vertical="center", wrap_text=True), thin)
    components = [
        ("Microservice, MVP", "Once", "MVP", "MicroItems"),
        ("Pilot brand website, MVP", "Once", "MVP", "BrandItems"),
        ("Microservice, Phase 2", "Once", "Phase 2", "MicroItems"),
        ("Shared website work, Phase 2", "Once", "Phase 2", "BrandItems"),
        ("Microservice, Phase 3", "Once", "Phase 3", "MicroItems"),
        ("Website, Phase 3", "Once", "Phase 3", "BrandItems"),
        ("Each additional brand website", "Per additional brand", "Phase 2", "BrandItems"),
        ("Native apps (excluded)", "Once", "Phase 2", "BrandItems"),
    ]
    for i, (name, basis, phase, table) in enumerate(components):
        r = 14 + i
        include = "No" if "excluded" in name else "Yes"
        stripe = fill_orange if include == "No" else (fill_white if i % 2 == 0 else fill_pale)
        paint(ws.cell(r, 1, name), font_bold, stripe, wrap_mid, thin)
        paint(ws.cell(r, 2, basis), font_small, stripe, Alignment(vertical="center", horizontal="center", wrap_text=True), thin)
        paint(ws.cell(r, 3, phase), font_small, PHASE_FILL.get(phase, stripe), Alignment(vertical="center", horizontal="center"), thin)
        for col, hour in enumerate(["Min (hrs)", "Mid (hrs)", "Max (hrs)"], 4):
            formula = f'={sumifs(table, hour, phase, basis, include)}'
            paint(ws.cell(r, col, formula), font_bold, stripe, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        ws.row_dimensions[r].height = 22

    # Workstream mix for MVP
    section_bar(ws, 23, 1, 8, "MVP effort by workstream  ·  so operations work is visible beside engineering")
    for col, header in enumerate(["Workstream", "Microservice mid (hrs)", "Brand website mid (hrs)", "MVP mid (hrs)", "Share of MVP mid"], 1):
        paint(ws.cell(24, col, header), font_white_s, fill_navy, Alignment(wrap_text=True, vertical="center", horizontal="center"), thin)
    ws.merge_cells("E24:H24")
    for col in range(5, 9):
        paint(ws.cell(24, col), fill=fill_navy, border=thin)
    workstreams = [
        "Discovery and design",
        "Backend engineering",
        "Frontend engineering",
        "Quality assurance",
        "Security and privacy",
        "Business operations",
    ]
    for i, name in enumerate(workstreams):
        r = 25 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        paint(ws.cell(r, 1, name), font_small, stripe, wrap_mid, thin)
        ms_formula = (
            f'=SUMIFS(MicroItems[Mid (hrs)],MicroItems[Workstream],A{r},MicroItems[Phase],"MVP",'
            f'MicroItems[Include in rollup],"Yes")'
        )
        br_formula = (
            f'=SUMIFS(BrandItems[Mid (hrs)],BrandItems[Workstream],A{r},BrandItems[Phase],"MVP",'
            f'BrandItems[Include in rollup],"Yes")'
        )
        paint(ws.cell(r, 2, ms_formula), font_small, stripe, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        paint(ws.cell(r, 3, br_formula), font_small, stripe, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        paint(ws.cell(r, 4, f"=B{r}+C{r}"), font_bold, stripe, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
        paint(ws.cell(r, 5, f'=IF($D$31=0,"—",D{r}/$D$31)'), font_bold, stripe, Alignment(horizontal="center", vertical="center"), thin, "0%")
        for col in range(6, 9):
            paint(ws.cell(r, col), border=thin, fill=stripe)
        ws.row_dimensions[r].height = 20
    total_r = 31
    paint(ws.cell(total_r, 1, "MVP total"), font_white, fill_total, wrap_mid, thin)
    for col, formula in enumerate(["=SUM(B25:B30)", "=SUM(C25:C30)", "=SUM(D25:D30)", "=SUM(E25:E30)"], 2):
        paint(ws.cell(total_r, col, formula), font_white, fill_total, Alignment(horizontal="center", vertical="center"), thin, "#,##0" if col < 5 else "0%")
    ws.merge_cells(start_row=total_r, start_column=5, end_row=total_r, end_column=8)
    for col in range(5, 9):
        paint(ws.cell(total_r, col), fill=fill_total, border=thin, font=font_white)
    ws.cell(total_r, 5).number_format = "0%"
    ws.cell(total_r, 5).value = "=SUM(E25:E30)"

    # Narrative phases
    section_bar(ws, 33, 1, 8, "What each phase is for")
    phase_rows = [
        (
            "Phase 1 — MVP",
            "Prove the control on prizes worth holding.",
            "One pilot brand. Every physical prize. Digital vouchers at or above the signed threshold. Document-and-selfie enrollment for a winner with no biometric. Selfie authentication on the next in-scope prize. Server-side release. Warehouse and voucher issuer hold. Dual-controlled manual release. Support, fraud, and finance routines.",
            "Small vouchers. Further brands. Native apps. Staff review of Veriff images. A second selfie immediately after enrollment. Cross-brand identity.",
            "CR-08, CR-09, CR-02, and CR-17 are signed. Sandbox credentials exist. Downstream mint and ship can be held.",
            "A pilot winner with no biometric can enroll and receive one in-scope prize only after a signed approval. A second prize requires a selfie match. A small voucher never calls Veriff. A failed check does not ship or mint. Support cannot release alone.",
        ),
        (
            "Phase 2 — More brands, less redemption-day friction",
            "Repeat the pattern and shorten the return visit.",
            "Onboarding kit, threshold admin, operations dashboard, early enrollment from the win message, holding pages for name mismatch, and the per-brand website row. Assisted redemption only with a signed policy. Staff review (9121) only with a trained team. Native apps only if confirmed.",
            "Turning the check on for every small voucher. Silent re-enrollment that swaps the person on an account.",
            "Pilot scorecard exists. The first brand's warehouse and support procedures have been used on a real campaign.",
            "A second brand integrates without a new microservice design. A winner who enrolled from the win message completes redemption with a selfie.",
        ),
        (
            "Phase 3 — Evidence-led tightening",
            "Close the gaps the pilot actually shows.",
            "Duplicate-face policy across accounts, and across brands only with a privacy decision. A switch to include one abused low-value campaign. A reviewed re-enrollment path.",
            "A blanket biometric check on all rewards.",
            "Scorecard evidence: abuse on a named small-voucher campaign, or duplicate-face collisions that the MVP hold does not already explain.",
            "The business can opt a campaign in, and can repair an enrollment, without opening a self-serve bypass.",
        ),
    ]
    labels = ["Phase", "Intent", "In scope", "Left out", "Start when", "Done when"]
    for col, header in enumerate(labels, 1):
        paint(ws.cell(34, col, header), font_white_s, fill_navy, Alignment(horizontal="center", vertical="center", wrap_text=True), thin)
    ws.merge_cells("F34:H34")
    for col in range(6, 9):
        paint(ws.cell(34, col), fill=fill_navy, border=thin)
    for i, prow in enumerate(phase_rows):
        r = 35 + i
        fills = [fill_green, fill_pale_navy, fill_amber][i]
        for c, value in enumerate(prow, 1):
            paint(ws.cell(r, c, value), font_small, fills if c == 1 else fill_white, wrap, thin)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        for col in range(6, 9):
            paint(ws.cell(r, col), border=thin, fill=fill_white)
        ws.cell(r, 1).font = font_bold
        ws.row_dimensions[r].height = 96

    ws.row_dimensions[34].height = 22
    ws.freeze_panes = "A8"
    ws.print_title_rows = "1:2"


def build_cost(ws):
    apply_header_footer(ws, "Cost model")
    ws.sheet_properties.tabColor = "BF8F00"
    set_widths(ws, {"A": 42, "B": 22, "C": 22, "D": 22, "E": 22, "F": 22, "G": 36, "H": 24})
    write_title(ws, 1, 1, "Cost model for the estimate workshop", 8)
    ws.merge_cells("A2:H2")
    paint(
        ws.cell(
            2,
            1,
            "Yellow cells are client inputs. Hour cells are drawn from the work-item tabs. This sheet does not contain Veriff's prices; the client types the fees they were quoted. Currency is whatever the client uses, and it must be the same in every money cell.",
        ),
        font_small,
        fill_amber,
        wrap,
    )
    ws.row_dimensions[2].height = 32

    section_bar(ws, 4, 1, 6, "Internal effort  ·  enter an hourly rate to see a cost")
    for col, header in enumerate(
        ["Workstream", "MVP mid (hrs)", "Your rate", "MVP mid cost", "Phase 2 mid (hrs)", "Phase 2 mid cost"],
        1,
    ):
        paint(ws.cell(5, col, header), font_white_s, fill_navy, Alignment(wrap_text=True, horizontal="center", vertical="center"), thin)
    workstreams = [
        "Discovery and design",
        "Backend engineering",
        "Frontend engineering",
        "Quality assurance",
        "Security and privacy",
        "Business operations",
    ]
    for i, name in enumerate(workstreams):
        r = 6 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        paint(ws.cell(r, 1, name), font_small, stripe, wrap_mid, thin)
        hours = (
            f'=SUMIFS(MicroItems[Mid (hrs)],MicroItems[Workstream],A{r},MicroItems[Phase],"MVP",MicroItems[Basis],"Once",MicroItems[Include in rollup],"Yes")'
            f'+SUMIFS(BrandItems[Mid (hrs)],BrandItems[Workstream],A{r},BrandItems[Phase],"MVP",BrandItems[Basis],"Once",BrandItems[Include in rollup],"Yes")'
        )
        p2 = (
            f'=SUMIFS(MicroItems[Mid (hrs)],MicroItems[Workstream],A{r},MicroItems[Phase],"Phase 2",MicroItems[Basis],"Once",MicroItems[Include in rollup],"Yes")'
            f'+SUMIFS(BrandItems[Mid (hrs)],BrandItems[Workstream],A{r},BrandItems[Phase],"Phase 2",BrandItems[Basis],"Once",BrandItems[Include in rollup],"Yes")'
        )
        paint(ws.cell(r, 2, hours), font_bold, stripe, Alignment(horizontal="center"), thin, "#,##0")
        rate = ws.cell(r, 3, None)
        style_input(rate)
        rate.number_format = "#,##0.00"
        paint(ws.cell(r, 4, f'=IF(C{r}="", "", B{r}*C{r})'), font_bold, stripe, Alignment(horizontal="center"), thin, "#,##0")
        paint(ws.cell(r, 5, p2), font_small, stripe, Alignment(horizontal="center"), thin, "#,##0")
        paint(ws.cell(r, 6, f'=IF(C{r}="", "", E{r}*C{r})'), font_small, stripe, Alignment(horizontal="center"), thin, "#,##0")
        ws.row_dimensions[r].height = 20
    paint(ws.cell(12, 1, "Total"), font_white, fill_total, wrap_mid, thin)
    paint(ws.cell(12, 2, "=SUM(B6:B11)"), font_white, fill_total, Alignment(horizontal="center"), thin, "#,##0")
    paint(ws.cell(12, 3, ""), font_white, fill_total, border=thin)
    paint(ws.cell(12, 4, '=IF(COUNT(C6:C11)=0,"",SUM(D6:D11))'), font_white, fill_total, Alignment(horizontal="center"), thin, "#,##0")
    paint(ws.cell(12, 5, "=SUM(E6:E11)"), font_white, fill_total, Alignment(horizontal="center"), thin, "#,##0")
    paint(ws.cell(12, 6, '=IF(COUNT(C6:C11)=0,"",SUM(F6:F11))'), font_white, fill_total, Alignment(horizontal="center"), thin, "#,##0")

    ws.merge_cells("A13:F13")
    paint(
        ws.cell(
            13,
            1,
            "Phase 2 hours here are the shared one-off items. Add brands with the row below. Native apps are excluded.",
        ),
        font_muted,
        alignment=wrap,
    )

    section_bar(ws, 15, 1, 6, "Additional brands")
    paint(ws.cell(16, 1, "Brands beyond the pilot"), font_bold, fill_pale, left_mid, thin)
    # Link to phases input so there is one place to type. Also allow override? Single source: Phases!C5
    paint(ws.cell(16, 2, "='5. Phases and MVP'!C5"), font_bold, fill_pale_navy, Alignment(horizontal="center"), thin, "0")
    paint(ws.cell(16, 3, "Taken from the Phases tab"), font_muted, alignment=wrap)
    ws.merge_cells("C16:F16")
    paint(ws.cell(17, 1, "Mid hours, all extra brands"), font_small, fill_white, left_mid, thin)
    paint(
        ws.cell(
            17,
            2,
            '=B16*SUMIFS(BrandItems[Mid (hrs)],BrandItems[Basis],"Per additional brand",BrandItems[Include in rollup],"Yes")',
        ),
        font_bold,
        fill_white,
        Alignment(horizontal="center"),
        thin,
        "#,##0",
    )
    paint(ws.cell(18, 1, "Blended front-end rate for those brands (optional)"), font_small, fill_white, left_mid, thin)
    style_input(ws.cell(18, 2, None))
    paint(ws.cell(19, 1, "Extra-brand cost"), font_bold, fill_pale, left_mid, thin)
    paint(ws.cell(19, 2, '=IF(B18="","",B17*B18)'), font_bold, fill_pale, Alignment(horizontal="center"), thin, "#,##0")

    section_bar(ws, 21, 1, 8, "Veriff fee illustration  ·  type the quoted fees  ·  this is not a price list")
    labels = [
        (22, "A22", "Fee per enrollment session (document and selfie)", "C22"),
        (23, "A23", "Fee per biometric authentication session (selfie)", "C23"),
        (24, "A24", "Monthly physical-prize redemptions", "C24"),
        (25, "A25", "Average physical prize value", "C25"),
        (26, "A26", "Monthly in-scope digital redemptions", "C26"),
        (27, "A27", "Average in-scope digital prize value", "C27"),
        (28, "A28", "Share of in-scope redemptions that are a first enrollment", "C28"),
        (29, "A29", "Monthly small-voucher redemptions left out of scope", "C29"),
    ]
    for row, _addr, label, _target in labels:
        paint(ws.cell(row, 1, label), font_small, fill_white, left_mid, thin)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        paint(ws.cell(row, 2), fill=fill_white, border=thin)
        cell = ws.cell(row, 3, None)
        style_input(cell)
        ws.row_dimensions[row].height = 20
    ws["C28"].number_format = "0%"
    ws["C28"] = 0.85
    ws["C28"].comment = Comment(
        "First month of a new biometric program is mostly enrollment. Later months fall as winners return. 85% is a starting illustration for month one, not a measured rate.",
        "Estimate",
        width=240,
        height=70,
    )
    for addr in ["C22", "C23", "C24", "C25", "C26", "C27", "C29"]:
        ws[addr].number_format = "#,##0.00"

    paint(ws.cell(31, 1, "In-scope redemptions per month"), font_bold, fill_pale_navy, left_mid, thin)
    ws.merge_cells("A31:B31")
    paint(ws.cell(31, 2), border=thin, fill=fill_pale_navy)
    paint(ws.cell(31, 3, '=IF(COUNT(C24,C26)<2,"",C24+C26)'), font_bold, fill_pale_navy, Alignment(horizontal="center"), thin, "#,##0")

    paint(ws.cell(32, 1, "Enrollment sessions"), font_small, fill_white, left_mid, thin)
    ws.merge_cells("A32:B32")
    paint(ws.cell(32, 3, '=IF(OR(C31="",C28=""),"",C31*C28)'), font_small, fill_white, Alignment(horizontal="center"), thin, "#,##0.0")

    paint(ws.cell(33, 1, "Authentication sessions"), font_small, fill_white, left_mid, thin)
    ws.merge_cells("A33:B33")
    paint(ws.cell(33, 3, '=IF(OR(C31="",C28=""),"",C31*(1-C28))'), font_small, fill_white, Alignment(horizontal="center"), thin, "#,##0.0")

    paint(ws.cell(34, 1, "Illustrative Veriff fee per month"), font_bold, fill_amber, left_mid, thin)
    ws.merge_cells("A34:B34")
    paint(ws.cell(34, 2), border=thin, fill=fill_amber)
    paint(
        ws.cell(34, 3, '=IF(OR(C22="",C23="",C32=""),"",C32*C22+C33*C23)'),
        font_bold,
        fill_amber,
        Alignment(horizontal="center"),
        thin,
        "#,##0",
    )

    paint(ws.cell(35, 1, "Face value of in-scope prizes per month"), font_small, fill_white, left_mid, thin)
    ws.merge_cells("A35:B35")
    paint(
        ws.cell(35, 3, '=IF(OR(C24="",C25="",C26="",C27=""),"",C24*C25+C26*C27)'),
        font_small,
        fill_white,
        Alignment(horizontal="center"),
        thin,
        "#,##0",
    )

    paint(ws.cell(36, 1, "Fee as a share of in-scope prize value"), font_bold, fill_green, left_mid, thin)
    ws.merge_cells("A36:B36")
    paint(ws.cell(36, 2), border=thin, fill=fill_green)
    paint(
        ws.cell(36, 3, '=IF(OR(C34="",C35="",C35=0),"",C34/C35)'),
        font_bold,
        fill_green,
        Alignment(horizontal="center"),
        thin,
        "0.0%",
    )

    ws.merge_cells("A38:H40")
    paint(
        ws.cell(
            38,
            1,
            "How to use the share. Physical prizes stay in scope even when a single check is a small fraction of a large prize, because a mis-shipped parcel is hard to reverse and the volume is lower. "
            "For digital vouchers, keep the threshold at a face value where this share, plus expected drop-off and support time, is still acceptable. "
            "Small vouchers are counted in C29 so the workshop can see the volume that would be added if someone proposes checking all of them. "
            "They are not multiplied into the fee until the business opts a campaign in. The 85% first-enrollment figure is an illustration for month one; replace it with the pilot measurement.",
        ),
        font_small,
        fill_pale,
        wrap,
    )
    ws.row_dimensions[38].height = 20
    ws.row_dimensions[39].height = 20
    ws.row_dimensions[40].height = 20
    ws.freeze_panes = "A4"


def build_assumptions(ws):
    apply_header_footer(ws, "Assumptions")
    ws.sheet_properties.tabColor = "7F7F7F"
    set_widths(ws, {"A": 28, "B": 100})
    write_title(ws, 1, 1, "Assumptions, boundaries, and how to read the hours", 2)
    ws.merge_cells("A2:B2")
    paint(
        ws.cell(
            2,
            1,
            "Read this tab before turning the Mid column into a commercial quote. The hours are a solution-architecture baseline for one pilot brand and a shared service.",
        ),
        font_small,
        fill_pale,
        wrap,
    )
    ws.row_dimensions[2].height = 28

    blocks = [
        (
            "What the hours mean",
            [
                "Each number is effort, in person-hours, for the workstream on that row. A 40-hour Mid is one person-week, or two people for half a week. It is not a calendar duration and it does not include waiting for Veriff, counsel, or a warehouse change board.",
                "Min assumes a known rewards API, one website stack, and reviewers who are available. Mid is the planning number while discovery is unfinished. Max covers a tightly coupled redemption page, extra legal review, or downstream systems that cannot hold a mint or a shipment cleanly.",
                "The range is the contingency. Adding another percentage on top double-counts uncertainty.",
                "Business-operations rows are in the same total as engineering on purpose. Training, warehouse rules, voucher-issuer rules, terms workshops, and the pilot scorecard are delivery work.",
            ],
        ),
        (
            "Recommended MVP boundary",
            [
                "One shared microservice performs enrollment checks, biometric authentication, and the release decision. Many brand websites can call it later. The MVP wires one pilot website.",
                "In scope for the check: every physical prize, and every digital voucher at or above a threshold finance sets. Out of scope: the long tail of small vouchers, which keep today's controls.",
                "A winner with no biometric is the normal first session. They complete Veriff document-and-selfie enrollment. An approved enrollment on that challenge can release that prize. The next in-scope prize is a selfie matched to the stored template.",
                "Selfie-only enrollment is not the recommendation. It would let a taken-over account bind a new face and then pass every later prize.",
                "The browser's return from Veriff never releases a prize. A signed decision plus binding checks do.",
            ],
        ),
        (
            "What 'a legitimate winner' is securing",
            [
                "The person redeeming is the signed-in winner of that prize, on that brand.",
                "The face presented now matches the face enrolled on that account, with liveness and presentation-attack checks as Veriff returns them.",
                "Where a legal name is on the prize (shipping name, or a name collected for a high-value voucher), the document matches that name.",
                "The payout details frozen at the start of the check are the details that ship or the account that receives the code.",
                "One face or one document used across many accounts is held when Veriff returns a duplicate or farming reason, and the business rule for cross-brand customers is explicit before brand two.",
                "A decline, an outage, a missing consent, or a killed switch leaves the prize unredeemed. It does not fall through to an unchecked redemption.",
            ],
        ),
        (
            "Included",
            [
                "Discovery, threat model, journeys, API contract, and the identity standard.",
                "The multi-tenant service, Veriff session handling, webhooks, reconciliation, the state machine, release callbacks, audit, and dual control.",
                "The pilot website flow, including the not-enrolled path, status screens, address lock, voucher reveal, analytics, and device tests.",
                "Support, fraud, warehouse, voucher-issuer, and finance procedures for the pilot, plus hypercare engineering effort.",
                "Phase 2 and Phase 3 items, priced separately so they are not silently inside the MVP.",
            ],
        ),
        (
            "Not included",
            [
                "Veriff session fees, minimum commitments, or face-collection add-ons. Enter quoted fees on the Cost model.",
                "Legal counsel fees, the privacy impact assessment writing, and the drafting of promotion terms. Delivery supports the cycle; counsel writes the text.",
                "Call-centre headcount, warehouse labour, parcel postage, and voucher face value.",
                "A second microservice per brand, a rewrite of the rewards catalogue, or a new account-login system.",
                "Native iOS and Android SDK integration, unless BW-23 is switched on.",
                "Penetration-test fees. Engineering support for a test the client commissions is inside the security-review item; the tester is not.",
                "Migration of a large historical image archive. If that archive exists, ask Veriff for a migration approach and raise a separate item; MS-07 only covers the design conversation.",
                "Production support after the pilot hypercare window.",
            ],
        ),
        (
            "Dependencies that move the date without moving the hours",
            [
                "Veriff creating and linking the enrollment and authentication integrations.",
                "Counsel approving the notice, the regions, and the retention period.",
                "The voucher issuer proving they can mint late and void a mistake.",
                "The warehouse agreeing to hold pick and pack.",
                "Portal access and a sandbox where decisions can be set by testers.",
            ],
        ),
        (
            "Product behaviour this estimate relies on",
            [
                "Biometric authentication requires an approved session in the connected selfie enrollment integration, and the same endUserId or vendorData on both.",
                "This estimate uses endUserId (UUID v4), generated and stored by the microservice.",
                "Decisions arrive on the decision webhook and can be read back with GET /sessions/{id}/decision. Signature verification is mandatory.",
                "Statuses used: approved 9001, declined 9102, resubmission_requested 9103, expired 9104, abandoned 9104, review 9121. Expired and abandoned share a code; the status text is authoritative.",
                "Review (9121) is sent only when Veriff enables it, and the client's staff make that decision.",
                "A resubmission continues on the same session URL. Veriff caps resubmissions and then returns reason 539. SDK resubmission webhooks can be delayed by about five minutes.",
                "Sessions expire or abandon after seven days. That timer is not the business claim deadline.",
                "Reason codes called out in the decision flow, including 104, 105, 106, 120, 121, 122, 124, and 128, follow Veriff's published tables. Which of the velocity and duplicate-face checks are enabled is an integration setting to confirm.",
            ],
        ),
        (
            "How to maintain the file",
            [
                "Edit hours on Tab 1 and Tab 2. Totals, the phase scenarios, the workstream mix, and the cost model are formulas.",
                "Type additional brands and rates only in the yellow cells.",
                "Leave Include in rollup as No on the native-app row until that scope is real.",
                "If a row is added, add it inside the Excel table so the formulas keep covering it. Do not paste a total over a formula.",
                "Do not store API keys, HMAC secrets, or production personal data in this workbook.",
            ],
        ),
    ]
    row = 4
    for title, lines in blocks:
        section_bar(ws, row, 1, 2, title)
        row += 1
        for line in lines:
            paint(ws.cell(row, 1, ""), font_small, fill_pale_teal, border=thin)
            paint(ws.cell(row, 2, line), font_small, fill_white, wrap, thin)
            ws.row_dimensions[row].height = min(72, max(32, 16 * ((len(line) // 110) + 1)))
            row += 1
        row += 1
    ws.freeze_panes = "A4"
    ws.print_title_rows = "1:2"
    ws.page_setup.orientation = "landscape"


def build_start(ws):
    apply_header_footer(ws, "Start here")
    ws.sheet_properties.tabColor = NAVY
    set_widths(ws, {"A": 28, "B": 88, "C": 24, "D": 22})
    write_title(ws, 1, 1, "Veriff biometric authentication for rewards redemption", 4)
    ws.merge_cells("A2:D2")
    paint(
        ws.cell(
            2,
            1,
            "Solution estimate for a shared biometric microservice used by brand websites. Physical prizes and digital vouchers. Planning baseline, 12 October 2026.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[2].height = 22

    section_bar(ws, 4, 1, 4, "Recommendation")
    recs = [
        "Build one microservice that owns Veriff enrollment, Veriff biometric authentication, and the decision to release a prize. Brand websites start the check and show the result. They do not authorize the prize.",
        "Fund the MVP as one pilot brand. Put every physical prize through the check, and put digital vouchers through it only at or above a face value finance sets from the real Veriff fee. Leave the high volume of small vouchers on today's redemption controls.",
        "Plan for winners who have never enrolled. That is the first-session path: consent, then document and selfie. An approval on that challenge can release that prize and stores the template. The next in-scope prize is a selfie match, which is the biometric authentication product.",
        "Treat a passed selfie as proof that this person matches the identity enrolled on the winning account. Tie that identity to the parcel or the voucher by sending the legal name already required for shipping, and by collecting a legal name for a high-value digital voucher. A check without a name still stops account takeover; it does not prove who originally entered.",
        "Fail closed. A missing webhook signature, a Veriff outage, a declined consent, or a kill switch leaves the prize won and unredeemed. A manual release exists for operations and requires two different named people.",
        "Add the next websites in Phase 2 against the same service. Open small-voucher campaigns in Phase 3 only when the pilot scorecard shows abuse on a named campaign.",
    ]
    for i, text in enumerate(recs):
        r = 5 + i
        paint(ws.cell(r, 1, f"{i + 1}"), font_white, fill_teal, Alignment(horizontal="center", vertical="center"), thin)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        rec_fill = fill_pale_teal if i % 2 == 0 else fill_white
        paint(ws.cell(r, 2, text), font_small, rec_fill, wrap, thin)
        for col in range(3, 5):
            paint(ws.cell(r, col), border=thin, fill=rec_fill)
        ws.row_dimensions[r].height = 36

    section_bar(ws, 12, 1, 4, "Where to go in this workbook")
    guide = [
        ("Start Here", "This page. The recommendation and the shape of the estimate."),
        ("1. Bio Auth Microservice", "Tab 1. Shared service work items, with Min, Mid, and Max hours. Full title: Microservice Veriff Biometric Authentication. Excel limits a tab name to 31 characters."),
        ("2. Brand Websites", "Tab 2. Pilot website work, business routines on the site, and a per-brand row for every website after the pilot."),
        ("3. Client Requirements", "Tab 3. What the client must provide, why it blocks the work, and a recommended default."),
        ("4. Decision Flow", "The runtime flow, the ten decisions to lock before build, Veriff statuses, and the reason codes that change the path."),
        ("5. Phases and MVP", "Scenario totals. Enter extra brands in the yellow cell. Mid on scenario A is the MVP planning number."),
        ("6. Cost model", "Drop in internal rates and the quoted Veriff fees. The sheet shows fee as a share of prize value so the threshold is a business decision."),
        ("7. Assumptions", "What is inside the hours, what is outside, and the Veriff behaviour the flow relies on."),
    ]
    paint(ws.cell(13, 1, "Tab"), font_white_s, fill_navy, Alignment(horizontal="center"), thin)
    ws.merge_cells("B13:D13")
    paint(ws.cell(13, 2, "What it is for"), font_white_s, fill_navy, Alignment(horizontal="center"), thin)
    for col in range(3, 5):
        paint(ws.cell(13, col), fill=fill_navy, border=thin)
    for i, (tab, desc) in enumerate(guide):
        r = 14 + i
        stripe = fill_white if i % 2 == 0 else fill_pale
        cell = ws.cell(r, 1, tab)
        paint(cell, font_bold, stripe, wrap_mid, thin)
        cell.hyperlink = f"#'{tab}'!A1"
        cell.font = Font(name="Calibri", size=11, bold=True, color="0B5CAB", underline="single")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        paint(ws.cell(r, 2, desc), font_small, stripe, wrap, thin)
        for col in range(3, 5):
            paint(ws.cell(r, col), border=thin, fill=stripe)
        ws.row_dimensions[r].height = 32

    section_bar(ws, 23, 1, 4, "MVP planning total")
    paint(ws.cell(24, 1, "Scenario A mid (hrs)"), font_bold, fill_green, left_mid, thin)
    ws.merge_cells("A24:B24")
    paint(ws.cell(24, 2), fill=fill_green, border=thin)
    paint(ws.cell(24, 3, "='5. Phases and MVP'!E8"), font_title, fill_yellow, Alignment(horizontal="center", vertical="center"), thin, "#,##0")
    paint(ws.cell(24, 4, "Open Phases for Min, Max, and later phases."), font_muted, alignment=wrap)
    ws.row_dimensions[24].height = 36
    ws.merge_cells("A25:D25")
    paint(
        ws.cell(
            25,
            1,
            "That figure is the pilot brand plus the shared microservice, including discovery, security, quality, and business operations. It updates if hours on Tab 1 or Tab 2 change.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[25].height = 28

    section_bar(ws, 27, 1, 4, "How the control meets the goal")
    controls = [
        ("Right account", "The brand backend names the winning account. A forwarded link cannot start a check for someone else."),
        ("Right person on that account", "First visit enrolls a document and a live face. Later visits match a new selfie to that enrollment, with the same endUserId."),
        ("Right name on the prize", "The shipping name, or a name collected for a high-value voucher, is sent with enrollment so a document mismatch is held."),
        ("Right payout", "Address and voucher minting are frozen before the check and released once, by the microservice, after the decision."),
        ("Hard stops", "Presentation attacks, known fraud, duplicate documents, duplicate faces, and an identifier mismatch do not get a self-serve retry."),
        ("No quiet bypass", "Small vouchers stay on their current controls. In-scope prizes have no unchecked fallback. Two people are required for a manual release."),
    ]
    for i, (title, body) in enumerate(controls):
        r = 28 + i
        paint(ws.cell(r, 1, title), font_bold, fill_pale_navy, wrap_mid, thin)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        paint(ws.cell(r, 2, body), font_small, fill_white, wrap, thin)
        for col in range(3, 5):
            paint(ws.cell(r, col), border=thin, fill=fill_white)
        ws.row_dimensions[r].height = 28

    ws.merge_cells("A35:D35")
    paint(
        ws.cell(
            35,
            1,
            "Nothing in this workbook is a deployment, a Veriff configuration, or a quote. Confirm integration settings with Veriff before build. Do not put secrets or customer data into the yellow input cells.",
        ),
        font_muted,
        alignment=wrap,
    )
    ws.row_dimensions[35].height = 28
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:2"
    ws.oddHeader.right.text = "Start here"


if __name__ == "__main__":
    main()
